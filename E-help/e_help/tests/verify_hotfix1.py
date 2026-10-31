"""
v1.5.2 Hotfix1 — 两项验收 Bug 修复验证测试

  Bug 1（T1-T8） : 动态字段数据丢失 → 按字段名+列索引读取，前两个必需字段必须识别
  Bug 2（T9-T15）: 重启 AttributeError → 状态变量先于 _load_output_dir_setting 初始化

  （Bug 3 自制文件夹选择器已于 Hotfix2 移除，回归测试见 verify_hotfix2.py）

运行:
  cd "Day07(小试牛刀)源文件"
  python -m e_help.tests.verify_hotfix1
"""

import os
import sys
import json
import shutil
import tempfile

# TEST-PHONE-* 是字段保留测试的合成占位内容，不是联系电话。
THIS_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(THIS_DIR)
PARENT_DIR = os.path.dirname(PROJECT_ROOT)
if PARENT_DIR not in sys.path:
    sys.path.insert(0, PARENT_DIR)

import tkinter as tk
from openpyxl import Workbook, load_workbook

from e_help import utils
from e_help.excel_reader import read_and_aggregate
from e_help.field_mapper import (
    detect_columns, save_field_config, reset_field_config,
)
from e_help.gui import SummaryApp

n_passed = 0
n_total = 0


def check(name, condition, detail=""):
    global n_passed, n_total
    n_total += 1
    if condition:
        n_passed += 1
        print(f"  PASS  {name}")
    else:
        print(f"  FAIL  {name}  {detail}")


TEST_DIR = os.path.join(PROJECT_ROOT, "tests", "_hotfix1_test")
if os.path.isdir(TEST_DIR):
    shutil.rmtree(TEST_DIR)
os.makedirs(TEST_DIR)


def make_xlsx(path, headers, rows):
    """创建测试用 .xlsx 文件。"""
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(path)
    wb.close()


print("=" * 55)
print("  v1.5.2 Hotfix1 — 两项验收 Bug 修复验证")
print("=" * 55)

# ══════════════════════════════════════════════════════════════
#  Bug 1: 动态字段数据丢失（T1-T8）
# ══════════════════════════════════════════════════════════════
print("\n[BUG 1] 动态字段数据丢失")

reset_field_config(TEST_DIR)
CONFIG_3 = [
    {"name": "客户", "keywords": ["客户", "客户名称", "公司"]},
    {"name": "金额", "keywords": ["金额", "总金额"]},
    {"name": "电话", "keywords": ["电话", "联系电话"]},
]

# T1: 3 字段（客户/金额/电话）干净输入 → 三个客户全部保留
print("\n[T1] 3 字段干净输入 — 全部客户保留")
save_field_config(TEST_DIR, CONFIG_3)
make_xlsx(
    os.path.join(TEST_DIR, "t1.xlsx"),
    ["客户", "金额", "电话"],
    [["公司A", 100.0, "TEST-PHONE-02"],
     ["公司B", 200.0, "TEST-PHONE-03"],
     ["公司C", 300.0, "TEST-PHONE-04"]],
)
total, count, fields = read_and_aggregate(TEST_DIR, TEST_DIR)
check("T1a: 字段顺序正确", fields == ["客户", "金额", "电话"], f"实际 {fields}")
check("T1b: 3 个客户全部保留", len(total) == 3, f"实际 {len(total)} keys={list(total.keys())}")
check("T1c: 公司A 金额正确", total.get("公司A", {}).get("金额") == 100.0, f"实际 {total.get('公司A')}")
check("T1d: 公司B 金额正确", total.get("公司B", {}).get("金额") == 200.0, f"实际 {total.get('公司B')}")
check("T1e: 公司C 金额正确", total.get("公司C", {}).get("金额") == 300.0, f"实际 {total.get('公司C')}")
check("T1f: 电话字段被保留", total.get("公司A", {}).get("电话") == "TEST-PHONE-02",
      f"实际 {total.get('公司A', {}).get('电话')}")
os.remove(os.path.join(TEST_DIR, "t1.xlsx"))

# T2: 4 字段（客户/金额/电话/日期）→ 4 字段全部保留
print("\n[T2] 4 字段输入 — 全部保留")
CONFIG_4 = [
    {"name": "客户", "keywords": ["客户", "客户名称", "公司"]},
    {"name": "金额", "keywords": ["金额", "总金额"]},
    {"name": "电话", "keywords": ["电话", "联系电话"]},
    {"name": "日期", "keywords": ["日期", "下单日期"]},
]
save_field_config(TEST_DIR, CONFIG_4)
make_xlsx(
    os.path.join(TEST_DIR, "t2.xlsx"),
    ["客户", "金额", "电话", "日期"],
    [["公司A", 100.0, "TEST-PHONE-02", "2026-01-01"],
     ["公司B", 200.0, "TEST-PHONE-03", "2026-01-02"]],
)
total, count, fields = read_and_aggregate(TEST_DIR, TEST_DIR)
check("T2a: 字段顺序正确", fields == ["客户", "金额", "电话", "日期"], f"实际 {fields}")
check("T2b: 2 个客户全部保留", len(total) == 2, f"实际 {len(total)}")
check("T2c: 日期字段被保留", total.get("公司A", {}).get("日期") == "2026-01-01",
      f"实际 {total.get('公司A', {}).get('日期')}")
check("T2d: 电话字段被保留", total.get("公司B", {}).get("电话") == "TEST-PHONE-03",
      f"实际 {total.get('公司B', {}).get('电话')}")
os.remove(os.path.join(TEST_DIR, "t2.xlsx"))

# T3: 2 字段默认（客户/金额）→ 仍正常工作
print("\n[T3] 2 字段默认输入")
CONFIG_2 = [
    {"name": "客户", "keywords": ["客户", "客户名称", "公司"]},
    {"name": "金额", "keywords": ["金额", "总金额"]},
]
save_field_config(TEST_DIR, CONFIG_2)
make_xlsx(
    os.path.join(TEST_DIR, "t3.xlsx"),
    ["客户", "金额"],
    [["公司A", 150.0], ["公司B", 250.0]],
)
total, count, fields = read_and_aggregate(TEST_DIR, TEST_DIR)
check("T3a: 字段顺序正确", fields == ["客户", "金额"], f"实际 {fields}")
check("T3b: 2 个客户保留", len(total) == 2, f"实际 {len(total)}")
check("T3c: 金额正确", total.get("公司A", {}).get("金额") == 150.0, f"实际 {total.get('公司A')}")
os.remove(os.path.join(TEST_DIR, "t3.xlsx"))

# T4: 表头列顺序打乱（电话/金额/客户）→ 按列名读取仍正确
print("\n[T4] 表头列顺序打乱 — 按列名读取")
save_field_config(TEST_DIR, CONFIG_3)
make_xlsx(
    os.path.join(TEST_DIR, "t4.xlsx"),
    ["电话", "金额", "客户"],
    [["TEST-PHONE-02", 100.0, "公司A"],
     ["TEST-PHONE-03", 200.0, "公司B"],
     ["TEST-PHONE-04", 300.0, "公司C"]],
)
total, count, fields = read_and_aggregate(TEST_DIR, TEST_DIR)
check("T4a: 3 个客户全部保留", len(total) == 3, f"实际 {len(total)} keys={list(total.keys())}")
check("T4b: 公司C 金额正确", total.get("公司C", {}).get("金额") == 300.0, f"实际 {total.get('公司C')}")
check("T4c: 公司A 电话正确", total.get("公司A", {}).get("电话") == "TEST-PHONE-02",
      f"实际 {total.get('公司A', {}).get('电话')}")
os.remove(os.path.join(TEST_DIR, "t4.xlsx"))

# T5: 必需字段「金额」缺失 → detect_columns 返回 None（不再错位）
print("\n[T5] 必需字段「金额」缺失 → 识别失败")
save_field_config(TEST_DIR, CONFIG_3)
make_xlsx(
    os.path.join(TEST_DIR, "t5.xlsx"),
    ["客户", "电话"],
    [["公司A", "TEST-PHONE-02"]],
)
wb = load_workbook(os.path.join(TEST_DIR, "t5.xlsx"), data_only=True)
det = detect_columns(wb.active)
wb.close()
check("T5: detect_columns 返回 None", det is None, f"实际 {det}")
os.remove(os.path.join(TEST_DIR, "t5.xlsx"))

# T6: 可选字段「电话」缺失 → 前两个字段仍识别成功
print("\n[T6] 可选字段「电话」缺失 → 仍正常识别")
save_field_config(TEST_DIR, CONFIG_3)
make_xlsx(
    os.path.join(TEST_DIR, "t6.xlsx"),
    ["客户", "金额"],
    [["公司A", 100.0]],
)
wb = load_workbook(os.path.join(TEST_DIR, "t6.xlsx"), data_only=True)
det = detect_columns(wb.active)
wb.close()
names = [d.field_name for d in (det or [])]
check("T6: 客户+金额 识别成功", det is not None and "客户" in names and "金额" in names,
      f"实际 {names}")
os.remove(os.path.join(TEST_DIR, "t6.xlsx"))

# T7: 多文件 3 字段聚合 → 同一客户跨文件求和
print("\n[T7] 多文件 3 字段聚合")
save_field_config(TEST_DIR, CONFIG_3)
make_xlsx(
    os.path.join(TEST_DIR, "t7a.xlsx"),
    ["客户", "金额", "电话"],
    [["公司A", 100.0, "TEST-PHONE-02"]],
)
make_xlsx(
    os.path.join(TEST_DIR, "t7b.xlsx"),
    ["客户", "金额", "电话"],
    [["公司A", 50.0, "TEST-PHONE-02"], ["公司B", 200.0, "TEST-PHONE-03"]],
)
total, count, fields = read_and_aggregate(TEST_DIR, TEST_DIR)
check("T7a: 2 个文件都被读取", count == 2, f"实际 {count}")
check("T7b: 公司A 跨文件求和 = 150", total.get("公司A", {}).get("金额") == 150.0,
      f"实际 {total.get('公司A')}")
check("T7c: 公司B 正确", total.get("公司B", {}).get("金额") == 200.0, f"实际 {total.get('公司B')}")
os.remove(os.path.join(TEST_DIR, "t7a.xlsx"))
os.remove(os.path.join(TEST_DIR, "t7b.xlsx"))

# T8: 电话列为非数字文本 → 不影响金额汇总、不丢数据
print("\n[T8] 电话列为非数字 → 不影响金额汇总")
save_field_config(TEST_DIR, CONFIG_3)
make_xlsx(
    os.path.join(TEST_DIR, "t8.xlsx"),
    ["客户", "金额", "电话"],
    [["公司A", 100.0, "无"],
     ["公司B", 200.0, "TEST-PHONE-01"]],
)
total, count, fields = read_and_aggregate(TEST_DIR, TEST_DIR)
check("T8a: 2 个客户保留", len(total) == 2, f"实际 {len(total)} keys={list(total.keys())}")
check("T8b: 公司A 金额正确", total.get("公司A", {}).get("金额") == 100.0, f"实际 {total.get('公司A')}")
check("T8c: 公司B 金额正确", total.get("公司B", {}).get("金额") == 200.0, f"实际 {total.get('公司B')}")
os.remove(os.path.join(TEST_DIR, "t8.xlsx"))

reset_field_config(TEST_DIR)

# ══════════════════════════════════════════════════════════════
#  Bug 2: 重启 AttributeError（T9-T15）
# ══════════════════════════════════════════════════════════════
print("\n[BUG 2] 重启 AttributeError（custom_output_dir）")

# 把 get_base_dir 指到临时目录，隔离真实数据目录
TMP = tempfile.mkdtemp(prefix="hotfix1_cfg_")
_orig_base_dir = utils.get_base_dir
utils.get_base_dir = lambda: TMP


def config_dir():
    return utils.get_config_dir()


def settings_path():
    return os.path.join(config_dir(), "settings.json")


def write_settings(obj):
    d = config_dir()
    os.makedirs(d, exist_ok=True)
    with open(os.path.join(d, "settings.json"), "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False)


def fresh_app():
    """创建隐藏 Tk root 并实例化 SummaryApp。"""
    root = tk.Tk()
    root.withdraw()
    try:
        app = SummaryApp(root)
        return root, app
    except Exception:
        root.destroy()
        raise


# T9 (场景 A): 无 settings.json（首次运行）→ 无异常，默认输出目录
print("\n[T9] 场景 A: 无 settings.json（首次运行）")
if os.path.isfile(settings_path()):
    os.remove(settings_path())
root, app = fresh_app()
check("T9a: 初始化无异常", app.custom_output_dir is not None)
check("T9b: 默认输出目录为空", app.custom_output_dir.get() == "",
      f"实际 '{app.custom_output_dir.get()}'")
root.destroy()

# T10 (场景 B): settings.json 存在且含有效 output_dir → 正确加载
print("\n[T10] 场景 B: 有效 output_dir 被加载")
out_dir = os.path.join(TMP, "我的输出")
os.makedirs(out_dir, exist_ok=True)
write_settings({"output_dir": out_dir})
root, app = fresh_app()
check("T10: 自定义目录已加载", app.custom_output_dir.get() == out_dir,
      f"实际 '{app.custom_output_dir.get()}' 期望 '{out_dir}'")
root.destroy()

# T11 (场景 C): 自定义 output_dir 持久化链（模拟重启再次加载）
print("\n[T11] 场景 C: 重启后自定义目录仍有效")
root, app = fresh_app()
check("T11: 重启后仍加载同一目录", app.custom_output_dir.get() == out_dir,
      f"实际 '{app.custom_output_dir.get()}'")
root.destroy()

# T12 (场景 D): 损坏的 settings.json → 无异常，静默回退
print("\n[T12] 场景 D: 损坏的 settings.json")
with open(settings_path(), "w", encoding="utf-8") as f:
    f.write("{ not valid json")
root, app = fresh_app()
check("T12: 损坏文件不崩溃且目录为空", app.custom_output_dir.get() == "",
      f"实际 '{app.custom_output_dir.get()}'")
root.destroy()

# T13 (场景 E): 删除 settings.json 后重启 → 无异常
print("\n[T13] 场景 E: 删除 settings.json 后重启")
if os.path.isfile(settings_path()):
    os.remove(settings_path())
root, app = fresh_app()
check("T13: 删除后无异常且目录为空", app.custom_output_dir.get() == "",
      f"实际 '{app.custom_output_dir.get()}'")
root.destroy()

# T14 (场景 F): settings.json 存在但无 output_dir 键 → 无异常
print("\n[T14] 场景 F: settings.json 无 output_dir 键")
write_settings({"field_mapping": {"某表头": "客户"}})
root, app = fresh_app()
check("T14: 无 output_dir 键不崩溃且目录为空", app.custom_output_dir.get() == "",
      f"实际 '{app.custom_output_dir.get()}'")
root.destroy()

# T15 (场景 G): output_dir 指向不存在目录 → 忽略该值
print("\n[T15] 场景 G: output_dir 指向不存在目录")
write_settings({"output_dir": os.path.join(TMP, "不存在的目录")})
root, app = fresh_app()
check("T15: 无效目录被忽略", app.custom_output_dir.get() == "",
      f"实际 '{app.custom_output_dir.get()}'")
root.destroy()

utils.get_base_dir = _orig_base_dir
shutil.rmtree(TMP, ignore_errors=True)

# 清理测试数据目录
shutil.rmtree(TEST_DIR, ignore_errors=True)

print("\n" + "=" * 55)
print(f"  {n_passed}/{n_total} 通过")
print("=" * 55)
if n_passed != n_total:
    sys.exit(1)
