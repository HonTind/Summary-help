"""
v1.4 Phase 1 — 多格式支持 + 文件扫描 验证测试

运行:
  cd "Day07(小试牛刀)源文件/e_help"
  python -m tests.verify_v14
"""

import os, sys, shutil, tempfile

THIS_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(THIS_DIR)
PARENT_DIR = os.path.dirname(PROJECT_ROOT)
if PARENT_DIR not in sys.path:
    sys.path.insert(0, PARENT_DIR)

from openpyxl import Workbook
import xlwt
import e_help.config as config_module
from e_help.excel_reader import (
    read_and_aggregate, _get_excel_files, _detect_columns_xls,
)

TEST_DIR = os.path.join(PROJECT_ROOT, "tests", "_v14_test")
if os.path.isdir(TEST_DIR):
    shutil.rmtree(TEST_DIR)
os.makedirs(TEST_DIR)

passed = 0
total = 0

def check(name, condition, detail=""):
    global passed, total
    total += 1
    if condition:
        passed += 1
        print(f"  PASS  {name}")
    else:
        print(f"  FAIL  {name}  {detail}")

print("=" * 55)
print("  v1.4 Phase 1 — 多格式兼容 验证测试")
print("=" * 55)

# ═══════════════════════════════════════════════════
#  T1: .xlsx
# ═══════════════════════════════════════════════════
print("\n[T1] .xlsx 基本兼容")
wb = Workbook()
ws = wb.active
ws.append(["客户", "金额"])
ws.append(["公司A", 1000.0])
xlsx_path = os.path.join(TEST_DIR, "t1.xlsx")
wb.save(xlsx_path)
wb.close()
sum_dict, count, fields = read_and_aggregate(TEST_DIR, TEST_DIR)
check("count=1", count == 1, f"实际: {count}")
check("sum=1000", sum_dict.get("公司A", {}).get("金额") == 1000.0, f"实际: {sum_dict}")
os.remove(xlsx_path)

# ═══════════════════════════════════════════════════
#  T2: .xls
# ═══════════════════════════════════════════════════
print("\n[T2] .xls 兼容")
wb2 = xlwt.Workbook()
ws2 = wb2.add_sheet("Sheet1")
ws2.write(0, 0, "客户")
ws2.write(0, 1, "金额")
ws2.write(1, 0, "公司B")
ws2.write(1, 1, 2000.0)
xls_path = os.path.join(TEST_DIR, "t2.xls")
wb2.save(xls_path)
sum_dict, count, fields = read_and_aggregate(TEST_DIR, TEST_DIR)
check("count=1", count == 1, f"实际: {count}")
check("sum=2000", sum_dict.get("公司B", {}).get("金额") == 2000.0, f"实际: {sum_dict}")
os.remove(xls_path)

# ═══════════════════════════════════════════════════
#  T3: .xlsm
# ═══════════════════════════════════════════════════
print("\n[T3] .xlsm 兼容")
wb3 = Workbook()
ws3 = wb3.active
ws3.append(["客户", "金额"])
ws3.append(["公司C", 3000.0])
xlsm_path = os.path.join(TEST_DIR, "t3.xlsm")
wb3.save(xlsm_path)
wb3.close()
sum_dict, count, fields = read_and_aggregate(TEST_DIR, TEST_DIR)
check("count=1", count == 1, f"实际: {count}")
check("sum=3000", sum_dict.get("公司C", {}).get("金额") == 3000.0, f"实际: {sum_dict}")
os.remove(xlsm_path)

# ═══════════════════════════════════════════════════
#  T4: 混合目录 (.xlsx + .xls + .xlsm)
# ═══════════════════════════════════════════════════
print("\n[T4] 混合目录")
# xlsx
wb4a = Workbook()
ws4a = wb4a.active
ws4a.append(["客户", "金额"]); ws4a.append(["公司A", 100.0])
wb4a.save(os.path.join(TEST_DIR, "a.xlsx")); wb4a.close()
# xls
wb4b = xlwt.Workbook()
ws4b = wb4b.add_sheet("S")
ws4b.write(0, 0, "客户"); ws4b.write(0, 1, "金额")
ws4b.write(1, 0, "公司B"); ws4b.write(1, 1, 200.0)
wb4b.save(os.path.join(TEST_DIR, "b.xls"))
# xlsm
wb4c = Workbook()
ws4c = wb4c.active
ws4c.append(["客户", "金额"]); ws4c.append(["公司C", 300.0])
wb4c.save(os.path.join(TEST_DIR, "c.xlsm")); wb4c.close()
sum_dict, count, fields = read_and_aggregate(TEST_DIR, TEST_DIR)
check("count=3", count == 3, f"实际: {count}")
# v1.5.0: sum_dict values are dicts {field_name: value}
total_sum = sum(item.get("金额", 0) for item in sum_dict.values())
check("sum=600", total_sum == 600.0, f"实际: {total_sum}")
for f in ["a.xlsx", "b.xls", "c.xlsm"]:
    os.remove(os.path.join(TEST_DIR, f))

# ═══════════════════════════════════════════════════
#  T5: ~$ 临时文件过滤
# ═══════════════════════════════════════════════════
print("\n[T5] ~$ 临时文件 + .开头的隐藏文件")
open(os.path.join(TEST_DIR, "~$hidden.xlsx"), "w").close()
open(os.path.join(TEST_DIR, ".gitkeep"), "w").close()
open(os.path.join(TEST_DIR, "readme.txt"), "w").close()
files = _get_excel_files(TEST_DIR)
check("~$ 被过滤", "~$hidden.xlsx" not in files)
check(". 隐藏被过滤", ".gitkeep" not in files)
check(".txt 被过滤", "readme.txt" not in files)
for f in ["~$hidden.xlsx", ".gitkeep", "readme.txt"]:
    os.remove(os.path.join(TEST_DIR, f))

# ═══════════════════════════════════════════════════
#  T6: 空目录
# ═══════════════════════════════════════════════════
print("\n[T6] 空目录")
empty = os.path.join(TEST_DIR, "_empty")
os.makedirs(empty, exist_ok=True)
files = _get_excel_files(empty)
check("空目录返回空列表", files == [], f"实际: {files}")
sum_dict, count, fields = read_and_aggregate(empty, TEST_DIR)
check("空目录 count=0", count == 0, f"实际: {count}")
shutil.rmtree(empty)

# ═══════════════════════════════════════════════════
#  T7: 损坏 Excel
# ═══════════════════════════════════════════════════
print("\n[T7] 损坏 Excel")
corrupt = os.path.join(TEST_DIR, "corrupt.xlsx")
with open(corrupt, "w") as f:
    f.write("not a zip file")
sum_dict, count, fields = read_and_aggregate(TEST_DIR, TEST_DIR)
check("损坏文件被找到但跳过（count=1, sum=0）",
      count == 1 and not sum_dict,
      f"实际: count={count}, sum={sum_dict}")
check("不崩溃", True)
os.remove(corrupt)

# ═══════════════════════════════════════════════════
#  T8: 非 Excel 文件 + 0字节锁定文件
# ═══════════════════════════════════════════════════
print("\n[T8] 非Excel + 0字节文件")
open(os.path.join(TEST_DIR, "notes.txt"), "w").close()
open(os.path.join(TEST_DIR, "locked.xlsx"), "w").close()  # 0 bytes
files = _get_excel_files(TEST_DIR)
check(".txt 不在列表中", "notes.txt" not in files)
check("0字节被过滤", "locked.xlsx" not in files, f"files={files}")
for f in ["notes.txt", "locked.xlsx"]:
    os.remove(os.path.join(TEST_DIR, f))

# ═══════════════════════════════════════════════════
#  T9: _detect_columns_xls 正确识别
# ═══════════════════════════════════════════════════
print("\n[T9] _detect_columns_xls 字段识别")
wb9 = xlwt.Workbook()
ws9 = wb9.add_sheet("S")
ws9.write(0, 0, "购货单位"); ws9.write(0, 1, "价税合计"); ws9.write(0, 2, "日期")
ws9.write(1, 0, "公司X"); ws9.write(1, 1, 999.0)
t9_path = os.path.join(TEST_DIR, "t9.xls")
wb9.save(t9_path)
import xlrd
wb9r = xlrd.open_workbook(t9_path)
m = _detect_columns_xls(wb9r.sheet_by_index(0))
# v1.5.0: returns list[FieldDetection] instead of ColumnMapping
check("custom: 购货单位→0", m is not None and len(m) >= 2 and m[0].col_index == 0,
      f"detections={m}")
check("custom: 价税合计→1", m is not None and len(m) >= 2 and m[1].col_index == 1,
      f"detections={m}")
os.remove(t9_path)

# T9b: 空表头
wb9b = xlwt.Workbook()
ws9b = wb9b.add_sheet("S")
t9b_path = os.path.join(TEST_DIR, "t9b.xls")
wb9b.save(t9b_path)
wb9br = xlrd.open_workbook(t9b_path)
m2 = _detect_columns_xls(wb9br.sheet_by_index(0))
check("空表头→None", m2 is None, f"实际: {m2}")
os.remove(t9b_path)

# ═══════════════════════════════════════════════════
#  T10: 回归 — 不影响已有功能
# ═══════════════════════════════════════════════════
print("\n[T10] 回归 — 已有功能")
from e_help.field_mapper import load_settings, save_field_mapping
_orig = dict(config_module.CUSTOM_HEADER_MAP)
config_module.CUSTOM_HEADER_MAP.clear(); config_module.CUSTOM_HEADER_MAP.update(_orig)

save_field_mapping(TEST_DIR, "test14", "customer")
check("field_mapper save ok",
      config_module.CUSTOM_HEADER_MAP.get("test14") == "customer")
config_module.CUSTOM_HEADER_MAP.clear(); config_module.CUSTOM_HEADER_MAP.update(_orig)

# ═══════════════════════════════════════════════════
print(f"\n{'=' * 55}")
print(f"  {passed}/{total} PASS")
print(f"{'=' * 55}")

shutil.rmtree(TEST_DIR)
