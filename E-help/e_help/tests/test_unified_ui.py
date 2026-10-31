"""Unified shell acceptance: synthetic local files, isolated paths and denied network calls."""
import ast
import hashlib
import importlib
import importlib.util
import json
from pathlib import Path
import socket
import sys
import tempfile
import threading
import time
import tkinter as tk
from tkinter import ttk
import unittest
import urllib.request
import webbrowser
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

PROJECT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT))
ui = importlib.import_module("e_help.desktop_ui")
help_content = importlib.import_module("e_help.offline_help")
workbench = importlib.import_module("e_help.workbench_gui")
monthly = importlib.import_module("e_help.monthly")
utils = importlib.import_module("e_help.utils")
task_history = importlib.import_module("e_help.history")
spec = importlib.util.spec_from_file_location("local_privacy_audit", PROJECT / "tools" / "audit_repository.py")
audit_module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(audit_module)


class UnifiedUiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ehelper-unified-")
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        def no_network(*_args, **_kwargs):
            raise AssertionError("Unexpected network/external browser access during startup or operation")
        for target, name in ((socket.socket, "connect"), (socket, "create_connection"),
                             (urllib.request, "urlopen"), (webbrowser, "open")):
            mock = patch.object(target, name, side_effect=no_network)
            mock.start()
            self.addCleanup(mock.stop)
        for target in (utils, ui.gui):
            mock = patch.object(target, "get_base_dir", return_value=str(self.folder))
            mock.start()
            self.addCleanup(mock.stop)
        for name, value in (("showinfo", None), ("showerror", None), ("showwarning", None), ("askyesno", True)):
            mock = patch.object(ui.messagebox, name, return_value=value)
            mock.start()
            self.addCleanup(mock.stop)
        self.root = tk.Tk()
        self.root.withdraw()
        self.errors = []
        self.root.report_callback_exception = lambda *error: self.errors.append(error)
        self.app = ui.UnifiedApp(self.root)
        self.root.update()
        self.addCleanup(self.cleanup)

    def cleanup(self):
        self.app.workbench.runner.cancel()
        if self.app.workbench.runner.thread:
            self.app.workbench.runner.thread.join(3)
        try:
            self.root.destroy()
        except tk.TclError:
            pass
        self.assertFalse(self.app.workbench.runner.busy)
        self.assertFalse(self.errors, self.errors)

    def file(self, name="input.xlsx", rows=(("A", "2026-10-01", "0.1"), ("A", "2026-09-01", "0.2"))):
        path = self.folder / name
        book = Workbook()
        book.active.append(("客户", "日期", "金额"))
        for row in rows:
            book.active.append(row)
        book.save(path)
        book.close()
        return path

    def drain(self):
        deadline = time.monotonic() + 6
        while self.app.workbench._busy and time.monotonic() < deadline:
            self.root.update()
            time.sleep(0.005)
        self.root.update()
        self.assertFalse(self.app.workbench._busy)
        self.assertFalse(self.errors)

    def preview(self):
        path = self.file()
        self.app.select_mode("monthly")
        self.app.workbench.monthly_input.set_files([path])
        self.app.workbench.month.set("2026-10")
        self.app.workbench.preview()
        self.drain()
        return path

    def test_one_root_all_pages_and_empty_contacts(self):
        for page in ("workbench", "history", "rules", "help", "feedback", "about"):
            self.assertTrue(self.app.show_page(page))
            self.root.update()
            self.assertEqual(self.app.active_page, page)
        self.assertEqual(help_content.PUBLIC_CONTACTS, ())
        self.assertEqual(self.app.contact_label.cget("text"), "联系方式暂未提供")
        self.assertFalse(any(isinstance(child, tk.Toplevel) for child in self.root.winfo_children()))
        self.assertFalse(any(isinstance(child, (tk.Entry, tk.Text, ttk.Entry))
                             for child in self.app.pages["feedback"].winfo_children()))

    def test_regular_view_keeps_state_and_legacy_engine(self):
        self.assertTrue(self.app.select_mode("regular"))
        self.app.legacy.selected_folder.set(str(self.folder))
        self.app.legacy.selected_excel_files = ["chosen.xlsx"]
        self.assertTrue(self.app.select_mode("monthly"))
        self.assertTrue(self.app.select_mode("regular"))
        self.assertEqual(self.app.legacy.selected_excel_files, ["chosen.xlsx"])
        self.assertEqual(self.app.legacy.selected_folder.get(), str(self.folder))
        self.assertIs(ui.EmbeddedSummary._handle_worker_msg, ui.gui.SummaryApp._handle_worker_msg)
        self.assertIs(ui.EmbeddedSummary._rescan_selected_files, ui.gui.SummaryApp._rescan_selected_files)

    def test_monthly_preview_export_and_date_filter_unchanged(self):
        path = self.preview()
        window = self.app.workbench
        self.assertEqual(str(window.result.total), "0.3")
        before = hashlib.sha256(path.read_bytes()).digest()
        window.date_filter.set(True)
        window._date_toggle()
        window.date_field.set("日期")
        self.assertIsNone(window.result)
        window.preview()
        self.drain()
        self.assertEqual(str(window.result.total), "0.1")
        self.assertEqual(window.result.report.excluded_rows, 1)
        with patch.object(workbench.filedialog, "askdirectory", return_value=str(self.folder)):
            window.export()
            self.drain()
        output = next(self.folder.glob("月度汇总-2026-10*.xlsx"))
        book = load_workbook(output)
        try:
            self.assertEqual(book.sheetnames, ["月度汇总", "来源明细", "异常清单", "运行配置"])
        finally:
            book.close()
        self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), before)

    def test_mode_change_clears_monthly_total_hint(self):
        self.preview()
        self.assertIn("金额合计", self.app.workbench.empty_hint.cget("text"))
        self.app.select_mode("compare")
        self.root.update()
        self.assertIsNone(self.app.workbench.result)
        self.assertNotIn("金额合计", self.app.workbench.empty_hint.cget("text"))

    def test_both_reconciliation_modes_run_in_same_window(self):
        left = self.file("left.xlsx", rows=(("A", "D1", "1"), ("B", "D1", "2")))
        right = self.file("right.xlsx", rows=(("A", "D1", "1"), ("B", "D1", "3")))
        self.app.select_mode("compare")
        window = self.app.workbench
        for panel, path in ((window.left_input, left), (window.right_input, right)):
            panel.set_files([path])
            panel.extra_keys.set(json.dumps(["日期"]))
        for mode in ("客户总额", "单据逐笔"):
            window.mode.set(mode)
            window.mode_changed()
            window.preview()
            self.drain()
            self.assertEqual(window.result.counts["consistent"], 1)
            self.assertEqual(window.result.counts["different"], 1)
        self.assertEqual([record["task_type"] for record in task_history.load_history()], ["customer", "document"])
        self.assertTrue(all(record["file_count"] == 2 for record in task_history.load_history()))

    def test_monthly_metadata_preview_export_and_incomplete_states(self):
        self.preview()
        window = self.app.workbench
        record = task_history.load_history()[0]
        self.assertEqual((record["task_type"], record["status"], record["result_count"]), ("monthly", "previewed", 1))
        self.assertEqual(record["output_file"], "")
        self.assertNotIn("input_files", record)
        self.assertNotIn("amount", record)
        with patch.object(workbench.filedialog, "askdirectory", return_value=str(self.folder)):
            window.export()
            self.drain()
        saved = task_history.load_history()[-1]
        self.assertEqual((saved["status"], saved["phase"]), ("success", "export"))
        self.assertTrue(Path(task_history.output_path(saved)).is_file())
        invalid = self.file("invalid.xlsx", rows=(("SYNTHETIC_PRIVATE_CUSTOMER", "2026-10-01", "not-money"),))
        window.monthly_input.set_files([invalid])
        window.preview()
        self.drain()
        self.assertEqual(task_history.load_history()[-1]["status"], "incomplete")
        self.assertEqual(task_history.load_history()[-1]["issue_count"], len(window.result.report.issues))
        with patch.object(workbench.filedialog, "askdirectory", return_value=str(self.folder)):
            window.export()
            self.drain()
        self.assertIn("不完整", task_history.status_label(task_history.load_history()[-1]))
        text = Path(task_history.get_history_path()).read_text(encoding="utf-8")
        self.assertNotIn("SYNTHETIC_PRIVATE_CUSTOMER", text)
        self.assertNotIn("invalid.xlsx", text)

    def test_legacy_metadata_is_minimal_and_terminal_callbacks_do_not_duplicate(self):
        legacy = self.app.legacy
        legacy._task_input_folder = "SYNTHETIC_PRIVATE_FOLDER"
        legacy._task_input_files = ["SYNTHETIC_PRIVATE_INPUT"]
        legacy._record_task("success", 1, ["SYNTHETIC_PRIVATE_FIELD"], str(self.folder / "saved.xlsx"),
                            "SYNTHETIC_PRIVATE_ERROR", result_count=3)
        legacy._record_task("failed", 0, [], "", "repeat")
        records = task_history.load_history()
        self.assertEqual(len(records), 1)
        self.assertEqual((records[0]["task_type"], records[0]["result_count"]), ("regular", 3))
        self.assertNotIn("SYNTHETIC_PRIVATE", json.dumps(records))
        legacy._task_history_written = False
        legacy._record_task("failed", 0, [], "", "SYNTHETIC_PRIVATE_ERROR")
        self.assertEqual(task_history.load_history()[-1]["file_count"], 1)

    def test_error_cancel_late_cancel_and_flush_record_once(self):
        window = self.app.workbench
        def fail(cancel, progress):
            raise RuntimeError("SYNTHETIC_PRIVATE_DIAGNOSTIC")
        window._start(fail, "preview")
        self.drain()
        self.assertEqual(task_history.load_history()[-1]["status"], "failed")
        def wait_for_cancel(cancel, progress):
            cancel.wait(2)
            workbench.check_cancel(cancel)
        window._start(wait_for_cancel, "preview")
        window.cancel()
        self.drain()
        self.assertEqual(task_history.load_history()[-1]["status"], "cancelled")
        result = monthly.run_monthly([self.file()], monthly.make_template("T", window.monthly_input.rules()), "2026-10")
        window._start(lambda cancel, progress: result, "preview")
        window.runner.thread.join(2)
        window.cancel()
        self.drain()
        self.assertEqual(task_history.load_history()[-1]["status"], "cancelled")
        window.flush_history()
        window.flush_history()
        self.assertEqual(len(task_history.load_history()), 3)
        self.assertNotIn("SYNTHETIC_PRIVATE_DIAGNOSTIC", json.dumps(task_history.load_history()))

    def test_late_cancel_after_export_return_still_records_actual_saved_file(self):
        self.preview()
        window = self.app.workbench
        saved_file = self.file("saved-result.xlsx")
        def saved_then_cancelled(cancel, progress):
            cancel.set()  # The completed exporter already returned its committed path.
            return str(saved_file)
        window._start(saved_then_cancelled, "export")
        self.drain()
        record = task_history.load_history()[-1]
        self.assertEqual((record["status"], record["phase"]), ("success", "export"))
        self.assertEqual(task_history.output_path(record), str(saved_file))
        self.assertTrue(saved_file.exists())

    def test_corrupted_history_retains_bytes_and_does_not_break_preview(self):
        path = Path(task_history.get_history_path())
        path.write_text("{broken synthetic history", encoding="utf-8")
        original = path.read_bytes()
        self.preview()
        self.assertEqual(str(self.app.workbench.result.total), "0.3")
        self.assertIn("记录未保存", self.app.workbench.status.get())
        self.app.show_page("history")
        self.root.update()
        self.assertTrue(self.app.history_notice.cget("text"))
        self.assertEqual(path.read_bytes(), original)

    def test_history_filter_open_and_missing_output(self):
        path = self.file("local-result.xlsx")
        task_history.add_record({"id": "legacy", "status": "success", "output_file": path.name,
                                  "output_folder": str(self.folder), "file_count": 1})
        task_history.save_task("monthly", "previewed", phase="preview", result_count=1)
        task_history.save_task("document", "success", phase="export", output_path=str(path))
        self.app.show_page("history")
        self.root.update()
        self.assertEqual(len(self.app.history_records), 3)
        self.app.history_filter.set("单据逐笔对账")
        self.app.refresh_history()
        self.assertEqual(len(self.app.history_records), 1)
        tree = self.app.history_tree.tree
        tree.selection_set(tree.get_children()[0])
        self.root.update()
        self.assertEqual(str(self.app.history_open_button.cget("state")), "normal")
        self.assertIn(str(self.folder), self.app.history_notice.cget("text"))
        self.assertNotIn("完整导出", self.app.history_tree.label.cget("text"))
        with patch.object(ui.os, "startfile") as open_file:
            self.app.open_history_result()
            open_file.assert_called_once_with(str(path))
        with patch.object(ui.os, "startfile", side_effect=OSError("synthetic association failure")), \
                patch.object(ui.messagebox, "showwarning") as warning:
            self.app.open_history_result()
            warning.assert_called_once()
        path.unlink()  # Only this synthetic test file, never a user output.
        with patch.object(ui.os, "startfile") as open_file, patch.object(ui.messagebox, "showinfo") as info:
            self.app.open_history_result()
            open_file.assert_not_called()
            info.assert_called_once()
        self.app.history_filter.set("月度汇总")
        self.app.refresh_history()
        tree.selection_set(tree.get_children()[0])
        self.root.update()
        self.assertEqual(str(self.app.history_open_button.cget("state")), "disabled")

    def test_history_second_page_selection_opens_correct_record(self):
        records = [{"id": str(index), "task_type": "regular", "status": "success", "time": "synthetic",
                    "output_folder": str(self.folder), "output_file": f"saved-{index}.xlsx"}
                   for index in range(102)]
        self.assertTrue(task_history._save(records))
        self.app.show_page("history")
        self.app.history_tree.move(1)
        tree = self.app.history_tree.tree
        tree.selection_set(tree.get_children()[0])
        self.root.update()
        self.assertEqual(self.app._selected_history()["id"], "1")
        output = self.folder / "saved-1.xlsx"
        output.write_bytes(b"synthetic test placeholder; not opened by Excel")
        with patch.object(ui.os, "startfile") as open_file:
            self.app.open_history_result()
            open_file.assert_called_once_with(str(output))

    def test_manage_dialog_shows_types_and_preview_status(self):
        task_history.save_task("monthly", "previewed", phase="preview", result_count=3)
        dialog = ui.gui.HistoryDialog(self.root)
        try:
            self.root.update()
            values = dialog.tree.item(dialog.tree.get_children()[0], "values")
            self.assertIn("月度汇总", values)
            self.assertIn("预览完成", values)
        finally:
            dialog.destroy()

    def test_safe_close_flushes_terminal_metadata_without_restarting_task(self):
        window = self.app.workbench
        def wait_for_cancel(cancel, progress):
            cancel.wait(2)
            workbench.check_cancel(cancel)
        window._start(wait_for_cancel, "preview")
        self.app.close()
        window.runner.thread.join(2)
        self.app._wait_close()
        records = task_history.load_history()
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["status"], "cancelled")

    def test_ordered_field_rows_and_template_roundtrip(self):
        panel = self.app.workbench.monthly_input
        panel.add_key()
        panel._extra_variables[0].set("单据号,批次")
        self.assertEqual(panel.rules().field_selections(), ("客户", "单据号,批次"))
        rules = panel.rules()
        panel.remove_key(0)
        self.assertEqual(panel.rules().field_selections(), ("客户",))
        panel.apply_rules(rules)
        self.assertEqual(panel._extra_variables[0].get(), "单据号,批次")
        window = self.app.workbench
        window.template_name.set("组合规则")
        window.save_template()
        panel.key.set("变更")
        window.load_template()
        self.assertEqual(panel.rules().field_selections(), ("客户", "单据号,批次"))

    def test_global_writer_lock_help_reading_and_return(self):
        path = self.file()
        window = self.app.workbench
        window.monthly_input.set_files([path])
        window.month.set("2026-10")
        release = threading.Event()
        original = workbench.run_monthly

        def held(*args, **kwargs):
            release.wait(3)
            return original(*args, **kwargs)

        with patch.object(workbench, "run_monthly", side_effect=held):
            window.preview()
            self.assertTrue(self.app.busy())
            self.assertFalse(self.app.select_mode("regular"))
            self.assertFalse(self.app.show_page("rules"))
            self.assertTrue(self.app.show_page("help"))
            self.app.open_article_mode()
            self.assertEqual(self.app.active_page, "help")
            self.assertTrue(self.app.show_page("workbench"))
            with patch.object(ui.gui, "_WorkerThread") as worker:
                self.app.legacy._on_start_summary()
                worker.assert_not_called()
            self.assertTrue(window.monthly_input.disabled)
            old_fields = window.monthly_input.extra_keys.get()
            window.monthly_input.add_key()
            self.assertEqual(window.monthly_input.extra_keys.get(), old_fields)
            release.set()
            self.drain()
        self.assertTrue(self.app.select_mode("regular"))

    def test_legacy_busy_rejects_workbench_programmatic_start(self):
        window = self.app.workbench
        window.monthly_input.set_files([self.file()])
        self.app.legacy.task_running = True
        try:
            window.preview()
            self.assertFalse(window.runner.busy)
            self.assertFalse(window._busy)
        finally:
            self.app.legacy.task_running = False

    def test_help_search_and_all_main_actions_without_network(self):
        def denied(*_args, **_kwargs):
            raise AssertionError("Unexpected network or external browser access")
        with patch.object(socket.socket, "connect", side_effect=denied), \
                patch.object(socket, "create_connection", side_effect=denied), \
                patch.object(urllib.request, "urlopen", side_effect=denied), \
                patch.object(webbrowser, "open", side_effect=denied):
            for page in ("help", "feedback", "about", "rules", "history", "workbench"):
                self.app.show_page(page)
                self.root.update()
            self.app.show_page("help")
            self.app.help_query.set("虚构")
            self.assertEqual(self.app.visible_articles[0].title, "日期筛选")
            self.app.help_query.set("not-present-unknown")
            self.assertEqual(self.app.visible_articles, ())
            self.assertEqual(str(self.app.article_button.cget("state")), "disabled")
            self.preview()

    def test_persistent_controls_reachable_at_two_sizes(self):
        self.root.deiconify()
        self.app.select_mode("monthly")
        for geometry in ("840x580", "1280x900"):
            self.root.geometry(geometry)
            self.root.update()
            for control in (*self.app.nav_buttons.values(), *self.app.mode_buttons.values(),
                            self.app.workbench.preview_button, self.app.workbench.export_button):
                x = control.winfo_rootx() - self.root.winfo_rootx()
                y = control.winfo_rooty() - self.root.winfo_rooty()
                self.assertTrue(control.winfo_ismapped(), (geometry, str(control)))
                self.assertGreaterEqual(x, 0)
                self.assertGreaterEqual(y, 0)
                self.assertLessEqual(x + control.winfo_width(), self.root.winfo_width())
                self.assertLessEqual(y + control.winfo_height(), self.root.winfo_height())
        self.app.select_mode("compare")
        self.root.update()
        self.app.select_mode("monthly")
        self.root.update()
        tree = self.app.workbench.summary.tree
        self.assertGreaterEqual(sum(tree.column(c, "width") for c in ("c0", "c2", "c5")), tree.winfo_width() - 30)
        self.app.show_page("history")
        for geometry in ("840x580", "1280x900"):
            self.root.geometry(geometry)
            self.root.update()
            for control in (self.app.history_filter_box, self.app.history_open_button,
                            self.app.history_tree.previous, self.app.history_tree.next):
                x = control.winfo_rootx() - self.root.winfo_rootx()
                y = control.winfo_rooty() - self.root.winfo_rooty()
                self.assertTrue(control.winfo_ismapped(), (geometry, str(control)))
                self.assertGreaterEqual(x, 0)
                self.assertGreaterEqual(y, 0)
                self.assertLessEqual(x + control.winfo_width(), self.root.winfo_width())
                self.assertLessEqual(y + control.winfo_height(), self.root.winfo_height())

    def test_privacy_scanner_never_echoes_matched_value(self):
        fake_value = "s" + "k-" + "x" * 32
        data = audit_module.scan_text(fake_value, "sample.py", "synthetic")
        self.assertEqual(data[0]["category"], "token-shaped-literal")
        self.assertNotIn(fake_value, json.dumps(data))
        self.assertEqual(data[0]["line"], 1)

    def test_static_pages_at_minimum_size_and_no_feedback_inputs(self):
        self.root.deiconify()
        self.root.geometry("840x580")
        for page in ("help", "feedback", "about", "rules"):
            self.app.show_page(page)
            self.root.update()
        self.app.show_page("feedback")
        self.root.update()
        pending = [self.app.pages["feedback"]]
        while pending:
            widget = pending.pop()
            self.assertNotIsInstance(widget, (tk.Entry, tk.Text, ttk.Entry))
            pending.extend(widget.winfo_children())
        self.assertTrue(self.app.contact_label.winfo_ismapped())
        self.assertLessEqual(self.app.contact_label.winfo_rootx() + self.app.contact_label.winfo_width(),
                             self.root.winfo_rootx() + self.root.winfo_width())
        self.assertTrue(audit_module.noreply_identity("bot@noreply.github.com"))
        self.assertTrue(audit_module.protected_path("nested/.env.production"))
        self.assertFalse(audit_module.protected_path(".env.example"))

    def test_close_cannot_destroy_running_legacy_task(self):
        self.app.legacy.task_running = True
        try:
            self.app.close()
            self.assertTrue(self.root.winfo_exists())
            self.assertFalse(self.app._closing)
        finally:
            self.app.legacy.task_running = False

    def test_new_ui_modules_have_no_network_imports_or_private_contacts(self):
        forbidden = {"requests", "socket", "urllib", "http", "webbrowser", "smtplib"}
        for name in ("desktop_ui.py", "offline_help.py"):
            source = (PROJECT / "e_help" / name).read_text(encoding="utf-8")
            tree = ast.parse(source)
            for node in ast.walk(tree):
                if isinstance(node, ast.Import):
                    self.assertFalse(any(alias.name.split(".")[0] in forbidden for alias in node.names))
                if isinstance(node, ast.ImportFrom):
                    self.assertNotIn((node.module or "").split(".")[0], forbidden)
        self.assertEqual(help_content.PUBLIC_CONTACTS, ())


if __name__ == "__main__":
    unittest.main()
