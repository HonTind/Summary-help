"""
v1.5.3 — 内存压力测试脚本

测试大量 Excel 文件汇总时的内存 / 稳定性 / 结果正确性。

覆盖：
    - 100 / 300 / 500 / 1000 个文件
    - .xlsx / .xls / .xlsm 三种格式
    - 多字段（客户 + 金额 + 电话 + 供应商，4 个字段）
    - 自定义字段配置（field_config 4 字段）
    - 多文件聚合（同名客户跨文件求和）
    - 不同文件大小（每文件 3~50 行不等）

内存记录：
    - 开始内存（进程 RSS）
    - 峰值内存（tracemalloc 追踪的 Python 堆峰值）
    - 结束内存（汇总 + gc 后进程 RSS）
    - 内存增量（结束 RSS - 开始 RSS）

结果正确性：
    - 汇总结果逐客户比对预期（金额 / 电话 / 供应商）
    - 输出 Excel 读回后与预期一致（客户数 / 金额 / 自定义字段 / 合计）

输出：
    - 控制台表格
    - stress_test_report.txt（与脚本同目录）

运行：
    cd "Day07(小试牛刀)源文件"
    python -m e_help.tests.stress_test_v153
"""

import os
import sys
import gc
import time
import random
import shutil
import tempfile
import tracemalloc
import datetime

THIS_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(THIS_DIR)
PARENT_DIR = os.path.dirname(PROJECT_ROOT)
if PARENT_DIR not in sys.path:
    sys.path.insert(0, PARENT_DIR)

from e_help.excel_reader import read_and_aggregate
from e_help.processor import sort_by_amount
from e_help.excel_writer import generate_summary
from e_help.field_mapper import save_field_config, reset_field_config

REPORT_PATH = os.path.join(THIS_DIR, "stress_test_report.txt")

# 多字段配置：客户 + 金额 + 电话 + 供应商
FIELD_CONFIG = [
    {"name": "客户", "keywords": ["客户"]},
    {"name": "金额", "keywords": ["金额"]},
    {"name": "电话", "keywords": ["电话"]},
    {"name": "供应商", "keywords": ["供应商"]},
]
FIELD_NAMES = ["客户", "金额", "电话", "供应商"]

N_CUSTOMERS = 60
CUSTOMERS = [f"客户{i:03d}" for i in range(1, N_CUSTOMERS + 1)]
PHONE = {c: f"138{i:08d}" for i, c in enumerate(CUSTOMERS, start=1)}
SUPPLIER = {c: f"供应商{(i % 12) + 1:02d}" for i, c in enumerate(CUSTOMERS)}


# ------------------------------------------------------------------
#  内存测量辅助
# ------------------------------------------------------------------

def _get_rss():
    """Windows 下通过 ctypes 读取当前进程工作集（RSS，字节）。失败返回 0。"""
    try:
        import ctypes
        from ctypes import wintypes

        class _PMC(ctypes.Structure):
            _fields_ = [
                ("cb", wintypes.DWORD),
                ("PageFaultCount", wintypes.DWORD),
                ("PeakWorkingSetSize", ctypes.c_size_t),
                ("WorkingSetSize", ctypes.c_size_t),
                ("QuotaPeakPagedPoolUsage", ctypes.c_size_t),
                ("QuotaPagedPoolUsage", ctypes.c_size_t),
                ("QuotaPeakNonPagedPoolUsage", ctypes.c_size_t),
                ("QuotaNonPagedPoolUsage", ctypes.c_size_t),
                ("PagefileUsage", ctypes.c_size_t),
                ("PeakPagefileUsage", ctypes.c_size_t),
            ]

        psapi = ctypes.windll.psapi
        kernel32 = ctypes.windll.kernel32
        # 必须显式声明返回/参数类型，否则伪句柄会被截断为 32 位导致调用失败返回 0
        kernel32.GetCurrentProcess.restype = wintypes.HANDLE
        psapi.GetProcessMemoryInfo.argtypes = [
            wintypes.HANDLE, ctypes.POINTER(_PMC), wintypes.DWORD,
        ]
        psapi.GetProcessMemoryInfo.restype = wintypes.BOOL
        pmc = _PMC()
        pmc.cb = ctypes.sizeof(pmc)
        psapi.GetProcessMemoryInfo(kernel32.GetCurrentProcess(), ctypes.byref(pmc), pmc.cb)
        return int(pmc.WorkingSetSize)
    except Exception:
        return 0


def _mb(n):
    return f"{n / (1024 * 1024):.2f} MB"


# ------------------------------------------------------------------
#  测试数据生成
# ------------------------------------------------------------------

def _rows_for_index(i):
    """按文件序号返回行数，制造不同大小的文件。"""
    if i % 100 == 0:
        return 50            # 每 100 个文件放一个较大的文件
    return (i % 15) + 3      # 3~17 行


def _write_xlsx(path, headers, rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(path)
    wb.close()


def _write_xls(path, headers, rows):
    import xlwt
    wb = xlwt.Workbook()
    sh = wb.add_sheet("S")
    for c, h in enumerate(headers):
        sh.write(0, c, h)
    for ri, r in enumerate(rows, start=1):
        for ci, v in enumerate(r):
            sh.write(ri, ci, v)
    wb.save(path)


def generate_files(folder, n_files, expected_amount):
    """生成 n_files 个 Excel 文件（.xlsx/.xls/.xlsm），并累积预期金额。

    Returns:
        total_size: 所有输入文件的总字节数
    """
    headers = FIELD_NAMES
    total_size = 0

    for i in range(n_files):
        if i % 5 == 0:
            ext = ".xls"
        elif i % 5 == 1:
            ext = ".xlsm"
        else:
            ext = ".xlsx"
        fpath = os.path.join(folder, f"f_{i:04d}{ext}")

        n_rows = _rows_for_index(i)
        rows = []
        for r in range(n_rows):
            cidx = (i * 3 + r * 7) % N_CUSTOMERS
            cust = CUSTOMERS[cidx]
            amt = random.randint(10, 5000)
            rows.append([cust, amt, PHONE[cust], SUPPLIER[cust]])
            expected_amount[cust] = expected_amount.get(cust, 0) + amt

        if ext == ".xls":
            _write_xls(fpath, headers, rows)
        else:  # .xlsx / .xlsm
            _write_xlsx(fpath, headers, rows)

        total_size += os.path.getsize(fpath)

    return total_size


# ------------------------------------------------------------------
#  结果正确性验证
# ------------------------------------------------------------------

def verify_output(output_path, expected_amount, n_files, file_count, field_names):
    """读回输出 Excel，与预期结果逐项比对。

    Returns:
        (ok: bool, detail: str)
    """
    from openpyxl import load_workbook

    if file_count != n_files:
        return False, f"文件数不符：file_count={file_count} != 期望 {n_files}"
    if list(field_names) != FIELD_NAMES:
        return False, f"字段顺序不符：{field_names}"
    if not os.path.isfile(output_path):
        return False, "输出文件未生成"

    try:
        wb = load_workbook(output_path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception as e:
        return False, f"输出文件读取失败：{e}"

    if not rows:
        return False, "输出文件为空"

    header = list(rows[0])
    if header != FIELD_NAMES:
        return False, f"输出表头不符：{header}"

    data_rows = rows[1:-1]   # 去掉表头行和合计行
    sum_row = rows[-1]

    if len(data_rows) != len(expected_amount):
        return False, f"客户数不符：输出 {len(data_rows)} != 期望 {len(expected_amount)}"

    out = {}
    for row in data_rows:
        cust, amt, phone, sup = row[0], row[1], row[2], row[3]
        out[cust] = (amt, phone, sup)

    for c, exp_amt in expected_amount.items():
        if c not in out:
            return False, f"客户 [{c}] 在输出中缺失"
        amt, phone, sup = out[c]
        if abs(float(amt) - float(exp_amt)) > 0.001:
            return False, f"客户 [{c}] 金额不符：{amt} != {exp_amt}"
        if str(phone) != PHONE[c]:
            return False, f"客户 [{c}] 电话不符：{phone} != {PHONE[c]}"
        if str(sup) != SUPPLIER[c]:
            return False, f"客户 [{c}] 供应商不符：{sup} != {SUPPLIER[c]}"

    exp_sum = sum(expected_amount.values())
    if abs(float(sum_row[1]) - float(exp_sum)) > 0.5:
        return False, f"合计行不符：{sum_row[1]} != {exp_sum}"

    return True, "OK"


# ------------------------------------------------------------------
#  单个批次：生成 → 汇总 → 写输出 → 验证 + 内存测量
# ------------------------------------------------------------------

def run_batch(n_files, stress_root):
    """运行一个批次，返回结果 dict。"""
    folder = os.path.join(stress_root, f"files_{n_files}")
    os.makedirs(folder, exist_ok=True)
    output_path = os.path.join(stress_root, f"out_{n_files}.xlsx")

    expected_amount = {}
    total_size = generate_files(folder, n_files, expected_amount)

    # 稳定后开始测量
    gc.collect()
    start_rss = _get_rss()
    tracemalloc.start()
    t0 = time.time()

    success = False
    error = ""
    try:
        total, file_count, field_names = read_and_aggregate(folder, stress_root)
        sorted_data = sort_by_amount(total, field_names)
        generate_summary(sorted_data, total, output_path, field_names)
        success = True
    except Exception as e:
        import traceback
        error = f"{type(e).__name__}: {e}"
        traceback.print_exc()

    t1 = time.time()
    peak_heap = tracemalloc.get_traced_memory()[1]
    tracemalloc.stop()

    gc.collect()
    end_rss = _get_rss()

    output_size = os.path.getsize(output_path) if os.path.isfile(output_path) else 0

    correct = False
    detail = ""
    if success:
        correct, detail = verify_output(
            output_path, expected_amount, n_files, file_count, field_names,
        )

    return {
        "n_files": n_files,
        "total_size": total_size,
        "start_rss": start_rss,
        "peak_heap": peak_heap,
        "end_rss": end_rss,
        "rss_delta": end_rss - start_rss,
        "elapsed": t1 - t0,
        "success": success,
        "correct": correct,
        "detail": detail,
        "error": error,
        "output_size": output_size,
        "customers": len(expected_amount),
    }


# ------------------------------------------------------------------
#  报告
# ------------------------------------------------------------------

def _fmt_mb(n):
    return f"{n / (1024 * 1024):.2f}"


def build_report(results):
    lines = []
    lines.append("=" * 70)
    lines.append("E-help v1.5.3 — 内存压力测试报告")
    lines.append("=" * 70)
    lines.append(f"生成时间：{datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append(f"字段配置：{' / '.join(FIELD_NAMES)}（4 个字段）")
    lines.append("覆盖格式：.xlsx / .xls / .xlsm；不同文件大小（3~50 行/文件）")
    lines.append("")

    lines.append("-" * 70)
    lines.append("汇总表")
    lines.append("-" * 70)
    header = (
        f"{'文件数量':>6} | {'输入大小':>10} | {'峰值内存':>10} | "
        f"{'结束内存':>10} | {'内存增量':>10} | {'耗时':>8} | {'结果':>6}"
    )
    lines.append(header)
    lines.append("-" * 70)

    for r in results:
        result_text = "PASS" if (r["success"] and r["correct"]) else "FAIL"
        lines.append(
            f"{r['n_files']:>6} | {_fmt_mb(r['total_size']):>10} | "
            f"{_fmt_mb(r['peak_heap']):>10} | {_fmt_mb(r['end_rss']):>10} | "
            f"{_fmt_mb(r['rss_delta']):>10} | {r['elapsed']:>6.2f}s | {result_text:>6}"
        )

    lines.append("-" * 70)
    lines.append("")

    lines.append("详细记录")
    lines.append("~" * 70)
    for r in results:
        lines.append(f"批次 {r['n_files']} 文件：")
        lines.append(f"  总输入大小 : {_mb(r['total_size'])}")
        lines.append(f"  开始内存   : {_mb(r['start_rss'])}（进程 RSS）")
        lines.append(f"  峰值内存   : {_mb(r['peak_heap'])}（Python 堆，tracemalloc）")
        lines.append(f"  结束内存   : {_mb(r['end_rss'])}（进程 RSS）")
        lines.append(f"  内存增量   : {_mb(r['rss_delta'])}（结束 RSS - 开始 RSS）")
        lines.append(f"  耗时       : {r['elapsed']:.2f} 秒")
        lines.append(f"  输出大小   : {_mb(r['output_size'])}")
        lines.append(f"  客户数     : {r['customers']}")
        lines.append(f"  是否成功   : {'是' if r['success'] else '否'}")
        lines.append(f"  结果正确   : {'是' if r['correct'] else '否'}（{r['detail']}）")
        if r["error"]:
            lines.append(f"  异常信息   : {r['error']}")
        lines.append("")
    lines.append("=" * 70)
    return "\n".join(lines)


def run_full_stress(stress_root=None):
    """运行 100/300/500/1000 完整压力测试，返回 results 列表。

    若未传入 stress_root，则自动创建临时目录并在结束后清理。
    字段配置在批处理前写入、结束后复位（不影响调用方环境）。
    """
    sizes = [100, 300, 500, 1000]

    own_root = stress_root is None
    if stress_root is None:
        stress_root = tempfile.mkdtemp(prefix="stress_v153_")

    random.seed(20260814)

    # 4 字段配置写入内存（read_and_aggregate 通过 _get_effective_config 读取）
    save_field_config(stress_root, FIELD_CONFIG)

    results = []
    for n in sizes:
        print(f"  正在运行批次：{n} 个文件 ...")
        r = run_batch(n, stress_root)
        results.append(r)
        status = "PASS" if (r["success"] and r["correct"]) else "FAIL"
        print(f"  完成：峰值堆 {_fmt_mb(r['peak_heap'])}，"
              f"耗时 {r['elapsed']:.2f}s，结果 {status}"
              + ("" if r["correct"] else f"（{r['detail']}）"))

    reset_field_config(stress_root)

    if own_root:
        shutil.rmtree(stress_root, ignore_errors=True)

    return results


def main():
    print("=" * 70)
    print("E-help v1.5.3 — 内存压力测试")
    print("=" * 70)

    stress_root = tempfile.mkdtemp(prefix="stress_v153_")
    print(f"压力测试数据目录：{stress_root}")
    print(f"字段配置：{' / '.join(FIELD_NAMES)}")
    print()

    results = run_full_stress(stress_root=stress_root)

    report = build_report(results)
    with open(REPORT_PATH, "w", encoding="utf-8") as f:
        f.write(report + "\n")

    print(report)
    print()
    print(f"报告已写入：{REPORT_PATH}")

    # 清理临时数据
    shutil.rmtree(stress_root, ignore_errors=True)

    # 任一失败则返回非零退出码
    all_pass = all(r["success"] and r["correct"] for r in results)
    return 0 if all_pass else 1


if __name__ == "__main__":
    sys.exit(main())
