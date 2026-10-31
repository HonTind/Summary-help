"""Hidden Tk tests with explicitly injected temporary data/configuration paths."""
import hashlib
import importlib
import json
from pathlib import Path
import sys
import tempfile
import threading
import time
import tkinter as tk
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
ui = importlib.import_module("e_help.workbench_gui")
monthly = importlib.import_module("e_help.monthly")
analysis = importlib.import_module("e_help.analysis")


class WorkbenchGuiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ehelper-workbench-ui-")
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.root = tk.Tk()
        self.root.withdraw()
        self.tk_errors = []
        self.root.report_callback_exception = lambda *args: self.tk_errors.append(args)
        self.window = ui.WorkbenchWindow(self.root, self.folder / "config" / "monthly_templates.json")
        self.window.withdraw()
        self.root.update()
        self.addCleanup(self.cleanup_window)
        for name, value in (("showinfo", None), ("showerror", None), ("showwarning", None), ("askyesno", True)):
            mock = patch.object(ui.messagebox, name, return_value=value)
            mock.start()
            self.addCleanup(mock.stop)

    def cleanup_window(self):
        self.window.runner.cancel()
        if self.window.runner.thread:
            self.window.runner.thread.join(2)
        try:
            self.root.destroy()
        except tk.TclError:
            pass
        self.assertFalse(self.window.runner.busy, "测试线程未释放")
        self.assertFalse(self.tk_errors, self.tk_errors)

    def make_file(self, name, rows, headers=("客户", "金额")):
        path = self.folder / name
        book = Workbook()
        sheet = book.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        book.save(path)
        book.close()
        return path

    def drain(self):
        deadline = time.monotonic() + 5
        while self.window._busy and not self.window._closed and time.monotonic() < deadline:
            self.root.update()
            time.sleep(0.005)
        self.root.update()
        self.assertFalse(self.window._busy if not self.window._closed else False, "工作台未结束")
        self.assertFalse(self.tk_errors)

    def monthly_preview(self, rows=(("A", "0.1"), ("A", "0.2"))):
        self.window.monthly_input.set_files([self.make_file("month.xlsx", rows)])
        self.window.month.set("2026-10")
        self.window.preview()
        self.drain()
        return self.window.result

    def test_monthly_preview_and_full_export(self):
        result = self.monthly_preview()
        self.assertEqual(str(result.total), "0.3")
        self.assertEqual(self.window.summary.rows[0][2], "0.3")
        self.assertEqual(len(self.window.sources.rows), 2)
        with patch.object(ui.filedialog, "askdirectory", return_value=str(self.folder)):
            self.window.export()
            self.drain()
        outputs = list(self.folder.glob("月度汇总-2026-10*.xlsx"))
        self.assertEqual(len(outputs), 1)
        self.assertIn("已导出", self.window.status.get())
        book = load_workbook(outputs[0])
        self.addCleanup(book.close)
        self.assertEqual(book["来源明细"].max_row, 3)

    def test_both_modes_and_independent_fields(self):
        self.window.tabs.select(1)
        self.root.update()
        left = self.make_file("left.xlsx", [("D", 1), ("D", 2)], ("编号", "货款"))
        right = self.make_file("right.xlsx", [(3, "D")], ("支付", "单号"))
        self.window.left_input.set_files([left])
        self.window.right_input.set_files([right])
        self.window.left_input.key.set("编号")
        self.window.left_input.amount.set("货款")
        self.window.right_input.key.set("单号")
        self.window.right_input.amount.set("支付")
        self.window.preview()
        self.drain()
        self.assertTrue(self.window.result.all_consistent)
        self.window.mode.set("单据逐笔")
        self.window.mode_changed()
        self.window.preview()
        self.drain()
        self.assertEqual(self.window.result.rows[0].category, "ambiguous")
        self.assertIn("结果不完整", self.window.status.get())
        self.assertNotIn("全部一致", self.window.status.get())
        self.assertEqual(len(self.window.sources.rows), 3)

    def test_incomplete_export_requires_confirmation(self):
        self.monthly_preview([("A", 1), ("B", "bad")])
        with patch.object(ui.messagebox, "askyesno", return_value=False) as confirm, patch.object(
            ui.filedialog, "askdirectory") as choose:
            self.window.export()
        confirm.assert_called_once()
        choose.assert_not_called()
        self.assertEqual(len(self.window.issues.rows), 1)

    def test_rule_change_invalidates_old_result(self):
        self.monthly_preview()
        self.window.month.set("2026-11")
        self.assertIsNone(self.window.result)
        self.assertEqual(str(self.window.export_button.cget("state")), "disabled")
        with patch.object(ui.filedialog, "askdirectory") as choose:
            self.window.export()
        choose.assert_not_called()

    def test_template_save_load_and_declined_update(self):
        self.window.template_name.set("隔离模板")
        self.window.save_template()
        path = self.window.template_path
        self.assertTrue(path.exists())
        before = path.read_bytes()
        self.window.monthly_input.key.set("ID")
        with patch.object(ui.messagebox, "askyesno", return_value=False):
            self.window.save_template()
        self.assertEqual(path.read_bytes(), before)
        self.window.load_template()
        self.assertEqual(self.window.monthly_input.key.get(), "客户")
        self.assertFalse(monthly.load_templates(path)["隔离模板"].get("input_files"))
        self.window.refresh_templates()
        self.assertIn("隔离模板", self.window.template_box.cget("values"))

    def test_broken_template_file_is_not_reset(self):
        path = self.window.template_path
        path.parent.mkdir()
        path.write_text("{broken", encoding="utf-8")
        self.window.save_template()
        self.assertEqual(path.read_text(encoding="utf-8"), "{broken")

    def test_template_name_whitespace_is_consistent(self):
        self.window.template_name.set("  月报  ")
        self.window.save_template()
        self.window.monthly_input.key.set("ID")
        self.window.load_template()
        self.assertEqual(self.window.monthly_input.key.get(), "客户")

    def test_worker_single_writer_cancel_and_ui_restored(self):
        def slow(cancel, progress):
            progress("等待取消")
            cancel.wait(2)
            analysis.check_cancel(cancel)
            return None
        self.assertTrue(self.window._start(slow, "preview"))
        self.assertFalse(self.window._start(slow, "preview"))
        self.assertEqual(str(self.window.monthly_input.key_entry.cget("state")), "disabled")
        self.window.cancel()
        self.drain()
        self.assertIsNone(self.window.result)
        self.assertIn("取消", self.window.status.get())
        self.assertEqual(str(self.window.preview_button.cget("state")), "normal")

    def test_late_cancel_discards_completed_preview(self):
        file = self.make_file("late.xlsx", [("A", 1)])
        result = monthly.run_monthly([file], monthly.make_template("T", analysis.SheetRules()), "2026-10")
        self.window._start(lambda cancel, progress: result, "preview")
        self.window.runner.thread.join(2)
        self.window.cancel()
        self.drain()
        self.assertIsNone(self.window.result)
        self.assertIn("取消", self.window.status.get())

    def test_worker_failure_restores_controls(self):
        def fail(cancel, progress):
            raise RuntimeError("模拟失败")
        self.window._start(fail, "preview")
        self.drain()
        self.assertIn("模拟失败", self.window.status.get())
        self.assertEqual(str(self.window.preview_button.cget("state")), "normal")

    def test_close_waits_for_cancelled_worker(self):
        self.window._start(lambda cancel, progress: cancel.wait(2) or analysis.check_cancel(cancel), "preview")
        self.window.close()
        self.assertTrue(self.window._closing)
        self.drain()
        self.assertTrue(self.window._closed)

    def test_parent_destroy_signals_cancel_without_tk_in_worker(self):
        self.window._start(lambda cancel, progress: cancel.wait(2) or analysis.check_cancel(cancel), "preview")
        self.root.destroy()
        self.window.runner.thread.join(2)
        self.assertTrue(self.window.runner.cancel_event.is_set())
        self.assertFalse(self.window.runner.busy)
        self.assertFalse(self.tk_errors)

    def test_close_confirmation_completion_race(self):
        result = monthly.run_monthly([self.make_file("close-race.xlsx", [("A", 1)])],
                                     monthly.make_template("T", analysis.SheetRules()), "2026-10")
        self.window._start(lambda cancel, progress: result, "preview")
        def answer(*args, **kwargs):
            self.drain()  # Simulate the native dialog's nested Tk event loop.
            return True
        with patch.object(ui.messagebox, "askyesno", side_effect=answer):
            self.window.close()
        self.assertTrue(self.window._closed)

    def test_template_list_available_after_reopening(self):
        monthly.save_template(self.window.template_path, monthly.make_template("我的月报", analysis.SheetRules()))
        another = ui.WorkbenchWindow(self.root, self.window.template_path)
        another.withdraw()
        self.assertIn("我的月报", another.template_box.cget("values"))
        another.destroy()

    def test_worker_uses_snapshot_and_not_changed_tk_variables(self):
        path = self.make_file("snapshot.xlsx", [("A", 1)])
        self.window.monthly_input.set_files([path])
        release = threading.Event()
        original = ui.run_monthly
        def delayed(paths, template, month, **kwargs):
            release.wait(2)
            return original(paths, template, month, **kwargs)
        with patch.object(ui, "run_monthly", side_effect=delayed):
            self.window.preview()
            self.window.monthly_input.key.set("不应影响运行")
            release.set()
            self.drain()
        self.assertEqual(self.window.result.template["rules"]["key_column"], "客户")
        self.assertEqual(len(self.window.result.rows), 1)

    def test_pagination_does_not_truncate_export(self):
        self.monthly_preview([(f"K{i:03}", i) for i in range(205)])
        self.assertEqual(len(self.window.summary.rows), 205)
        self.assertEqual(len(self.window.summary.tree.get_children()), 100)
        self.window.summary.move(2)
        self.assertEqual(len(self.window.summary.tree.get_children()), 5)
        with patch.object(ui.filedialog, "askdirectory", return_value=str(self.folder)):
            self.window.export()
            self.drain()
        output = next(self.folder.glob("月度汇总-2026-10*.xlsx"))
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book["来源明细"].max_row, 206)

    def test_column_index_does_not_confuse_numeric_header(self):
        self.assertEqual(ui.column_selection("#2"), 2)
        self.assertEqual(ui.column_selection("2"), "2")
        with self.assertRaises(ValueError):
            ui.column_selection("#abc")

    def test_existing_main_entry_is_guarded_and_single_instance(self):
        legacy = importlib.import_module("e_help.gui")
        app = legacy.SummaryApp.__new__(legacy.SummaryApp)
        app.root = self.root
        app.task_running = True
        with patch.object(ui, "WorkbenchWindow") as create:
            app._on_open_workbench()
            create.assert_not_called()
            app.task_running = False
            with patch.object(legacy, "get_config_dir", return_value=str(self.folder)):
                app._on_open_workbench()
                app._on_open_workbench()
            create.assert_called_once()
            app._workbench_window.grab_set.assert_called_once()
            app._workbench_window.lift.assert_called_once()

    def test_window_layout_at_default_size(self):
        self.window.geometry("1040x720")
        self.window.update_idletasks()
        self.assertEqual(self.window.summary.PAGE_SIZE, 100)
        self.assertLess(self.window.preview_button.winfo_reqheight(), 60)
        self.assertTrue(self.window.tabs.tab(0, "text").startswith("每月"))

    def test_composite_preview_export_snapshot_and_invalidation(self):
        path = self.make_file("composite.xlsx", [("甲", "D1", 1), ("甲", "D2", 2), ("乙", "D1", 3)],
                              headers=("客户", "单据号", "金额"))
        before = hashlib.sha256(path.read_bytes()).digest()
        self.window.monthly_input.set_files([path])
        self.window.monthly_input.key.set("客户")
        self.window.monthly_input.extra_keys.set('["单据号"]')
        self.window.month.set("2026-10")
        self.window.preview()
        self.drain()
        result = self.window.result
        self.assertEqual(result.template["rules"]["key_fields"], ("客户", "单据号"))
        self.assertEqual({row.parts for row in result.rows},
                         {("甲", "D1"), ("甲", "D2"), ("乙", "D1")})
        self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), before)
        self.window.monthly_input.extra_keys.set('["单据号", "批次"]')
        self.assertIsNone(self.window.result)
        self.window.monthly_input.extra_keys.set('["单据号"]')
        self.window.preview()
        self.drain()
        with patch.object(ui.filedialog, "askdirectory", return_value=str(self.folder)):
            self.window.export()
            self.drain()
        output = next(self.folder.glob("月度汇总-2026-10*.xlsx"))
        book = load_workbook(output)
        self.addCleanup(book.close)
        config = dict(book["运行配置"].values)
        self.assertEqual(config["key_fields[1]"], "客户")
        self.assertEqual(config["key_fields[2]"], "单据号")
        self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), before)

    def test_composite_run_disabled_mode_switch_and_bad_json(self):
        self.window.tabs.select(1)
        self.root.update()
        self.window.left_input.key.set("客户")
        self.window.left_input.extra_keys.set('["单据号"]')
        self.window.right_input.key.set("客户")
        self.window.right_input.extra_keys.set('["单据号"]')
        self.window.mode.set("单据逐笔")
        self.window.mode_changed()  # Must not overwrite a user-configured composite.
        self.assertEqual(self.window.left_input.key.get(), "客户")
        self.assertEqual(self.window.left_input.extra_keys.get(), '["单据号"]')
        left = self.make_file("pl.xlsx", [("甲", "D1", 1)], headers=("客户", "单据号", "金额"))
        right = self.make_file("pr.xlsx", [("甲", "D1", 1)], headers=("客户", "单据号", "金额"))
        self.window.left_input.set_files([left])
        self.window.right_input.set_files([right])
        self.window.left_input.extra_keys.set("[not json")
        with patch.object(ui.messagebox, "showerror") as error:
            self.window.preview()
        error.assert_called_once()
        self.window.left_input.extra_keys.set('["单据号"]')
        release = threading.Event()
        original = ui.run_reconciliation

        def slow(*args, **kwargs):
            release.wait(2)
            return original(*args, **kwargs)

        with patch.object(ui, "run_reconciliation", side_effect=slow):
            self.window.preview()
            self.assertEqual(str(self.window.left_input.extra_entry.cget("state")), "disabled")
            release.set()
            self.drain()
        self.assertIsNotNone(self.window.result)
        self.assertTrue(self.window.result.all_consistent)

    # --- composite-key UI regressions (round 2) -----------------------------
    def assert_controls_reachable(self, geometry):
        """Real Tk layout evidence: composite inputs and footer buttons are on-screen."""
        self.root.deiconify()  # A transient Toplevel only maps when its master is mapped.
        self.window.deiconify()
        self.window.geometry(geometry)
        self.window.update_idletasks()
        self.window.update()
        try:
            width, height = self.window.winfo_width(), self.window.winfo_height()
            widgets = {
                "键字段①": self.window.monthly_input.key_entry,
                "组合键②③…": self.window.monthly_input.extra_entry,
                "预览结果": self.window.preview_button,
                "导出完整 Excel": self.window.export_button,
            }
            bounds = {}
            for name, widget in widgets.items():
                x = widget.winfo_rootx() - self.window.winfo_rootx()
                y = widget.winfo_rooty() - self.window.winfo_rooty()
                w, h = widget.winfo_width(), widget.winfo_height()
                bounds[name] = (x, y, w, h, bool(widget.winfo_ismapped()))
                self.assertTrue(widget.winfo_ismapped(), f"{geometry} {name} 未映射: {bounds}")
                self.assertGreaterEqual(x, 0, f"{geometry} {name} 越出左边界: {bounds}")
                self.assertGreaterEqual(y, 0, f"{geometry} {name} 越出上边界: {bounds}")
                self.assertLessEqual(x + w, width, f"{geometry} {name} 越出右边界: {bounds}")
                self.assertLessEqual(y + h, height, f"{geometry} {name} 越出下边界: {bounds}")
        finally:
            self.window.withdraw()
            self.root.withdraw()
            self.window.update_idletasks()

    def test_composite_controls_reachable_at_min_and_default_size(self):
        self.window.tabs.select(0)
        self.root.update()
        for geometry in ("840x580", "1040x720"):
            with self.subTest(geometry=geometry):
                self.assert_controls_reachable(geometry)

    def test_composite_worker_uses_frozen_snapshot(self):
        path = self.make_file("csnap.xlsx", [("甲", "D1", 1)], headers=("客户", "单据号", "金额"))
        self.window.monthly_input.set_files([path])
        self.window.monthly_input.key.set("客户")
        self.window.monthly_input.extra_keys.set('["单据号"]')
        self.window.month.set("2026-10")
        release = threading.Event()
        original = ui.run_monthly

        def delayed(paths, template, month, **kwargs):
            release.wait(2)
            return original(paths, template, month, **kwargs)

        with patch.object(ui, "run_monthly", side_effect=delayed):
            self.window.preview()
            # Mutate both composite Tk variables while the worker is still running.
            self.window.monthly_input.key.set("金额")
            self.window.monthly_input.extra_keys.set('["批次"]')
            release.set()
            self.drain()
        self.assertEqual(self.window.result.template["rules"]["key_fields"], ("客户", "单据号"))
        self.assertEqual({row.parts for row in self.window.result.rows}, {("甲", "D1")})

    def test_composite_field_count_mismatch_rejected_before_worker(self):
        self.window.tabs.select(1)
        self.root.update()
        left = self.make_file("mm-l.xlsx", [("甲", "D1", 1)], headers=("客户", "单据号", "金额"))
        right = self.make_file("mm-r.xlsx", [("甲", 1)], headers=("客户", "金额"))
        self.window.left_input.set_files([left])
        self.window.right_input.set_files([right])
        self.window.left_input.key.set("客户")
        self.window.left_input.extra_keys.set('["单据号"]')
        self.window.right_input.key.set("客户")
        self.window.right_input.extra_keys.set("")
        with patch.object(self.window.runner, "start") as start, \
                patch.object(ui, "run_reconciliation") as run, \
                patch.object(ui.messagebox, "showerror") as error:
            self.window.preview()
        start.assert_not_called()   # Worker never starts.
        run.assert_not_called()     # No file is read.
        error.assert_called_once()
        self.assertFalse(self.window._busy)

    def test_composite_template_reopen_preserves_order_and_special_chars(self):
        rules = analysis.SheetRules("客户,名", "金额", None, 1, ("客户,名", '单据"号', "批次\n号"))
        monthly.save_template(self.window.template_path, monthly.make_template("组合模板", rules))
        another = ui.WorkbenchWindow(self.root, self.window.template_path)
        another.withdraw()
        try:
            another.template_name.set("组合模板")
            another.load_template()
            self.assertEqual(another.monthly_input.key.get(), "客户,名")
            self.assertEqual(json.loads(another.monthly_input.extra_keys.get()),
                             ['单据"号', "批次\n号"])
            self.assertEqual(another.monthly_input.rules().field_selections(),
                             ("客户,名", '单据"号', "批次\n号"))
        finally:
            another.destroy()


if __name__ == "__main__":
    unittest.main()
