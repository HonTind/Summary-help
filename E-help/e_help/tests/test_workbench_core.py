"""Synthetic, isolated workbench acceptance tests. Never read real user data."""
import hashlib
import importlib
import json
from pathlib import Path
import sys
import tempfile
import threading
import unittest
import datetime
import zipfile
import xml.etree.ElementTree as ET
from unittest.mock import patch
from decimal import Decimal

from openpyxl import Workbook, load_workbook
import xlwt

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
PACKAGE = "e_help"
analysis = importlib.import_module(PACKAGE + ".analysis")
monthly = importlib.import_module(PACKAGE + ".monthly")
reconciliation = importlib.import_module(PACKAGE + ".reconciliation")
export = importlib.import_module(PACKAGE + ".workbench_export")
Rules = analysis.SheetRules


class WorkbenchCoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ehelper-workbench-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.rules = Rules()
        self.template = monthly.make_template("默认月报", self.rules)

    def workbook(self, name, rows, headers=("客户", "金额"), sheet="账单", title_rows=()):
        path = self.root / name
        if path.suffix == ".xls":
            book = xlwt.Workbook()
            ws = book.add_sheet(sheet)
            all_rows = list(title_rows) + [headers] + list(rows)
            for i, row in enumerate(all_rows):
                for j, value in enumerate(row):
                    if value is not None:
                        ws.write(i, j, value)
            book.save(str(path))
        else:
            book = Workbook()
            ws = book.active
            ws.title = sheet
            for row in title_rows:
                ws.append(row)
            ws.append(headers)
            for row in rows:
                ws.append(row)
            book.save(path)
            book.close()
        return path

    def reports(self, left, right):
        return (analysis.read_inputs([self.workbook("left.xlsx", left)], self.rules),
                analysis.read_inputs([self.workbook("right.xlsx", right)], self.rules))

    def test_three_formats_sources_and_read_only(self):
        for extension in ("xlsx", "xlsm", "xls"):
            with self.subTest(extension=extension):
                path = self.workbook("book." + extension, [(" 001 ", "0.1"), (None, None),
                                                          ("1", "-0.2"), ("zero", 0)],
                                     sheet="明细", title_rows=[("标题",)])
                before = hashlib.sha256(path.read_bytes()).digest()
                report = analysis.read_inputs([path], Rules(sheet="明细", header_row=2))
                self.assertFalse(report.issues)
                self.assertEqual([r.key for r in report.records], ["001", "1", "zero"])
                self.assertEqual([r.source.row for r in report.records], [3, 5, 6])
                self.assertTrue(all(r.source.sheet == "明细" and r.source.file == str(path) for r in report.records))
                self.assertEqual(report.scanned_rows, 4)
                self.assertEqual(report.blank_rows, 1)
                self.assertEqual(report.invalid_rows, 0)
                self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), before)

    def test_invalid_values_and_row_accounting(self):
        path = self.workbook("bad.xlsx", [("A", True), ("B", "=SUM(1,2)"), ("C", None),
                                            ("D", "NaN"), ("E", "Infinity"), ("F", "1,000"),
                                            ("G", "$1"), ("", 1), ("valid", -1), (None, None)])
        report = analysis.read_inputs([path], self.rules)
        self.assertEqual(report.invalid_rows, 8)
        self.assertEqual(len(report.records), 1)
        self.assertEqual(report.scanned_rows, len(report.records) + report.blank_rows + report.invalid_rows)
        self.assertFalse(report.complete)
        self.assertIn("公式没有缓存", report.issues[1].message)

    def test_xls_boolean_rejected(self):
        report = analysis.read_inputs([self.workbook("bool.xls", [("A", True), ("B", 0)])], self.rules)
        self.assertEqual(report.invalid_rows, 1)
        self.assertEqual(len(report.records), 1)

    def rewrite_xml(self, original, name, transform):
        target = self.root / name
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(target, "w") as output:
            for entry in source.infolist():
                content = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    tree = ET.fromstring(content)
                    transform(tree)
                    content = ET.tostring(tree, encoding="utf-8")
                output.writestr(entry, content)
        return target

    def test_absent_or_underreported_dimensions_do_not_drop_rows(self):
        namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        original = self.workbook("dimension.xlsx", [("A", 1), ("B", 2), ("C", 3)])
        for missing in (True, False):
            def transform(tree):
                dimension = tree.find(namespace + "dimension")
                if missing:
                    tree.remove(dimension)
                else:
                    dimension.set("ref", "A1:A2")
            path = self.rewrite_xml(original, f"dimension-{missing}.xlsx", transform)
            report = analysis.read_inputs([path], self.rules)
            self.assertTrue(report.complete, report.issues)
            self.assertEqual(len(report.records), 3)
            self.assertEqual(report.records[-1].source.row, 4)

    def test_cached_formula_error_key_is_invalid(self):
        namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        original = self.workbook("formula-error.xlsx", [("=NA()", 2)])
        def transform(tree):
            cell = tree.find(".//" + namespace + "c[@r='A2']")
            cell.set("t", "e")
            cell.find(namespace + "v").text = "#N/A"
        path = self.rewrite_xml(original, "cached-error.xlsx", transform)
        report = analysis.read_inputs([path], self.rules)
        self.assertEqual(report.invalid_rows, 1)
        self.assertFalse(report.records)
        self.assertIn("Excel 错误", report.issues[0].message)

    def test_dates_rejected_as_required_fields_across_formats(self):
        path = self.workbook("dates.xlsx", [(datetime.datetime(2026, 10, 8), 1)])
        report = analysis.read_inputs([path], self.rules)
        self.assertEqual(report.invalid_rows, 1)
        path = self.root / "dates.xls"
        book = xlwt.Workbook()
        sheet = book.add_sheet("账单")
        sheet.write(0, 0, "客户")
        sheet.write(0, 1, "金额")
        sheet.write(1, 0, datetime.datetime(2026, 10, 8), xlwt.easyxf(num_format_str="YYYY-MM-DD"))
        sheet.write(1, 1, 1)
        book.save(str(path))
        self.assertEqual(analysis.read_inputs([path], self.rules).invalid_rows, 1)

    def test_formula_key_without_cache_is_issue(self):
        report = analysis.read_inputs([self.workbook("formula.xlsx", [("=1+1", 2)])], self.rules)
        self.assertEqual(report.invalid_rows, 1)
        self.assertIn("公式", report.issues[0].message)

    def test_structured_file_errors(self):
        path = self.workbook("missing-header.xlsx", [("A", 1)], headers=("其他", "金额"))
        report = analysis.read_inputs([path, self.root / "not-here.xlsx"], self.rules)
        self.assertEqual(len(report.issues), 2)
        self.assertTrue(all(issue.code == "file_error" for issue in report.issues))
        self.assertFalse(report.complete)

    def test_sheet_header_and_column_validation(self):
        path = self.workbook("headers.xlsx", [("A", 1)], headers=("ID", "值"), title_rows=[("说明",)])
        good = analysis.read_inputs([path], Rules(1, 2, header_row=2))
        self.assertTrue(good.complete)
        for rules in (Rules(sheet="不存在"), Rules(header_row=9), Rules(1, 1), Rules(1, 8)):
            self.assertEqual(analysis.read_inputs([path], rules).issues[0].code, "file_error")
        for kwargs in ({"header_row": True}, {"header_row": 0}, {"key_column": 0}, {"amount_column": ""}):
            with self.assertRaises(ValueError):
                Rules(**kwargs)

    def test_duplicate_header_requires_index(self):
        path = self.workbook("duplicates.xlsx", [("A", 1)], headers=("值", "值"))
        self.assertTrue(analysis.read_inputs([path], Rules("值", 2)).issues)
        self.assertTrue(analysis.read_inputs([path], Rules(1, 2)).complete)

    def test_duplicate_files_not_double_counted(self):
        path = self.workbook("duplicate.xlsx", [("A", 1)])
        report = analysis.read_inputs([path, path], self.rules)
        self.assertEqual(len(report.records), 1)
        self.assertEqual(report.issues[0].code, "duplicate_file")

    def test_empty_or_no_inputs_incomplete(self):
        self.assertFalse(analysis.read_inputs([], self.rules).complete)
        report = analysis.read_inputs([self.workbook("empty.xlsx", [])], self.rules)
        self.assertFalse(report.complete)
        self.assertEqual(report.issues[0].code, "no_valid_rows")

    def test_amounts_plain_finite_and_bounded(self):
        for value in (None, True, "", "nan", float("inf"), float("nan"), "1_000", "￥1", "1e1001"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                analysis.parse_amount(value)
        self.assertEqual(analysis.parse_amount("-0.001"), Decimal("-0.001"))
        self.assertEqual(analysis.parse_amount("1e3"), Decimal("1000"))

    def test_exact_decimal_monthly_and_more_than_28_digits(self):
        path = self.workbook("decimal.xlsx", [("A", "0.1"), ("A", "0.2"),
                                             ("B", "123456789012345678901234567890.001"),
                                             ("B", "0.009")])
        result = monthly.run_monthly([path], self.template, "2026-10")
        self.assertEqual(result.rows[0].amount, Decimal("0.3"))
        self.assertEqual(str(result.rows[1].amount), "123456789012345678901234567890.010")

    def test_month_validation(self):
        for value in ("2026-00", "2026-13", "0000-01", "../2026-10", "2026-1", None):
            with self.assertRaises(ValueError):
                monthly.validate_month(value)

    def test_month_template_reuses_rules_not_filenames(self):
        first = self.workbook("October.xlsx", [("A", 1)])
        second = self.workbook("November.xlsx", [("A", 2)])
        self.assertEqual(monthly.run_monthly([first], self.template, "2026-10").total, Decimal("1"))
        self.assertEqual(monthly.run_monthly([second], self.template, "2026-11").total, Decimal("2"))
        result = monthly.run_monthly([second], self.template, "2026-11")
        self.template["rules"]["key_column"] = "变更"
        self.assertEqual(result.template["rules"]["key_column"], "客户")
        snapshot = result.config
        snapshot["template"]["rules"]["key_column"] = "变更"
        self.assertEqual(result.template["rules"]["key_column"], "客户")

    def test_customer_all_categories_and_exact_keys(self):
        left, right = self.reports([("A", "0.1"), ("A", "0.2"), ("B", "1.001"), ("L", -1), ("001", 0)],
                                   [("A", "0.3"), ("B", "1.002"), ("R", 3), ("1", 0)])
        result = reconciliation.compare_reports(left, right)
        rows = {r.key: r for r in result.rows}
        self.assertEqual(rows["A"].category, "consistent")
        self.assertEqual(rows["B"].difference, Decimal("-0.001"))
        self.assertEqual(rows["B"].category, "different")
        self.assertEqual(rows["L"].category, "left_only")
        self.assertEqual(rows["R"].category, "right_only")
        self.assertEqual(rows["001"].category, "left_only")
        self.assertEqual(rows["1"].category, "right_only")
        self.assertIsNone(rows["L"].right_amount)
        self.assertIsNone(rows["L"].difference)

    def test_tolerance_boundary_no_quantization(self):
        left, right = self.reports([("A", "1.001")], [("A", "1.002")])
        self.assertFalse(reconciliation.compare_reports(left, right, tolerance="0").all_consistent)
        self.assertTrue(reconciliation.compare_reports(left, right, tolerance="0.001").all_consistent)
        self.assertFalse(reconciliation.compare_reports(left, right, tolerance="0.000999").all_consistent)
        for value in ("-1", "NaN", "Infinity", True):
            with self.assertRaises(ValueError):
                reconciliation.compare_reports(left, right, tolerance=value)

    def test_high_precision_difference_no_context_rounding(self):
        left, right = self.reports([("A", "123456789012345678901234567890.001")],
                                   [("A", "123456789012345678901234567890.002")])
        result = reconciliation.compare_reports(left, right, tolerance="0")
        self.assertEqual(result.rows[0].difference, Decimal("-0.001"))

    def test_document_duplicates_both_sides_all_rows(self):
        left, right = self.reports([("D", 1), ("D", 2), ("U", 3)], [("D", 3), ("U", 3)])
        for lreport, rreport in ((left, right), (right, left)):
            result = reconciliation.compare_reports(lreport, rreport, mode="document")
            duplicate = next(r for r in result.rows if r.key == "D")
            self.assertEqual(duplicate.category, "ambiguous")
            self.assertEqual(len(duplicate.left) + len(duplicate.right), 3)
            self.assertIsNone(duplicate.difference)
            self.assertFalse(result.all_consistent)
            classified = [(side, record.source) for row in result.rows
                          for side, records in (("left", row.left), ("right", row.right)) for record in records]
            self.assertEqual(len(classified), len(set(classified)))
            self.assertEqual(len(classified), len(left.records) + len(right.records))

    def test_read_error_prevents_all_consistent(self):
        left, right = self.reports([("A", 1), ("B", "坏金额")], [("A", 1)])
        result = reconciliation.compare_reports(left, right)
        self.assertEqual(result.rows[0].category, "consistent")
        self.assertFalse(result.complete)
        self.assertFalse(result.all_consistent)

    def test_independent_left_right_fields_and_old_globals(self):
        mapper = importlib.import_module(PACKAGE + ".field_mapper")
        old_config = json.dumps(mapper._FIELD_CONFIG, ensure_ascii=False)
        left = self.workbook("left.xlsx", [("A", 2)], headers=("买方", "货款"))
        right = self.workbook("right.xlsx", [(2, "A")], headers=("支付", "对象"))
        result = reconciliation.run_reconciliation([left], [right], Rules("买方", "货款"),
                                                   Rules("对象", "支付"))
        self.assertTrue(result.all_consistent)
        self.assertEqual(json.dumps(mapper._FIELD_CONFIG, ensure_ascii=False), old_config)

    def test_templates_roundtrip_update_and_isolation(self):
        path = self.root / "config" / "monthly_templates.json"
        monthly.save_template(path, self.template)
        before = path.read_bytes()
        with self.assertRaises(FileExistsError):
            monthly.save_template(path, self.template)
        self.assertEqual(path.read_bytes(), before)
        loaded = monthly.load_templates(path)
        loaded["默认月报"]["rules"]["key_column"] = "ID"
        self.assertEqual(monthly.load_templates(path)["默认月报"]["rules"]["key_column"], "客户")
        new = monthly.make_template("默认月报", Rules("ID", "金额"))
        monthly.save_template(path, new, allow_update=True)
        self.assertEqual(monthly.load_templates(path)["默认月报"], new)

    def test_bad_json_and_version_preserved(self):
        path = self.root / "bad.json"
        for text in ("{broken", '{"schema_version":2,"templates":[]}'):
            path.write_text(text, encoding="utf-8")
            before = path.read_bytes()
            with self.assertRaises((ValueError, json.JSONDecodeError)):
                monthly.save_template(path, self.template, allow_update=True)
            self.assertEqual(path.read_bytes(), before)

    def test_failed_atomic_save_preserves_old_file(self):
        path = self.root / "template.json"
        monthly.save_template(path, self.template)
        before = path.read_bytes()
        with patch.object(monthly.os, "replace", side_effect=PermissionError("模拟失败")):
            with self.assertRaises(PermissionError):
                monthly.save_template(path, self.template, allow_update=True)
        self.assertEqual(path.read_bytes(), before)
        self.assertFalse(list(self.root.glob(".monthly-*.tmp")))

    def test_template_rejects_extra_data_and_names(self):
        for value in ("../escape", "C:/escape", "CON", "bad.", "", "A/B"):
            with self.assertRaises(ValueError):
                monthly.make_template(value, self.rules)
        template = dict(self.template, input_files=["private.xlsx"])
        with self.assertRaises(ValueError):
            monthly.validate_template(template)

    def test_export_traceability_precision_text_and_no_overwrite(self):
        path = self.workbook("source.xlsx", [("placeholder", "1.001"), ("bad", "NaN")], sheet="=来源")
        # Literal malicious-looking key rather than an input formula.
        book = load_workbook(path)
        book.active["A2"] = "=danger"
        book.active["A2"].data_type = "s"
        book.save(path)
        book.close()
        before = path.read_bytes()
        result = monthly.run_monthly([path], self.template, "2026-10")
        output = export.export_result(result, self.root, prefix="source")
        again = export.export_result(result, self.root, prefix="source")
        self.assertNotEqual(output, again)
        self.assertNotEqual(output, str(path))
        self.assertEqual(path.read_bytes(), before)
        with_read = load_workbook(output, data_only=False)
        self.addCleanup(with_read.close)
        self.assertEqual(with_read.sheetnames, ["月度汇总", "来源明细", "异常清单", "运行配置"])
        self.assertEqual(with_read["月度汇总"]["A2"].value, "=danger")
        self.assertEqual(with_read["月度汇总"]["A2"].data_type, "s")
        self.assertEqual(with_read["月度汇总"]["B2"].value, "1.001")
        self.assertEqual(with_read["来源明细"]["D2"].data_type, "s")
        self.assertEqual(with_read["来源明细"]["E2"].value, 2)
        self.assertEqual(with_read["异常清单"]["D2"].value, 3)
        for sheet in with_read:
            self.assertFalse(any(cell.data_type == "f" for row in sheet for cell in row))

    def test_export_reconciliation_ambiguity_and_configuration(self):
        left, right = self.reports([("D", 1), ("D", 2)], [("D", 3)])
        result = reconciliation.compare_reports(left, right, mode="document", tolerance="0.01")
        output = export.export_result(result, self.root)
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book["对账结果"]["B2"].value, "重复歧义")
        self.assertEqual(book["来源明细"].max_row, 4)
        config = dict(book["运行配置"].values)
        self.assertEqual(config["mode"], "document")
        self.assertEqual(config["tolerance"], "0.01")
        self.assertEqual(config["全部一致"], "否")

    def test_export_paths_and_limits_fail_without_output(self):
        result = monthly.run_monthly([self.workbook("s.xlsx", [("A", 1), ("B", 2)])], self.template, "2026-10")
        for prefix in ("../escape", "A/B"):
            with self.assertRaises(ValueError):
                export.export_result(result, self.root, prefix)
        with patch.object(export, "MAX_ROWS", 2), self.assertRaises(ValueError):
            export.export_result(result, self.root, "row-limit")
        with patch.object(export, "MAX_TEXT", 5), self.assertRaises(ValueError):
            export.export_result(result, self.root, "text-limit")
        self.assertFalse(list(self.root.glob("*limit*.xlsx")))

    def test_long_valid_template_prefix_can_export_with_month(self):
        template = monthly.make_template("长前缀", self.rules, output_prefix="a" * 80)
        result = monthly.run_monthly([self.workbook("long.xlsx", [("A", 1)])], template, "2026-10")
        output = Path(export.export_result(result, self.root))
        self.assertTrue(output.is_file())
        self.assertEqual(output.name, "a" * 80 + "-2026-10.xlsx")

    def test_cancel_before_read_run_and_export(self):
        event = threading.Event()
        event.set()
        path = self.workbook("cancel.xlsx", [("A", 1)])
        with self.assertRaises(analysis.CancelledError):
            analysis.read_inputs([path], self.rules, cancel=event)
        result = monthly.run_monthly([path], self.template, "2026-10")
        with self.assertRaises(analysis.CancelledError):
            export.export_result(result, self.root, "cancel-output", cancel=event)
        self.assertFalse(list(self.root.glob("cancel-output*")))

    def test_workbooks_closed_on_partial_open_failure(self):
        from unittest.mock import MagicMock
        book = MagicMock()
        with patch.object(analysis, "load_workbook", side_effect=[book, ValueError("坏公式读取")]):
            report = analysis.read_inputs([self.root / "broken.xlsx"], self.rules)
        book.close.assert_called_once()
        self.assertTrue(report.issues)

    def test_legacy_three_return_values_and_summary(self):
        reader = importlib.import_module(PACKAGE + ".excel_reader")
        writer = importlib.import_module(PACKAGE + ".excel_writer")
        mapper = importlib.import_module(PACKAGE + ".field_mapper")
        source = self.workbook("legacy.xlsx", [("A", 1), ("A", 2), ("B", -1)])
        with patch.object(mapper, "_FIELD_CONFIG", []):
            total, count, fields = reader.read_and_aggregate(str(self.root), str(self.root), [source.name])
        self.assertEqual(count, 1)
        self.assertEqual(total["A"]["金额"], 3)
        output = self.root / "legacy-result.xlsx"
        writer.generate_summary(sorted(total.items()), total, output, fields)
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book.active["B2"].value, 3)
        self.assertEqual(book.active["B4"].value, 2)


    def test_handbuilt_report_without_rules_is_viewable_and_exportable(self):
        source = analysis.Source("旧.xlsx", "表1", 2)
        record = analysis.Record("A", Decimal("1.5"), source)
        left = analysis.ReadReport(records=[record], files_loaded=1, input_files=["旧.xlsx"])
        right = analysis.ReadReport(records=[record], files_loaded=1, input_files=["旧.xlsx"])
        self.assertEqual(left.rules, {})  # No recorded rules metadata.
        result = reconciliation.compare_reports(left, right)
        self.assertTrue(result.all_consistent)
        self.assertEqual(result.config["left_key_fields"], [])   # Never fabricate fields.
        self.assertEqual(result.config["right_key_fields"], [])
        output = export.export_result(result, self.root, prefix="无规则")
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book["对账结果"]["A2"].value, "A")
        config = dict(book["运行配置"].values)
        self.assertEqual(config["left_rules"], "{}")
        self.assertNotIn("left_key_fields[1]", config)


if __name__ == "__main__":
    unittest.main()
