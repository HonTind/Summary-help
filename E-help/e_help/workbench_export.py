"""Traceable workbench reports; exclusive output creation, never overwrite input."""
from __future__ import annotations

import json
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analysis import check_cancel, date_trace_text
from .edition_policy import guarded_export
from .monthly import MonthlyResult, safe_name
from .reconciliation import CATEGORIES, ReconciliationResult

MAX_ROWS = 1048576
MAX_TEXT = 32767


def _append(sheet, values):
    if sheet.max_row >= MAX_ROWS:
        raise ValueError("数据超出 Excel 工作表行数限制，未导出；请缩小输入范围")
    # Validate before openpyxl silently truncates long strings.
    for value in values:
        if isinstance(value, str) and len(value) > MAX_TEXT:
            raise ValueError("单元格文字超出 Excel 长度限制，未导出")
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"  # Includes =... keys, formula text, file/sheet names.
            cell.number_format = "@"


def _sheet(book, name, headers):
    sheet = book.create_sheet(name)
    _append(sheet, headers)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 26
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="344054")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    return sheet


def _text(number):
    return "" if number is None else str(number)


def _selection_text(value):
    return f"#{value}" if type(value) is int else str(value)


def _populate(book, result, cancel):
    if isinstance(result, MonthlyResult):
        summary = _sheet(book, "月度汇总", ["键", "金额（精确文本）", "来源行数"])
        for row in result.rows:
            check_cancel(cancel)
            _append(summary, [row.key, str(row.amount), len(row.records)])
        _append(summary, ["合计", str(result.total), len(result.report.records)])
        reports = [("本月", result.report)]
        category = {(side, row.source): "月度汇总" for side, report in reports for row in report.records}
    elif isinstance(result, ReconciliationResult):
        summary = _sheet(book, "对账结果", ["键", "分类", "左金额（精确文本）", "右金额（精确文本）",
                                         "左减右（精确文本）", "左行数", "右行数"])
        category = {}
        for row in result.rows:
            check_cancel(cancel)
            _append(summary, [row.key, CATEGORIES[row.category], _text(row.left_amount), _text(row.right_amount),
                              _text(row.difference), len(row.left), len(row.right)])
            for side, records in (("左侧", row.left), ("右侧", row.right)):
                for record in records:
                    category[side, record.source] = CATEGORIES[row.category]
        reports = [("左侧", result.left), ("右侧", result.right)]
    else:
        book.close()
        raise ValueError("不支持的结果类型")
    # Leading columns keep their old positions; the date trace is appended, never borrowed
    # into the key column.
    details = _sheet(book, "来源明细", ["侧别", "分类", "文件", "工作表", "原行号", "键",
                                     "金额（精确文本）", "日期（原｜规范）"])
    issues = _sheet(book, "异常清单", ["侧别", "文件", "工作表", "原行号", "异常类型", "说明",
                                     "原始键", "原始金额", "原日期", "规范日期"])
    config = _sheet(book, "运行配置", ["项目", "值"])
    _append(config, ["结果完整", "是" if result.complete else "否：存在异常、歧义或无有效输入；请先核查"])
    _append(config, ["金额说明", "金额为精确文本，保留原精度；非 Excel 公式，不自动重算"])
    _append(config, ["前提", "同币种、同口径；键精确匹配，只去首尾空白"])
    if isinstance(result, MonthlyResult):
        date_column = result.report.rules.get("date_column")
        _append(config, ["日期筛选", "开启" if date_column is not None else "关闭"])
        if date_column is not None:
            _append(config, ["日期字段", _selection_text(date_column)])
            _append(config, ["日期范围", f"{result.month}-01 起，含当月全月；不含下月，跨年按自然月"])
            _append(config, ["外月跳过行数", result.report.excluded_rows])
            _append(config, ["日期排除说明", "「日期排除」行只校验了日期；键与金额为源表原值、未校验，"
                                            "原日期与规范日期见来源明细「日期（原｜规范）」列"])
    if isinstance(result, ReconciliationResult):
        _append(config, ["全部一致", "是" if result.all_consistent else "否"])
        for key, count in result.counts.items():
            _append(config, [CATEGORIES[key] + "键数", count])
    for side, report in reports:
        for record in report.records:
            check_cancel(cancel)
            source = record.source
            _append(details, [side, category[side, source], source.file, source.sheet, source.row,
                              record.key, str(record.amount),
                              date_trace_text(getattr(record, "raw_date", ""),
                                              getattr(record, "normalized_date", ""))])
        for issue in report.issues:
            check_cancel(cancel)
            source = issue.source
            _append(issues, [side, source.file, source.sheet, source.row, issue.code, issue.message,
                             issue.raw_key, issue.raw_amount,
                             getattr(issue, "raw_date", ""), getattr(issue, "normalized_date", "")])
        for excluded in getattr(report, "excluded", ()):
            check_cancel(cancel)
            source = excluded.source
            _append(details, [side, "日期排除", source.file, source.sheet, source.row,
                              getattr(excluded, "raw_key", ""), "未校验",
                              date_trace_text(excluded.raw_date, excluded.normalized)])
        for key in ("scanned_rows", "blank_rows", "invalid_rows", "files_loaded"):
            _append(config, [side + "." + key, getattr(report, key)])
    # Each file gets its own row so a long file list cannot be truncated into one cell.
    for key, value in result.config.items():
        if isinstance(value, list):
            for index, entry in enumerate(value, 1):
                _append(config, [f"{key}[{index}]", str(entry)])
        else:
            text = json.dumps(value, ensure_ascii=False, allow_nan=False) if isinstance(value, dict) else str(value)
            _append(config, [key, text])
    for sheet in book:
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(64, max(16, max(len(str(c.value or "")) for c in column) + 3))
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width
    return book


@guarded_export
def export_result(result, directory, prefix=None, *, cancel=None):
    check_cancel(cancel)
    if prefix is None:
        prefix = (f"{result.template['output_prefix']}-{result.month}" if isinstance(result, MonthlyResult)
                  else "客户总额对账" if result.mode == "customer" else "单据逐笔对账")
    safe_name(prefix, max_length=120)  # Room for a validated template prefix plus month.
    destination = Path(directory).resolve()
    if not destination.is_dir():
        raise ValueError("请选择已存在的输出目录")
    reports = [result.report] if isinstance(result, MonthlyResult) else [result.left, result.right]
    protected = {str(Path(p).resolve()).casefold() for report in reports for p in report.input_files}
    book = Workbook()
    book.remove(book.active)
    output = None
    try:
        _populate(book, result, cancel)
        for counter in range(10000):
            check_cancel(cancel)
            suffix = "" if counter == 0 else f"({counter})"
            candidate = destination / f"{prefix}{suffix}.xlsx"
            if str(candidate).casefold() in protected:
                continue
            try:
                descriptor = os.open(candidate, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
            except FileExistsError:
                continue
            output = candidate
            with os.fdopen(descriptor, "wb") as handle:
                book.save(handle)
                handle.flush()
                os.fsync(handle.fileno())
            check_cancel(cancel)
            return str(output)
        raise FileExistsError("无法分配唯一输出文件名，请更换输出前缀")
    except Exception:
        if output is not None:
            output.unlink(missing_ok=True)  # Only this invocation's exclusively-created output.
        raise
    finally:
        book.close()
