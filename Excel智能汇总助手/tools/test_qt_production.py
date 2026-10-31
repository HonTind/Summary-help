"""Real-engine Qt regressions using exclusively synthetic, temporary data."""
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
import ast
from pathlib import Path
import socket
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
import webbrowser

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT))
from openpyxl import Workbook, load_workbook
from PySide6.QtWidgets import QApplication, QDialog, QMessageBox
from Excel自动汇总助手 import field_mapper, history
from Excel自动汇总助手.analysis import SheetRules
from Excel自动汇总助手.qt_ui.backend import Backend, Request
from Excel自动汇总助手.qt_ui.application import ApplicationWindow
from Excel自动汇总助手.qt_ui.dialogs import FieldConfigDialog, MappingDialog, RulesDialog, selection
from Excel自动汇总助手.qt_ui.theme import configure_application
from tools import check_dist


class ProductionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])
        configure_application(cls.app)

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="qt-production-test-")
        self.root = Path(self.temporary.name)
        self.backend = Backend(self.root, self.root / "config", self.root / "logs")
        self.backend.config_dir.mkdir()
        self.backend.initialize()
        self.patches = [patch.object(history, "get_history_path", return_value=str(self.root / "history.json")),
            patch.object(socket.socket, "connect", side_effect=AssertionError("network forbidden")),
            patch.object(webbrowser, "open", side_effect=AssertionError("browser forbidden"))]
        for p in self.patches:
            p.start()
        self.left = self.book("left.xlsx", [("A", 10, "2026-10-01", "001", "X"),
                                             ("A", .25, "2026-10-02", "002", "X"),
                                             ("B", 9, "2026-09-01", "003", "Y")])
        self.right = self.book("right.xlsx", [("A", 10, "2026-10-01", "001", "X"),
                                               ("A", .25, "2026-10-02", "002", "X"),
                                               ("B", 8, "2026-09-01", "003", "Y")])
        self.windows = []

    def tearDown(self):
        for window in self.windows:
            if window.runner.busy:
                window.runner.cancel()
                window.runner.thread.join(timeout=5)
            window.busy = False
            window.close()
        self.app.processEvents()
        for p in reversed(self.patches):
            p.stop()
        self.temporary.cleanup()

    def book(self, name, rows):
        book = Workbook()
        book.active.append(["客户", "金额", "日期", "单据号", "部门"])
        for row in rows:
            book.active.append(row)
        path = self.root / name
        book.save(path)
        book.close()
        return str(path)

    def window(self):
        window = ApplicationWindow(self.backend)
        self.windows.append(window)
        window.show()
        self.app.processEvents()
        return window

    def wait_job(self, window):
        deadline = time.monotonic() + 5
        while window.busy and time.monotonic() < deadline:
            self.app.processEvents()
            time.sleep(.005)
        self.assertFalse(window.busy, "job failed to stop")

    def test_feedback_only_opens_fixed_url_after_confirmation(self):
        from PySide6.QtGui import QDesktopServices
        with patch.object(QDesktopServices, "openUrl", return_value=True) as opening:
            window = self.window()
            window.show_page("feedback")
            opening.assert_not_called()
            with patch.object(window, "confirm_online_feedback", return_value=False):
                window.feedback_button.click()
            opening.assert_not_called()
            with patch.object(window, "confirm_online_feedback", return_value=True):
                window.feedback_button.click()
            opening.assert_called_once()
            self.assertEqual(opening.call_args.args[0].toString(), "https://wj.qq.com/s2/28242546/gc6x/")

    def test_feedback_browser_failure_is_recoverable(self):
        from PySide6.QtGui import QDesktopServices
        window = self.window()
        for failure in (False, RuntimeError("synthetic failure")):
            with patch.object(QDesktopServices, "openUrl", return_value=False,
                              side_effect=failure if isinstance(failure, Exception) else None), \
                 patch.object(window, "confirm_online_feedback", return_value=True), \
                 patch.object(QMessageBox, "information") as message:
                window.open_online_feedback()
            self.assertIn("https://wj.qq.com/s2/28242546/gc6x/", message.call_args.args[2])
            self.assertFalse(window.busy)

    def test_feedback_confirmation_chinese_buttons_and_safe_default(self):
        window = self.window()
        for action in ("取消", "打开问卷", None):
            def interact(dialog):
                buttons = {button.text(): button for button in dialog.buttons()}
                self.assertEqual(set(buttons), {"打开问卷", "取消"})
                self.assertIs(dialog.defaultButton(), buttons["取消"])
                self.assertIs(dialog.escapeButton(), buttons["取消"])
                if action is not None:
                    buttons[action].click()
                else:
                    dialog.reject()
                return 0
            with patch.object(QMessageBox, "exec", new=interact):
                self.assertEqual(window.confirm_online_feedback(), action == "打开问卷")

    def test_edition_status_close_button_is_chinese(self):
        window = self.window()
        def inspect(dialog):
            self.assertEqual(dialog.button(QMessageBox.StandardButton.Close).text(), "关闭")
            return 0
        with patch("Excel自动汇总助手.qt_ui.application.edition_status", return_value="合成状态"), \
             patch.object(QMessageBox, "exec", new=inspect):
            window.show_edition_status()

    def test_missing_input_message_matches_mode(self):
        window = self.window()
        for mode, expected in (("regular", "输入文件夹"), ("monthly", "本月账单"), ("compare", "左侧账单")):
            window.select_mode(mode)
            window.selected_files = []
            with patch.object(QMessageBox, "warning") as warning:
                window.preview()
            self.assertIn(expected, warning.call_args.args[2])
            self.assertFalse(window.busy)
        window.selected_files = [self.left]
        window.right_files = []
        with patch.object(QMessageBox, "warning") as warning:
            window.preview()
        self.assertIn("右侧账单", warning.call_args.args[2])

    def test_compare_input_buttons_order_and_width(self):
        window = self.window()
        window.select_mode("compare")
        self.app.processEvents()
        layout = window.file_card.parentWidget().layout()
        left_index = next(i for i in range(layout.count()) if layout.itemAt(i).layout() is not None
                          and layout.itemAt(i).layout().indexOf(window.choose_files_button) >= 0)
        self.assertEqual(layout.indexOf(window.choose_right), left_index + 1)
        self.assertEqual(window.choose_files_button.width(), window.choose_right.width())

    def test_history_column_minimums_and_full_tooltips(self):
        window = self.window()
        from Excel自动汇总助手.qt_ui.demo import TableData
        window.refresh_history()
        panel = window.history_panel
        stamp = "2026-10-19 12:34:56"
        name = "很长的月度汇总结果文件名称-2026-10.xlsx"
        panel.set_data(TableData(("时间", "类型", "状态", "文件数", "结果文件"),
                                 ((stamp, "月度汇总", "已导出", 1, name),)))
        panel.fit_columns()
        self.assertGreaterEqual(panel.table.columnWidth(0), 210)
        self.assertGreaterEqual(panel.table.columnWidth(4), 300)
        self.assertEqual(panel.table.item(0, 0).toolTip(), stamp)
        self.assertEqual(panel.table.item(0, 4).toolTip(), name)

    def test_help_typography_preserves_text_and_escapes_markup(self):
        from types import SimpleNamespace
        window = self.window()
        window.article_list = [SimpleNamespace(title="安全 <标题>", body="第一步\n<script>不执行</script>\n第二步")]
        window.show_article(0)
        plain = window.help_text.toPlainText()
        self.assertIn("安全 <标题>", plain)
        self.assertIn("<script>不执行</script>", plain)
        self.assertIn("第二步", plain)
        window.show_article(-1)
        self.assertEqual(window.help_text.toPlainText(), "没有匹配的离线教程。")

    def test_sidebar_groups_secondary_navigation_at_bottom(self):
        window = self.window()
        self.assertNotIn("about", window.nav_buttons)
        self.assertNotIn("about", window.page_widgets)
        buttons = window.nav_buttons
        self.assertLess(buttons["workbench"].y(), buttons["history"].y())
        self.assertLess(buttons["history"].y(), buttons["settings"].y())
        layout = buttons["workbench"].parentWidget().layout()
        work_index = layout.indexOf(buttons["workbench"])
        history_index = layout.indexOf(buttons["history"])
        self.assertEqual(history_index, work_index + 2)
        self.assertIsNotNone(layout.itemAt(work_index + 1).spacerItem())
        window.show_page("settings")
        self.assertEqual(window.page, "settings")

    def test_product_copy_has_no_commercial_edition_labels(self):
        from PySide6.QtWidgets import QLabel, QPushButton
        from Excel自动汇总助手.config import APP_NAME, APP_VERSION, APP_FULL_NAME
        from Excel自动汇总助手.edition_policy import edition_status
        self.assertEqual(APP_FULL_NAME, f"{APP_NAME} {APP_VERSION}")
        window = self.window()
        texts = [window.windowTitle(), edition_status()]
        texts.extend(label.text() for label in window.findChildren(QLabel))
        texts.extend(button.text() for button in window.findChildren(QPushButton))
        for text in texts:
            for word in ("免费", "付费", "试用版", "赞助", "不限次数", "次数不限"):
                self.assertNotIn(word, text)

    def test_settings_shows_free_edition_without_activation(self):
        window = self.window()
        for mode in window.MODES:
            window.select_mode(mode)
            window.show_page("settings")
            self.app.processEvents()
            self.assertEqual(window.page, "settings")
            self.assertTrue(window.edition_button.isVisible())
            self.assertFalse(hasattr(window, "license_button"))
            self.assertFalse(hasattr(window, "import_paid_license"))
            self.assertEqual(window.edition_button.text(), "版本状态")
            self.assertFalse(window.advanced.body.isAncestorOf(window.edition_button))
        self.assertTrue(all("版本" not in b.text() and "授权" not in b.text()
                            for b in window.extra_buttons))

    def test_monthly_date_and_composite_rules_use_existing_engine(self):
        before = Path(self.left).read_bytes()
        request = Request("monthly", (self.left,), rules=SheetRules(key_fields=("客户", "部门"), date_column="日期"))
        result = self.backend.preview(request)
        self.assertEqual(str(result.raw.total), "10.25")
        self.assertEqual(result.counts["excluded_count"], 1)
        self.assertEqual(len(result.preview.sources.rows), 3)
        self.assertEqual(Path(self.left).read_bytes(), before)
        output = self.backend.export(result, self.root)
        again = self.backend.export(result, self.root)
        self.assertNotEqual(output, again)
        workbook = load_workbook(output, data_only=False)
        try:
            self.assertEqual(workbook["月度汇总"].cell(2, 2).value, "10.25")
        finally:
            workbook.close()

    def test_all_modes_export_more_than_ten_times_without_activation(self):
        requests = [Request("regular", (self.left,)), Request("monthly", (self.left,))]
        requests.extend(Request("compare", (self.left,), (self.right,),
            SheetRules(key_fields=keys), SheetRules(key_fields=keys), compare_mode=mode)
            for mode, keys in (("customer", ("客户",)), ("document", ("单据号",))))
        with patch.dict(os.environ, {"LOCALAPPDATA": str(self.root / "unused-profile")}):
            for request in requests:
                result = self.backend.preview(request)
                for _ in range(12):
                    self.assertTrue(self.backend.export(result, self.root).is_file())
            self.assertFalse((self.root / "unused-profile").exists())

    def test_both_reconciliation_modes_and_right_rules(self):
        for mode, keys in (("customer", ("客户", "部门")), ("document", ("单据号",))):
            result = self.backend.preview(Request("compare", (self.left,), (self.right,),
                SheetRules(key_fields=keys), SheetRules(key_fields=keys), compare_mode=mode))
            self.assertEqual(result.raw.counts["different"], 1)
            self.assertTrue(result.complete)
            self.assertTrue(self.backend.export(result, self.root).is_file())

    def test_regular_keeps_original_numeric_output_and_exclusive_names(self):
        result = self.backend.preview(Request("regular", (self.left,)))
        self.assertEqual(result.raw["total"]["A"]["金额"], 10.25)
        existing = self.root / "汇总.xlsx"
        existing.write_bytes(b"keep")
        output = self.backend.export(result, self.root)
        self.assertEqual(existing.read_bytes(), b"keep")
        workbook = load_workbook(output, data_only=False)
        try:
            self.assertIsInstance(workbook.active.cell(2, 2).value, (int, float))
        finally:
            workbook.close()

    def test_skipped_regular_file_blocks_export(self):
        broken = self.book("wrong.xlsx", [("A", 1, "2026-10-01", "1", "X")])
        workbook = load_workbook(broken)
        workbook.active.cell(1, 2).value = "未知字段"
        workbook.save(broken)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "阻止导出"):
            self.backend.preview(Request("regular", (self.left, broken)))

    def test_real_window_preview_export_history_and_invalidation(self):
        window = self.window()
        window.select_mode("monthly")
        window.selected_files = [self.left]
        window.month.setCurrentText("2026-10")
        window.preview()
        self.assertTrue(window.busy)
        self.assertFalse(window.key.isEnabled())
        window.select_mode("compare")
        self.assertEqual(window.mode, "monthly")
        self.wait_job(window)
        self.assertIsNotNone(window.result)
        self.assertTrue(window.export_button.isEnabled())
        window.explain_export()
        self.wait_job(window)
        self.assertTrue(Path(window.output_path).is_file())
        records, warning = history.read_history()
        self.assertFalse(warning)
        self.assertEqual([r["status"] for r in records], ["previewed", "success"])
        self.assertNotIn("input_files", str(records))
        self.assertNotIn("raw_amount", str(records))
        window.key.setCurrentText("单据号")
        self.assertIsNone(window.result)
        self.assertFalse(window.export_button.isEnabled())

    def test_error_clears_old_result_and_records_one_failure(self):
        window = self.window()
        window.select_mode("monthly")
        window.selected_files = [self.left]
        window.month.setCurrentText("2026-99")
        window.preview()
        self.wait_job(window)
        self.assertIsNone(window.result)
        self.assertEqual([r["status"] for r in history.load_history()], ["failed"])
        window.poll_job()
        self.assertEqual(len(history.load_history()), 1)

    def test_cancel_waits_for_actual_worker_exit(self):
        window = self.window()
        entered, release = threading.Event(), threading.Event()
        def slow(cancel, progress):
            entered.set()
            release.wait(timeout=2)
            return self.backend.preview(Request("monthly", (self.left,)), cancel)
        window.start_job(Request("monthly", (self.left,)), "preview", slow)
        self.assertTrue(entered.wait(timeout=1))
        window.cancel_job()
        window.poll_job()
        self.assertTrue(window.busy)
        self.assertFalse(window.key.isEnabled())
        release.set()
        self.wait_job(window)
        self.assertEqual(history.load_history()[0]["status"], "cancelled")
        self.assertIsNone(window.result)

    def test_completed_export_survives_late_cancel(self):
        window = self.window()
        result = self.backend.preview(Request("monthly", (self.left,)))
        def operation(cancel, progress):
            path = self.backend.export(result, self.root)
            cancel.set()
            return path
        window.result = result
        window.start_job(result.request, "export", operation, result=result)
        self.wait_job(window)
        self.assertTrue(Path(window.output_path).is_file())
        self.assertEqual(history.load_history()[0]["status"], "success")

    def test_preview_late_cancel_discards_queued_result(self):
        window = self.window()
        request = Request("monthly", (self.left,))
        window.start_job(request, "preview", lambda cancel, progress: self.backend.preview(request, cancel))
        window.runner.thread.join(timeout=5)
        self.assertFalse(window.runner.busy)
        window.cancel_job()
        self.wait_job(window)
        self.assertIsNone(window.result)
        self.assertEqual(history.load_history()[0]["status"], "cancelled")

    def test_incomplete_result_requires_export_confirmation(self):
        window = self.window()
        window.select_mode("monthly")
        window.selected_files = [self.left]
        window.amount.setCurrentText("不存在")
        window.preview()
        self.wait_job(window)
        self.assertFalse(window.result.complete)
        self.assertEqual(len(window.result.preview.issues.rows), 1)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No):
            window.explain_export()
        self.assertFalse(window.busy)
        self.assertFalse((self.root / "Excel汇总结果").exists())

    def test_missing_selected_regular_file_is_not_silently_ignored(self):
        with self.assertRaisesRegex(ValueError, "阻止导出"):
            self.backend.preview(Request("regular", (self.left, str(self.root / "missing.xlsx"))))

    def test_template_persistence_confirmation_and_corrupt_file_preserved(self):
        window = self.window()
        window.select_mode("monthly")
        window.template.setCurrentText("规则A")
        window.save_memory_template()
        before = self.backend.template_path.read_bytes()
        window.amount.setCurrentText("#2")
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No):
            window.save_memory_template()
        self.assertEqual(self.backend.template_path.read_bytes(), before)
        window.key.setCurrentText("#1")
        window.load_memory_template()
        self.assertEqual(window.key.currentText(), "客户")
        self.backend.template_path.write_text("broken", encoding="utf-8")
        with patch.object(QMessageBox, "warning"):
            window.save_memory_template()
        self.assertEqual(self.backend.template_path.read_text(), "broken")

    def test_history_delete_does_not_delete_output_and_bad_store_survives(self):
        window = self.window()
        result = self.backend.preview(Request("monthly", (self.left,)))
        output = self.backend.export(result, self.root)
        self.backend.record(result.request, "success", "export", result=result, output=output)
        window.show_page("history")
        window.history_panel.table.selectRow(0)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            window.delete_history()
        self.assertTrue(output.is_file())
        self.assertEqual(history.load_history(), [])
        path = self.root / "history.json"
        path.write_text('{"version":99,"records":[]}', encoding="utf-8")
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes), patch.object(QMessageBox, "warning"):
            window.clear_history()
        self.assertIn('99', path.read_text())

    def test_preferences_preserve_other_settings_and_reject_bad_json(self):
        path = self.backend.config_dir / "settings.json"
        path.write_text('{"field_mapping":{"旧字段":"客户"}}', encoding="utf-8")
        self.backend.save_preferences("custom", "local")
        self.assertEqual(self.backend.preferences()["field_mapping"], {"旧字段": "客户"})
        path.write_text("broken", encoding="utf-8")
        with self.assertRaises(ValueError):
            self.backend.save_preferences("default", "")
        self.assertEqual(path.read_text(), "broken")

    def test_native_rules_mapping_and_config_dialogs(self):
        window = self.window()
        dialog = RulesDialog(window, SheetRules(key_fields=("客户", "部门")))
        dialog.validate()
        self.assertEqual(dialog.value.key_fields, ("客户", "部门"))
        self.assertEqual(selection("#2"), 2)
        dialog.close()
        mapping = MappingDialog(window, self.left, {})
        mapping.validate()
        self.assertEqual(mapping.value["客户"], "客户")
        mapping.close()
        config = FieldConfigDialog(window)
        config.validate()
        self.assertGreaterEqual(len(config.value), 2)
        config.close()

    def test_minimum_window_actions_accessible_and_no_network_modules(self):
        window = self.window()
        window.resize(840, 580)
        self.app.processEvents()
        for button in (window.preview_button, window.export_button):
            top_left = button.mapTo(window, button.rect().topLeft())
            self.assertLessEqual(top_left.x() + button.width(), window.width())
            self.assertLessEqual(top_left.y() + button.height(), window.height())
        self.assertNotIn("PySide6.QtNetwork", sys.modules)
        self.assertNotIn("PySide6.QtWebEngineWidgets", sys.modules)
        self.assertNotIn("Excel自动汇总助手.gui", sys.modules)

    def test_qt_distribution_rejects_foreign_icu(self):
        folder = self.root / "fake-distribution"
        for relative, _description in [*check_dist.REQUIRED_FILES, *check_dist.QT_REQUIRED_FILES]:
            path = folder / relative
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(b"synthetic-not-executed")
        with patch.object(check_dist, "MIN_EXE_SIZE", 1), patch.object(check_dist, "MIN_PYD_COUNT", 0), patch.object(check_dist, "MIN_DLL_COUNT", 0):
            self.assertTrue(check_dist.check_integrity(folder, qt=True)[0])
            (folder / "_internal/icuuc.dll").write_bytes(b"foreign-icu-not-executed")
            self.assertFalse(check_dist.check_integrity(folder, qt=True)[0])

    def test_both_specs_filter_path_icu_and_preserve_qt_runtime(self):
        from types import SimpleNamespace
        for name in ("Excel智能汇总助手.spec", "Excel智能汇总助手_onedir.spec"):
            tree = ast.parse((PROJECT / name).read_text(encoding="utf-8"))
            assignments = [node for node in tree.body if isinstance(node, ast.Assign)
                and any(isinstance(target, ast.Attribute) and target.attr == "binaries"
                        for target in node.targets)]
            self.assertEqual(len(assignments), 1)
            module = ast.Module(body=assignments, type_ignores=[])
            context = {"os": os, "a": SimpleNamespace(binaries=[
                ("icuuc.dll", "foreign-source", "BINARY"),
                ("icudt78.dll", "foreign-data", "BINARY"),
                ("PySide6/Qt6Core.dll", "qt-source", "BINARY")])}
            exec(compile(module, name, "exec"), context)
            self.assertEqual(context["a"].binaries, [("PySide6/Qt6Core.dll", "qt-source", "BINARY")])


if __name__ == "__main__":
    unittest.main()
