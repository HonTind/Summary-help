"""
v1.3 GUI 服务层接口验证

测试:
  A1. get_file_headers — 正常文件 → 返回表头列表
  A2. get_file_headers — 损坏文件 → 返回 False + 错误信息
  A3. get_file_headers — 空表头文件 → 返回 False
  B1. get_all_mappings — 返回当前映射副本
  B2. get_all_mappings — 修改返回值不影响原字典
  C1. remove_field_mapping — 删除存在的映射 → True
  C2. remove_field_mapping — 删除不存在的映射 → False
  C3. remove_field_mapping — 内存和 JSON 同步删除
  D1. 回归: detect_columns 不受影响
  D2. 回归: save_field_mapping + load_settings 不受影响

运行:
  cd "Day07(小试牛刀)"
  python -m Excel自动汇总助手.tests.verify_v13_gui_api
"""

import os, sys, json, shutil

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PARENT_DIR = os.path.dirname(PROJECT_ROOT)
if PARENT_DIR not in sys.path:
    sys.path.insert(0, PARENT_DIR)

from openpyxl import Workbook
import Excel自动汇总助手.config as config_module
from Excel自动汇总助手.field_mapper import (
    get_file_headers,
    get_all_mappings,
    remove_field_mapping,
    save_field_mapping,
    load_settings,
    detect_columns,
    get_header_list,
)

TEST_DIR = os.path.join(PROJECT_ROOT, "tests", "_v13_gui_api_test")
if os.path.isdir(TEST_DIR):
    shutil.rmtree(TEST_DIR)
os.makedirs(TEST_DIR)

_original_map = dict(config_module.CUSTOM_HEADER_MAP)

def reset():
    config_module.CUSTOM_HEADER_MAP.clear()
    config_module.CUSTOM_HEADER_MAP.update(_original_map)

def count_pass(ok):
    global passed
    if ok:
        passed += 1

passed = 0
total = 0

print("=" * 55)
print("  v1.3 GUI 服务层接口验证")
print("=" * 55)

# ═══════════════════════════════════════════════════
#  A1. get_file_headers — 正常文件
# ═══════════════════════════════════════════════════
total += 1
print("\n[A1] get_file_headers — 正常文件")

normal = os.path.join(TEST_DIR, "normal.xlsx")
wb = Workbook()
ws = wb.active
ws.append(["客户名称", "金额", "日期"])
ws.append(["公司A", 1000, "2024-01-01"])
wb.save(normal)
wb.close()

ok, result = get_file_headers(normal)
assert ok, f"应返回 True，实际 {ok}: {result}"
assert result == [(0, "客户名称"), (1, "金额"), (2, "日期")], \
    f"表头不匹配: {result}"
count_pass(True)
print("  PASS")

# ═══════════════════════════════════════════════════
#  A2. get_file_headers — 损坏文件
# ═══════════════════════════════════════════════════
total += 1
print("\n[A2] get_file_headers — 损坏文件")

corrupt = os.path.join(TEST_DIR, "corrupt.xlsx")
with open(corrupt, "w", encoding="utf-8") as f:
    f.write("not a zip file")

ok, result = get_file_headers(corrupt)
assert not ok, f"损坏文件应返回 False: {ok}"
assert "无法读取文件" in result, f"错误信息格式不对: {result}"
count_pass(True)
print("  PASS")

# ═══════════════════════════════════════════════════
#  A3. get_file_headers — 空表头
# ═══════════════════════════════════════════════════
total += 1
print("\n[A3] get_file_headers — 空表头")

empty = os.path.join(TEST_DIR, "empty.xlsx")
wb2 = Workbook()
wb2.save(empty)
wb2.close()

ok, result = get_file_headers(empty)
assert not ok, f"空表头应返回 False: {ok}"
assert "表头为空" in result, f"错误信息不对: {result}"
count_pass(True)
print("  PASS")

# ═══════════════════════════════════════════════════
#  B1. get_all_mappings — 返回副本
# ═══════════════════════════════════════════════════
total += 1
print("\n[B1] get_all_mappings — 返回当前映射副本")

reset()
# v1.5.0: 使用字段名（中文）
config_module.CUSTOM_HEADER_MAP["测试A"] = "客户"
config_module.CUSTOM_HEADER_MAP["测试B"] = "金额"

mappings = get_all_mappings()
assert mappings.get("测试A") == "客户"
assert mappings.get("测试B") == "金额"
count_pass(True)
print("  PASS")

# ═══════════════════════════════════════════════════
#  B2. get_all_mappings — 修改返回值不影响原字典
# ═══════════════════════════════════════════════════
total += 1
print("\n[B2] get_all_mappings — 修改返回值不影响原字典")

mappings["测试A"] = "金额"  # 试图篡改
assert config_module.CUSTOM_HEADER_MAP["测试A"] == "客户", \
    "返回值是浅拷贝，不应影响 CUSTOM_HEADER_MAP"
count_pass(True)
print("  PASS")

# 清理
del config_module.CUSTOM_HEADER_MAP["测试A"]
del config_module.CUSTOM_HEADER_MAP["测试B"]

# ═══════════════════════════════════════════════════
#  C1. remove_field_mapping — 删除存在的
# ═══════════════════════════════════════════════════
total += 1
print("\n[C1] remove_field_mapping — 删除存在的映射")

reset()
save_field_mapping(TEST_DIR, "待删除列", "客户")
assert "待删除列" in config_module.CUSTOM_HEADER_MAP, "save 后应在内存中"

result = remove_field_mapping(TEST_DIR, "待删除列")
assert result is True, f"删除存在的映射应返回 True: {result}"
assert "待删除列" not in config_module.CUSTOM_HEADER_MAP, "remove 后不应在内存中"

# 验证 JSON 中也删除了
settings_path = os.path.join(TEST_DIR, "settings.json")
with open(settings_path, "r", encoding="utf-8") as f:
    s = json.load(f)
assert "待删除列" not in s.get("field_mapping", {}), "JSON 中也应删除"
count_pass(True)
print("  PASS")

# ═══════════════════════════════════════════════════
#  C2. remove_field_mapping — 删除不存在的
# ═══════════════════════════════════════════════════
total += 1
print("\n[C2] remove_field_mapping — 删除不存在的映射")

result = remove_field_mapping(TEST_DIR, "根本不存在的列")
assert result is False, f"删除不存在的映射应返回 False: {result}"
count_pass(True)
print("  PASS")

# ═══════════════════════════════════════════════════
#  C3. remove_field_mapping — 不影响同 JSON 中其他映射
# ═══════════════════════════════════════════════════
total += 1
print("\n[C3] remove_field_mapping — 不影响其他映射")

reset()
save_field_mapping(TEST_DIR, "保留A", "客户")
save_field_mapping(TEST_DIR, "保留B", "金额")
save_field_mapping(TEST_DIR, "删除我", "客户")

remove_field_mapping(TEST_DIR, "删除我")

m = get_all_mappings()
assert "保留A" in m and "保留B" in m, f"其他映射应保留: {m}"
assert "删除我" not in m, f"已删除的不应存在: {m}"

# JSON 验证
with open(settings_path, "r", encoding="utf-8") as f:
    s = json.load(f)
fm = s.get("field_mapping", {})
assert "保留A" in fm and "保留B" in fm
assert "删除我" not in fm
count_pass(True)
print("  PASS")

# ═══════════════════════════════════════════════════
#  D1. 回归: detect_columns 不受影响
# ═══════════════════════════════════════════════════
total += 1
print("\n[D1] 回归: detect_columns 正常识别")

wb3 = Workbook()
ws3 = wb3.active
ws3.append(["客户", "金额"])
ws3.append(["公司A", 1000])

detections = detect_columns(ws3)
assert detections is not None
# v1.5.0: detect_columns 返回 list[FieldDetection]
assert len(detections) >= 2
assert detections[0].col_index == 0 and detections[1].col_index == 1

# 新映射也应生效
reset()
save_field_mapping(TEST_DIR, "买方", "客户")
ws3["A1"] = "买方"
ws3["B1"] = "金额"
detections2 = detect_columns(ws3)
assert detections2 is not None
assert detections2[0].col_index == 0
wb3.close()
count_pass(True)
print("  PASS")

# ═══════════════════════════════════════════════════
#  D2. 回归: save + load 循环
# ═══════════════════════════════════════════════════
total += 1
print("\n[D2] 回归: save → 重启 → load 循环")

reset()
save_field_mapping(TEST_DIR, "循环测试", "金额")

# 模拟重启
config_module.CUSTOM_HEADER_MAP.clear()
config_module.CUSTOM_HEADER_MAP.update(_original_map)

load_settings(TEST_DIR)
assert config_module.CUSTOM_HEADER_MAP.get("循环测试") == "金额", \
    "重启后映射应恢复"
count_pass(True)
print("  PASS")

# ═══════════════════════════════════════════════════
#  清理
# ═══════════════════════════════════════════════════
reset()
shutil.rmtree(TEST_DIR)

print(f"\n{'=' * 55}")
print(f"  {passed}/{total} 全部通过")
print(f"{'=' * 55}")
