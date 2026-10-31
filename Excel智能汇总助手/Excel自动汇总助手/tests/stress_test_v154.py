# -*- coding: utf-8 -*-
"""
v1.5.4 — 正式压力测试脚本

测试目标（不改动任何生产代码，仅测试）：
  1. 文件数量压力（T1-T4：1000/3000/5000/10000 个 Excel）
  2. 单文件大数据压力（T5-T7：1 个 Excel × 1万/5万/10万 行）
  3. 多文件 × 多行混合压力（T8-T10：100×1万 / 500×5000 / 1000×1k~5k）
  4. 异常文件（空表 / 缺客户 / 缺金额 / 金额非数字 / 表头乱序 / 多余字段 /
     空白行 / 大量空白行 / .xlsx / .xlsm / .xls）
  5. 字段映射场景（临时映射不记住 / 下一任务不继承 / 核心字段无法识别时阻止汇总）
  6. 内存泄漏（同一任务连续运行 10 次，观察 RAM 是否持续增长）
  7. 长时间连续运行（单次大任务 RAM 曲线 / CPU / 完整输出）
  8. 汇总结果正确性（理论值 vs 实际输出逐客户比对）

测量指标（每组）：
  - 文件数 / 总行数 / 输入总大小
  - 开始 RAM / 峰值 RAM（后台采样线程）/ 结束 RAM / RAM 增量
  - CPU 时间（process_time）/ 总耗时（wall clock）
  - 是否崩溃 / 是否卡死 / 是否读取异常
  - 是否生成输出 / 输出文件大小 / 客户数量 / 汇总金额
  - 正确性：客户数一致 / 每客户金额一致 / 总金额一致 / 无字段错位 / 无丢失 / 无重复

运行（可按阶段分次运行）：
  cd "Day07(小试牛刀)源文件"
  python -m Excel自动汇总助手.tests.stress_test_v154 --tests count     # T1-T4
  python -m Excel自动汇总助手.tests.stress_test_v154 --tests single    # T5-T7
  python -m Excel自动汇总助手.tests.stress_test_v154 --tests mixed     # T8-T10
  python -m Excel自动汇总助手.tests.stress_test_v154 --tests abnormal  # 异常文件
  python -m Excel自动汇总助手.tests.stress_test_v154 --tests mapping   # 字段映射
  python -m Excel自动汇总助手.tests.stress_test_v154 --tests leak      # 10 次泄漏
  python -m Excel自动汇总助手.tests.stress_test_v154 --tests longrun   # 长时运行
  python -m Excel自动汇总助手.tests.stress_test_v154 --tests all       # 全部

测试数据与结果统一保留在 <工程根>/stress_v154_data/ 下（不自动删除）。
"""

import os
import sys
import gc
import json
import time
import shutil
import random
import ctypes
import threading
import tracemalloc
import datetime

THIS_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(THIS_DIR)
PARENT_DIR = os.path.dirname(PROJECT_ROOT)
if PARENT_DIR not in sys.path:
    sys.path.insert(0, PARENT_DIR)

from Excel自动汇总助手.excel_reader import aggregate_with_report
from Excel自动汇总助手.processor import sort_by_amount
from Excel自动汇总助手.excel_writer import generate_summary
from Excel自动汇总助手.field_mapper import (
    save_field_config, reset_field_config,
)
from Excel自动汇总助手 import field_mapper as fm

# ------------------------------------------------------------------
#  常量与配置
# ------------------------------------------------------------------

DATA_ROOT = os.path.join(PROJECT_ROOT, "stress_v154_data")
INPUT_DIR = os.path.join(DATA_ROOT, "input")
OUTPUT_DIR = os.path.join(DATA_ROOT, "output")
RESULT_DIR = os.path.join(DATA_ROOT, "results")
ABNORMAL_DIR = os.path.join(DATA_ROOT, "abnormal")
CONFIG_DIR = os.path.join(DATA_ROOT, "config")      # field_config.json
LOGS_DIR = os.path.join(DATA_ROOT, "logs")          # error.log（避免污染真实日志）

RESULT_JSON = os.path.join(RESULT_DIR, "stress_v154_results.json")
REPORT_TXT = os.path.join(RESULT_DIR, "stress_v154_report.txt")

# 4 字段配置：客户 + 金额 + 电话 + 供应商
FIELD_CONFIG = [
    {"name": "客户", "keywords": ["客户", "客户名称", "购货单位", "单位名称", "公司", "顾客", "收货人", "买方"]},
    {"name": "金额", "keywords": ["金额", "总金额", "合计金额", "价税合计", "货款", "费用", "小计", "应收金额"]},
    {"name": "电话", "keywords": ["电话", "联系电话", "手机", "联系方式"]},
    {"name": "供应商", "keywords": ["供应商", "供货单位", "厂商", "供应方"]},
]
FIELD_NAMES = ["客户", "金额", "电话", "供应商"]

N_CUSTOMERS = 200
CUSTOMERS = [f"客户{i:03d}" for i in range(1, N_CUSTOMERS + 1)]
PHONE = {c: f"1{i:010d}" for i, c in enumerate(CUSTOMERS, start=1)}
SUPPLIER = {c: f"供应商{(i % 20) + 1:02d}" for i, c in enumerate(CUSTOMERS)}


def _ensure_dirs():
    for d in (DATA_ROOT, INPUT_DIR, OUTPUT_DIR, RESULT_DIR, ABNORMAL_DIR, CONFIG_DIR, LOGS_DIR):
        os.makedirs(d, exist_ok=True)


# ------------------------------------------------------------------
#  内存 / CPU 测量辅助
# ------------------------------------------------------------------

# 注意：_PMC / _PMC_PTR 必须定义在模块级，只创建一次。
# ctypes 会把 Structure 子类和 POINTER 类型缓存起来并长期持有引用；
# 若把 class 定义放在 _get_rss 内部，高频采样线程每次调用都会新建并泄漏
# 一个类型对象，污染内存泄漏测量（这是测试工具自身的泄漏，非被测代码泄漏）。
from ctypes import wintypes


class _PMC(ctypes.Structure):
    _fields_ = [
        ("cb", wintypes.DWORD),
        ("PageFaultCount", wintypes.DWORD),
        ("PeakWorkingSetSize", ctypes.c_size_t),
        ("WorkingSetSize", ctypes.c_size_t),
        ("QuotaPeakPagedPoolUsage", ctypes.c_size_t),
        ("QuotaPagedPoolUsage", ctypes.c_size_t),
        ("QuotaPeakNonPagedPoolUsage", ctypes.c_size_t),
        ("QuotaNonPagedPoolUsage", ctypes.c_size_t),
        ("PagefileUsage", ctypes.c_size_t),
        ("PeakPagefileUsage", ctypes.c_size_t),
    ]


_PMC_PTR = ctypes.POINTER(_PMC)


def _get_rss():
    """Windows 下通过 ctypes 读取当前进程工作集（RSS，字节）。失败返回 0。"""
    try:
        psapi = ctypes.windll.psapi
        kernel32 = ctypes.windll.kernel32
        kernel32.GetCurrentProcess.restype = wintypes.HANDLE
        psapi.GetProcessMemoryInfo.argtypes = [
            wintypes.HANDLE, _PMC_PTR, wintypes.DWORD,
        ]
        psapi.GetProcessMemoryInfo.restype = wintypes.BOOL
        pmc = _PMC()
        pmc.cb = ctypes.sizeof(pmc)
        psapi.GetProcessMemoryInfo(kernel32.GetCurrentProcess(), ctypes.byref(pmc), pmc.cb)
        return int(pmc.WorkingSetSize)
    except Exception:
        return 0


class RssSampler(threading.Thread):
    """后台线程高频采样 RSS，捕获真实峰值内存（含 openpyxl C 层）。"""

    def __init__(self, interval=0.05):
        super().__init__(daemon=True)
        self.interval = interval
        self.peak = 0
        self.samples = []
        self._stop_event = threading.Event()

    def run(self):
        while not self._stop_event.is_set():
            rss = _get_rss()
            if rss > self.peak:
                self.peak = rss
            self.samples.append((time.time(), rss))
            self._stop_event.wait(self.interval)
        # 收尾再采一次
        rss = _get_rss()
        if rss > self.peak:
            self.peak = rss

    def stop(self):
        self._stop_event.set()
        self.join(timeout=5)


def _mb(n):
    return f"{n / (1024 * 1024):.2f}"


# ------------------------------------------------------------------
#  测试数据生成（写 .xlsx 用 write_only，大幅提速；理论答案同步累积）
# ------------------------------------------------------------------

def _write_xlsx_fast(path, headers, rows):
    from openpyxl import Workbook
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(path)


def _write_xls(path, headers, rows):
    import xlwt
    wb = xlwt.Workbook()
    sh = wb.add_sheet("S")
    for c, h in enumerate(headers):
        sh.write(0, c, h)
    for ri, r in enumerate(rows, start=1):
        for ci, v in enumerate(r):
            sh.write(ri, ci, v)
    wb.save(path)


def _gen_rows(rng, n_rows, expected):
    """生成 n_rows 行真实业务数据，并把金额累加到 expected。"""
    rows = []
    for _ in range(n_rows):
        c = CUSTOMERS[rng.randrange(N_CUSTOMERS)]
        amt = rng.randint(10, 5000)
        rows.append([c, amt, PHONE[c], SUPPLIER[c]])
        expected[c] = expected.get(c, 0) + amt
    return rows


def _rows_per_file(i, n_files, kind):
    """按文件序号返回行数（制造不同大小的文件）。kind: count/single/mixed。"""
    if kind == "count":
        if i % 100 == 0:
            return 50
        return (i % 15) + 3
    # 其它由调用方显式传 rows 列表，不使用此函数
    return 3


def generate_files(folder, n_files, expected, kind="count", rng_seed=20260822,
                   rows_list=None, ext_mix=(80, 10, 10)):
    """生成 n_files 个文件，返回 (总输入字节数, 总数据行数)。

    rows_list: 若提供，则 rows_list[i] 为第 i 个文件的行数（用于 mixed/single）。
    ext_mix: (xlsx%, xls%, xlsm%) 百分比。
    """
    headers = FIELD_NAMES
    total_size = 0
    total_rows = 0
    rng = random.Random(rng_seed)
    xlsx_p, xls_p, xlsm_p = ext_mix

    for i in range(n_files):
        n_rows = rows_list[i] if rows_list is not None else _rows_per_file(i, n_files, kind)

        roll = i % 100
        if roll < xlsx_p:
            ext = ".xlsx"
        elif roll < xlsx_p + xls_p:
            ext = ".xls"
        else:
            ext = ".xlsm"

        fpath = os.path.join(folder, f"f_{i:05d}{ext}")
        rows = _gen_rows(rng, n_rows, expected)
        total_rows += n_rows

        if ext == ".xls":
            _write_xls(fpath, headers, rows)
        else:
            _write_xlsx_fast(fpath, headers, rows)

        total_size += os.path.getsize(fpath)

    return total_size, total_rows


# ------------------------------------------------------------------
#  汇总 + 生成输出 + 验证（理论值 vs 实际）
# ------------------------------------------------------------------

def aggregate_and_generate(folder, output_path, base_dir=None, field_mappings=None):
    """调用生产链路：aggregate_with_report → sort_by_amount → generate_summary。

    返回 (report, sorted_data)。
    """
    base_dir = base_dir or LOGS_DIR
    report = aggregate_with_report(folder, base_dir, None, field_mappings)
    sorted_data = sort_by_amount(report["total"], report["field_names"])
    generate_summary(sorted_data, report["total"], output_path, report["field_names"])
    return report, sorted_data


def verify_output(output_path, expected_amount, n_files, file_count, field_names,
                  expected_phone=None, expected_supplier=None):
    """读回输出 Excel，逐客户比对理论值。

    Returns:
        (ok: bool, detail: str, extra: dict)  extra 含 customers/total_sum
    """
    from openpyxl import load_workbook

    extra = {"customers": len(expected_amount), "total_sum": sum(expected_amount.values())}

    if file_count != n_files:
        return False, f"文件数不符：file_count={file_count} != 期望 {n_files}", extra
    if list(field_names) != FIELD_NAMES:
        return False, f"字段顺序不符：{field_names}", extra
    if not os.path.isfile(output_path):
        return False, "输出文件未生成", extra

    try:
        wb = load_workbook(output_path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception as e:
        return False, f"输出文件读取失败：{e}", extra

    if not rows:
        return False, "输出文件为空", extra

    header = list(rows[0])
    if header != FIELD_NAMES:
        return False, f"输出表头不符：{header}", extra

    data_rows = rows[1:-1]
    sum_row = rows[-1]

    if len(data_rows) != len(expected_amount):
        return False, f"客户数不符：输出 {len(data_rows)} != 期望 {len(expected_amount)}", extra

    out = {}
    for row in data_rows:
        cust, amt = row[0], row[1]
        out[cust] = (amt, row[2], row[3])

    for c, exp_amt in expected_amount.items():
        if c not in out:
            return False, f"客户 [{c}] 在输出中缺失", extra
        amt, phone, sup = out[c]
        if abs(float(amt) - float(exp_amt)) > 0.001:
            return False, f"客户 [{c}] 金额不符：{amt} != {exp_amt}", extra
        if expected_phone is not None and str(phone) != expected_phone[c]:
            return False, f"客户 [{c}] 电话错位：{phone} != {expected_phone[c]}", extra
        if expected_supplier is not None and str(sup) != expected_supplier[c]:
            return False, f"客户 [{c}] 供应商错位：{sup} != {expected_supplier[c]}", extra

    exp_sum = sum(expected_amount.values())
    if abs(float(sum_row[1]) - float(exp_sum)) > 0.5:
        return False, f"合计行不符：{sum_row[1]} != {exp_sum}", extra

    return True, "OK", extra


def run_case(name, folder, n_files, expected, rows_list=None, kind="count",
             field_mappings=None, rng_seed=20260822, ext_mix=(85, 5, 10)):
    """运行单个压力用例，返回结果 dict。"""
    _ensure_dirs()
    os.makedirs(folder, exist_ok=True)
    output_path = os.path.join(OUTPUT_DIR, f"out_{name}.xlsx")

    expected.clear()
    gen_t0 = time.time()
    total_size, total_rows = generate_files(
        folder, n_files, expected, kind=kind, rng_seed=rng_seed, rows_list=rows_list,
        ext_mix=ext_mix,
    )
    gen_t1 = time.time()

    gc.collect()
    start_rss = _get_rss()
    sampler = RssSampler()
    p0 = time.process_time()
    t0 = time.time()
    sampler.start()

    success = False
    correct = False
    error = ""
    detail = ""
    extra = {}
    file_count = 0
    field_names = []
    skipped = []

    try:
        report, _ = aggregate_and_generate(folder, output_path, field_mappings=field_mappings)
        file_count = report["file_count"]
        field_names = report["field_names"]
        skipped = report["skipped"]
        success = True
    except Exception as e:
        import traceback
        error = f"{type(e).__name__}: {e}"
        traceback.print_exc()

    t1 = time.time()
    p1 = time.process_time()
    sampler.stop()
    peak_rss = sampler.peak

    gc.collect()
    end_rss = _get_rss()

    output_size = os.path.getsize(output_path) if os.path.isfile(output_path) else 0

    if success:
        correct, detail, extra = verify_output(
            output_path, expected, n_files, file_count, field_names,
            expected_phone=PHONE, expected_supplier=SUPPLIER,
        )

    return {
        "name": name,
        "n_files": n_files,
        "total_rows": total_rows,
        "total_size": total_size,
        "start_rss": start_rss,
        "peak_rss": peak_rss,
        "end_rss": end_rss,
        "rss_delta": end_rss - start_rss,
        "cpu_seconds": p1 - p0,
        "elapsed": t1 - t0,
        "gen_seconds": gen_t1 - gen_t0,
        "success": success,
        "correct": correct,
        "detail": detail,
        "error": error,
        "output_size": output_size,
        "customers": len(expected),
        "total_sum": sum(expected.values()),
        "skipped": len(skipped),
    }


# ------------------------------------------------------------------
#  结果持久化
# ------------------------------------------------------------------

def _load_results():
    if os.path.isfile(RESULT_JSON):
        with open(RESULT_JSON, "r", encoding="utf-8") as f:
            return json.load(f)
    return []


def _save_results(results):
    with open(RESULT_JSON, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)


def _fmt(n):
    return f"{n / (1024 * 1024):.2f}"


def _fmt_mb_opt(n):
    """格式化 MB，缺省（None/缺失）返回 '—'。"""
    if n is None:
        return "—"
    return _fmt(n)


def build_report(results):
    lines = []
    lines.append("=" * 92)
    lines.append("Excel智能汇总助手 v1.5.4 — 正式压力测试报告")
    lines.append("=" * 92)
    lines.append(f"生成时间：{datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append(f"字段配置：{' / '.join(FIELD_NAMES)}（4 字段）")
    lines.append(f"客户池：{N_CUSTOMERS} 个（有界，符合对账实际场景）")
    lines.append("")

    numeric = [r for r in results if "peak_rss" in r]
    special = [r for r in results if "peak_rss" not in r]

    lines.append("-" * 92)
    lines.append("A. 数据规模压力汇总表（T1-T10 + 长时间运行）")
    lines.append("-" * 92)
    hdr = (
        f"{'测试':<10} | {'文件数':>7} | {'总行数':>9} | {'输入MB':>8} | "
        f"{'峰值MB':>8} | {'RAM增MB':>8} | {'耗时s':>7} | {'结果':>6}"
    )
    lines.append(hdr)
    lines.append("-" * 92)
    for r in numeric:
        result_text = "PASS" if (r.get("success") and r.get("correct")) else "FAIL"
        ram_delta = r.get("rss_delta")
        if ram_delta is None and r.get("start_rss") is not None and r.get("end_rss") is not None:
            ram_delta = r["end_rss"] - r["start_rss"]
        lines.append(
            f"{r.get('name','?'):<10} | {r.get('n_files',0):>7} | {r.get('total_rows',0):>9} | "
            f"{_fmt_mb_opt(r.get('total_size')):>8} | {_fmt(r.get('peak_rss') or 0):>8} | "
            f"{_fmt_mb_opt(ram_delta):>8} | {r.get('elapsed',0):>7.1f} | {result_text:>6}"
        )
    lines.append("-" * 92)
    lines.append("")

    lines.append("B. 专项测试汇总")
    lines.append("-" * 92)
    for r in special:
        result_text = "PASS" if (r.get("success") and r.get("correct")) else "FAIL"
        lines.append(f"  [{r.get('name')}]  {result_text}  {r.get('detail','')}")
    lines.append("")

    lines.append("详细记录")
    lines.append("~" * 92)
    for r in results:
        name = r.get("name", "?")

        if name == "ABNORMAL":
            lines.append("【异常文件测试】")
            for c in r.get("cases", []):
                lines.append(f"  {c['file']:<24} processed={c['processed']} "
                             f"skipped={[s['reason'] for s in c['skipped']]}")
            lines.append(f"  混合场景（好文件+坏文件）跳过：{r.get('mixed_skipped')}，"
                         f"阻止汇总={r.get('mixed_blocked')}")
            lines.append(f"  结论：{r.get('detail','')}")
            lines.append("")
            continue

        if name == "MAPPING":
            lines.append("【字段映射测试】")
            lines.append(f"  临时映射（不记住）当前任务成功      : {r.get('temp_ok')}")
            lines.append(f"  下一任务不继承临时映射              : {r.get('no_inherit')}")
            lines.append(f"  不污染持久化 CUSTOM_HEADER_MAP      : {r.get('no_pollution')}")
            lines.append(f"  核心字段无法识别时阻止汇总          : {r.get('block_ok')}")
            lines.append(f"  结论：{r.get('detail','')}")
            lines.append("")
            continue

        if name == "LEAK":
            lines.append("【内存泄漏测试：同一任务连续 10 次】")
            for run in r.get("runs", []):
                lines.append(
                    f"  第{run['run']:>2}次：开始 {_mb(run['start_rss'])} / "
                    f"峰值 {_mb(run['peak_rss'])} / 结束 {_mb(run['end_rss'])} / "
                    f"{run['elapsed']:.2f}s"
                )
            lines.append(f"  结束 RAM 序列：{[_mb(v) for v in r.get('end_series', [])]}")
            lines.append(f"  峰值 RAM 序列：{[_mb(v) for v in r.get('peak_series', [])]}")
            lines.append(f"  末-首结束 RAM：{_mb(r.get('trend_bytes', 0))}")
            lines.append(f"  疑似泄漏判定：{r.get('suspected_leak')}")
            lines.append(f"  tracemalloc 保留 Top：{r.get('top_retained')}")
            lines.append(f"  结论：{r.get('detail','')}")
            lines.append("")
            continue

        if name == "LONGRUN":
            lines.append("【长时间连续运行】")
            lines.append(f"  文件数 / 总行数 : {r.get('n_files')} / {r.get('total_rows')}")
            lines.append(f"  开始 / 结束时间 : {datetime.datetime.fromtimestamp(r.get('start_time',0)).strftime('%H:%M:%S')}"
                         f" → {datetime.datetime.fromtimestamp(r.get('end_time',0)).strftime('%H:%M:%S')}")
            lines.append(f"  总耗时 / CPU 时间: {r.get('elapsed',0):.1f}s / {r.get('cpu_seconds',0):.1f}s")
            lines.append(f"  开始 / 峰值 / 结束 RAM: {_mb(r.get('start_rss',0))} / "
                         f"{_mb(r.get('peak_rss',0))} / {_mb(r.get('end_rss',0))}")
            lines.append(f"  RAM 曲线采样点数: {r.get('rss_curve_n')}")
            lines.append(f"  客户数量 / 汇总金额: {r.get('customers')} / {r.get('total_sum',0):.2f} 元")
            lines.append(f"  结果正确：{r.get('correct')}（{r.get('detail','')}）")
            lines.append("")
            continue

        # 通用数值型
        lines.append(f"测试 {name}：")
        lines.append(f"  Excel 文件数量 : {r.get('n_files')}")
        lines.append(f"  总数据行数     : {r.get('total_rows')}")
        lines.append(f"  总输入大小     : {_mb(r.get('total_size') or 0)}")
        lines.append(f"  开始 RAM       : {_mb(r.get('start_rss') or 0)}（进程 RSS）")
        lines.append(f"  峰值 RAM       : {_mb(r.get('peak_rss') or 0)}（后台采样）")
        lines.append(f"  结束 RAM       : {_mb(r.get('end_rss') or 0)}（进程 RSS）")
        lines.append(f"  RAM 增量       : {_fmt_mb_opt(r.get('rss_delta'))}")
        lines.append(f"  CPU 时间       : {r.get('cpu_seconds',0):.2f} 秒")
        lines.append(f"  总耗时         : {r.get('elapsed',0):.2f} 秒（生成 {r.get('gen_seconds',0):.1f}s + 汇总）")
        lines.append(f"  是否崩溃       : {'否' if r.get('success') else '是（' + r.get('error','') + '）'}")
        lines.append(f"  是否读取异常   : {'否' if not r.get('skipped') else '是（跳过 ' + str(r.get('skipped')) + ' 文件）'}")
        lines.append(f"  是否生成输出   : {'是' if r.get('output_size') else '否'}")
        lines.append(f"  输出文件大小   : {_mb(r.get('output_size') or 0)}")
        lines.append(f"  客户数量       : {r.get('customers')}")
        lines.append(f"  汇总金额       : {r.get('total_sum',0):.2f} 元")
        lines.append(f"  结果正确       : {'是' if r.get('correct') else '否（' + r.get('detail','') + '）'}")
        lines.append("")
    lines.append("=" * 92)
    return "\n".join(lines)


# ------------------------------------------------------------------
#  测试组
# ------------------------------------------------------------------

def _setup_field_config():
    _ensure_dirs()
    save_field_config(CONFIG_DIR, FIELD_CONFIG)


def _teardown_field_config():
    reset_field_config(CONFIG_DIR)


def test_count(results):
    """T1-T4 文件数量压力。"""
    _setup_field_config()
    folder = os.path.join(INPUT_DIR, "count")
    shutil.rmtree(folder, ignore_errors=True)
    cases = [("T1", 1000), ("T2", 3000), ("T3", 5000), ("T4", 10000)]
    for name, n in cases:
        print(f"[{name}] 生成并汇总 {n} 个文件 ...")
        expected = {}
        r = run_case(name, folder, n, expected, kind="count")
        results.append(r)
        _save_results(results)
        st = "PASS" if (r["success"] and r["correct"]) else "FAIL"
        print(f"  -> 峰值 {_mb(r['peak_rss'])}, 耗时 {r['elapsed']:.1f}s, {st}"
              + ("" if r["correct"] else f"（{r['detail']}）"))
    _teardown_field_config()


def test_single(results):
    """T5-T7 单文件大数据。"""
    _setup_field_config()
    folder = os.path.join(INPUT_DIR, "single")
    shutil.rmtree(folder, ignore_errors=True)
    cases = [("T5", 10000), ("T6", 50000), ("T7", 100000)]
    for name, n_rows in cases:
        print(f"[{name}] 1 个文件 × {n_rows} 行 ...")
        expected = {}
        rows_list = [n_rows]
        r = run_case(name, folder, 1, expected, rows_list=rows_list, kind="single")
        results.append(r)
        _save_results(results)
        st = "PASS" if (r["success"] and r["correct"]) else "FAIL"
        print(f"  -> 峰值 {_mb(r['peak_rss'])}, 耗时 {r['elapsed']:.1f}s, {st}"
              + ("" if r["correct"] else f"（{r['detail']}）"))
    _teardown_field_config()


def test_mixed(results):
    """T8-T10 多文件 × 多行混合。"""
    _setup_field_config()
    folder = os.path.join(INPUT_DIR, "mixed")
    shutil.rmtree(folder, ignore_errors=True)

    def _build_rows(n_files, lo, hi, seed):
        rng = random.Random(seed)
        return [rng.randint(lo, hi) for _ in range(n_files)]

    cases = [
        ("T8", 100, 10000, 10000),
        ("T9", 500, 5000, 5000),
        ("T10", 1000, 1000, 5000),
    ]
    for name, n_files, lo, hi in cases:
        print(f"[{name}] {n_files} 个文件 × 每 {lo}~{hi} 行 ...")
        rows_list = _build_rows(n_files, lo, hi, 20260822 + n_files)
        expected = {}
        r = run_case(name, folder, n_files, expected, rows_list=rows_list, kind="mixed")
        results.append(r)
        _save_results(results)
        st = "PASS" if (r["success"] and r["correct"]) else "FAIL"
        print(f"  -> 峰值 {_mb(r['peak_rss'])}, 耗时 {r['elapsed']:.1f}s, {st}"
              + ("" if r["correct"] else f"（{r['detail']}）"))
    _teardown_field_config()


# ------------------------------------------------------------------
#  异常文件测试
# ------------------------------------------------------------------

def test_abnormal(results):
    """异常文件：空表 / 缺客户 / 缺金额 / 金额非数字 / 表头乱序 / 多余字段 /
    空白行 / 大量空白行 / .xlsx / .xlsm / .xls。单独报告，不并入上面的 PASS 统计。"""
    _setup_field_config()
    folder = os.path.join(ABNORMAL_DIR, "cases")
    shutil.rmtree(folder, ignore_errors=True)
    os.makedirs(folder, exist_ok=True)

    from openpyxl import Workbook

    # 1. 空 Excel（无表头）
    wb = Workbook(); wb.save(os.path.join(folder, "ab01_empty.xlsx")); wb.close()

    # 2. 缺客户字段（只有 金额/电话）
    _write_xlsx_fast(os.path.join(folder, "ab02_nocust.xlsx"), ["金额", "电话"], [[100, "1"]])
    # 3. 缺金额字段（只有 客户/电话）
    _write_xlsx_fast(os.path.join(folder, "ab03_nomoney.xlsx"), ["客户", "电话"], [["A", "1"]])
    # 4. 金额含非数字
    _write_xlsx_fast(os.path.join(folder, "ab04_badmoney.xlsx"), ["客户", "金额"],
                     [["A", "abc"], ["A", 100]])
    # 5. 表头顺序不同（金额在前 客户在后）—— 应能按字段名正确识别
    _write_xlsx_fast(os.path.join(folder, "ab05_reorder.xlsx"), ["金额", "电话", "客户"],
                     [[100, "1", "A"]])
    # 6. 多余字段
    _write_xlsx_fast(os.path.join(folder, "ab06_extra.xlsx"),
                     ["客户", "金额", "电话", "供应商", "备注", "日期"],
                     [["A", 100, "1", "S1", "x", "2026"]])
    # 7. 空白行
    wb = Workbook(); ws = wb.active
    ws.append(["客户", "金额", "电话", "供应商"])
    ws.append(["A", 100, "1", "S1"])
    ws.append([None, None, None, None])
    ws.append(["B", 200, "2", "S2"])
    wb.save(os.path.join(folder, "ab07_blankrow.xlsx")); wb.close()
    # 8. 大量空白行
    wb = Workbook(); ws = wb.active
    ws.append(["客户", "金额", "电话", "供应商"])
    ws.append(["A", 100, "1", "S1"])
    for _ in range(5000):
        ws.append([None, None, None, None])
    ws.append(["B", 200, "2", "S2"])
    wb.save(os.path.join(folder, "ab08_manyblank.xlsx")); wb.close()
    # 9. .xlsm
    _write_xlsx_fast(os.path.join(folder, "ab09.xlsm"), ["客户", "金额"], [["A", 100]])
    # 10. .xls
    _write_xls(os.path.join(folder, "ab10.xls"), ["客户", "金额"], [["A", 100]])

    # 逐个文件用 aggregate_with_report 处理，观察是否崩溃/卡死/静默丢数据
    cases_report = []
    for f in sorted(os.listdir(folder)):
        if not f.lower().endswith((".xlsx", ".xlsm", ".xls")):
            continue
        report = aggregate_with_report(folder, LOGS_DIR, [f], {})
        cases_report.append({
            "file": f,
            "processed": report["processed_count"],
            "skipped": report["skipped"],
            "customers": list(report["total"].keys()),
        })
        print(f"  [{f}] processed={report['processed_count']} "
              f"skipped={[s['file'] + ':' + s['reason'] for s in report['skipped']]}")

    # 组合：好文件 + 异常文件（核心字段无法识别）应阻止汇总（skipped 非空）
    good_folder = os.path.join(ABNORMAL_DIR, "good")
    os.makedirs(good_folder, exist_ok=True)
    _write_xlsx_fast(os.path.join(good_folder, "good.xlsx"), ["客户", "金额"], [["好公司", 100]])
    # 把 ab02 复制进 good_folder 与 good 混合
    shutil.copy2(os.path.join(folder, "ab02_nocust.xlsx"), os.path.join(good_folder, "bad.xlsx"))
    mixed_report = aggregate_with_report(good_folder, LOGS_DIR, None, {})
    blocked = len(mixed_report["skipped"]) > 0

    result = {
        "name": "ABNORMAL",
        "n_files": len(cases_report),
        "cases": cases_report,
        "mixed_skipped": [s["file"] for s in mixed_report["skipped"]],
        "mixed_blocked": blocked,
        "success": True,
        "correct": True,
        "detail": f"异常文件逐个处理无崩溃；混合场景核心字段缺失时阻止汇总={blocked}",
    }
    results.append(result)
    _save_results(results)
    _teardown_field_config()
    print(f"  [异常文件] 混合场景阻止汇总={blocked}，跳过文件={result['mixed_skipped']}")


# ------------------------------------------------------------------
#  字段映射场景测试
# ------------------------------------------------------------------

def test_mapping(results):
    """字段映射：临时映射不记住 / 下一任务不继承 / 核心字段无法识别阻止汇总。"""
    _setup_field_config()
    _orig_custom_map = dict(fm.CUSTOM_HEADER_MAP)

    folder = os.path.join(INPUT_DIR, "mapping")
    shutil.rmtree(folder, ignore_errors=True)
    os.makedirs(folder, exist_ok=True)

    # A: 正常表头；B: ccc|金额|电话（ccc → 客户）；D: xxx|yyy 无法识别
    _write_xlsx_fast(os.path.join(folder, "a.xlsx"), ["客户", "金额", "电话"],
                     [["公司A", 100, "1"], ["公司A", 50, "1"]])
    _write_xlsx_fast(os.path.join(folder, "b.xlsx"), ["ccc", "金额", "电话"],
                     [["公司B", 200, "2"]])
    _write_xlsx_fast(os.path.join(folder, "d.xlsx"), ["xxx", "yyy"], [["公司D", 1]])

    # 1. 临时映射（不记住）：b.xlsx 用 {"ccc":"客户"}，当前任务成功
    report = aggregate_with_report(folder, LOGS_DIR, None, {"b.xlsx": {"ccc": "客户"}})
    temp_ok = (report["processed_count"] == 2 and "公司B" in report["total"]
               and "公司A" in report["total"])

    # 2. 下一任务不继承：不带映射再跑，b.xlsx 应被跳过（ccc 无法识别）。
    #    注意 folder 里还有 d.xlsx（xxx/yyy 同样无法识别），故只断言 b.xlsx 在跳过列表中。
    report2 = aggregate_with_report(folder, LOGS_DIR, None, {})
    skipped_files2 = [s["file"] for s in report2["skipped"]]
    no_inherit = "b.xlsx" in skipped_files2
    # 且 CUSTOM_HEADER_MAP 未被污染
    no_pollution = "ccc" not in fm.CUSTOM_HEADER_MAP

    # 3. 无法识别文件阻止汇总：a + d 一起跑，d 无法识别 → skipped 非空
    report3 = aggregate_with_report(folder, LOGS_DIR, ["a.xlsx", "d.xlsx"], {})
    block_ok = (len(report3["skipped"]) == 1
                and report3["skipped"][0]["file"] == "d.xlsx"
                and "客户" in report3["skipped"][0]["missing"])

    result = {
        "name": "MAPPING",
        "n_files": 3,
        "temp_ok": temp_ok,
        "no_inherit": no_inherit,
        "no_pollution": no_pollution,
        "block_ok": block_ok,
        "success": True,
        "correct": temp_ok and no_inherit and no_pollution and block_ok,
        "detail": (f"临时映射={temp_ok} 不继承={no_inherit} "
                   f"不污染持久映射={no_pollution} 阻止汇总={block_ok}"),
    }
    results.append(result)
    _save_results(results)

    fm.CUSTOM_HEADER_MAP.clear()
    fm.CUSTOM_HEADER_MAP.update(_orig_custom_map)
    _teardown_field_config()
    print(f"  [字段映射] {result['detail']}")


# ------------------------------------------------------------------
#  内存泄漏测试（同一任务连续 10 次）
# ------------------------------------------------------------------

def test_leak(results):
    """同一压力任务连续运行 10 次，观察 RAM 是否持续增长。"""
    _setup_field_config()
    folder = os.path.join(INPUT_DIR, "leak")
    shutil.rmtree(folder, ignore_errors=True)
    os.makedirs(folder, exist_ok=True)

    n_files = 500
    expected = {}
    generate_files(folder, n_files, expected, kind="count", rng_seed=20260822)

    output_path = os.path.join(OUTPUT_DIR, "out_leak.xlsx")

    runs = []
    tracemalloc.start()
    snap_before = tracemalloc.take_snapshot()
    for i in range(1, 11):
        gc.collect()
        start_rss = _get_rss()
        sampler = RssSampler()
        sampler.start()
        t0 = time.time()
        report = aggregate_with_report(folder, LOGS_DIR, None, {})
        sorted_data = sort_by_amount(report["total"], report["field_names"])
        generate_summary(sorted_data, report["total"], output_path, report["field_names"])
        t1 = time.time()
        sampler.stop()
        gc.collect()
        end_rss = _get_rss()
        runs.append({
            "run": i,
            "start_rss": start_rss,
            "peak_rss": sampler.peak,
            "end_rss": end_rss,
            "elapsed": t1 - t0,
        })
        print(f"  第{i}次：开始 {_mb(start_rss)} / 峰值 {_mb(sampler.peak)} / "
              f"结束 {_mb(end_rss)} / {t1 - t0:.2f}s")
    snap_after = tracemalloc.take_snapshot()
    tracemalloc.stop()

    top_stats = snap_after.compare_to(snap_before, "lineno")[:5]
    top_lines = [f"{s.traceback[0]}" for s in top_stats]

    end_series = [r["end_rss"] for r in runs]
    peak_series = [r["peak_rss"] for r in runs]
    # 判定：结束 RSS 是否有持续上升趋势（简单线性：末-首 超过 8MB 且单调上升）
    trend_up = end_series[-1] - end_series[0]
    monotonic = all(end_series[i] <= end_series[i + 1] for i in range(len(end_series) - 1))
    suspected = trend_up > 8 * 1024 * 1024 and monotonic

    result = {
        "name": "LEAK",
        "n_files": n_files,
        "runs": runs,
        "end_series": end_series,
        "peak_series": peak_series,
        "suspected_leak": suspected,
        "trend_bytes": trend_up,
        "top_retained": top_lines,
        "success": True,
        "correct": not suspected,
        "detail": (f"连续10次结束RAM {'持续上升(疑似泄漏)' if suspected else '正常波动'}"
                   f"，末-首={_mb(trend_up)}"),
    }
    results.append(result)
    _save_results(results)
    _teardown_field_config()
    print(f"  [内存泄漏] {result['detail']}")


# ------------------------------------------------------------------
#  长时间连续运行测试
# ------------------------------------------------------------------

def test_longrun(results):
    """长时间运行：一次大任务，采样 RAM 曲线，验证完整输出。"""
    _setup_field_config()
    folder = os.path.join(INPUT_DIR, "longrun")
    shutil.rmtree(folder, ignore_errors=True)
    os.makedirs(folder, exist_ok=True)

    n_files = 2000
    rows_list = [random.Random(20260822 + i).randint(500, 2000) for i in range(n_files)]
    expected = {}
    generate_files(folder, n_files, expected, kind="mixed", rng_seed=20260822,
                   rows_list=rows_list, ext_mix=(100, 0, 0))

    output_path = os.path.join(OUTPUT_DIR, "out_longrun.xlsx")

    gc.collect()
    start_rss = _get_rss()
    sampler = RssSampler(interval=0.2)
    p0 = time.process_time()
    t_start = time.time()
    sampler.start()

    report = aggregate_with_report(folder, LOGS_DIR, None, {})
    sorted_data = sort_by_amount(report["total"], report["field_names"])
    generate_summary(sorted_data, report["total"], output_path, report["field_names"])

    t_end = time.time()
    p1 = time.process_time()
    sampler.stop()
    gc.collect()
    end_rss = _get_rss()

    # RAM 曲线（等间隔抽样，最多 200 点）
    samples = sampler.samples
    curve = []
    if len(samples) > 200:
        step = len(samples) // 200
        curve = [(t - t_start, rss) for i, (t, rss) in enumerate(samples) if i % step == 0]
    else:
        curve = [(t - t_start, rss) for t, rss in samples]

    correct, detail, extra = verify_output(
        output_path, expected, n_files, report["file_count"], report["field_names"],
        expected_phone=PHONE, expected_supplier=SUPPLIER,
    )

    result = {
        "name": "LONGRUN",
        "n_files": n_files,
        "total_rows": sum(rows_list),
        "start_time": t_start,
        "end_time": t_end,
        "elapsed": t_end - t_start,
        "cpu_seconds": p1 - p0,
        "start_rss": start_rss,
        "peak_rss": sampler.peak,
        "end_rss": end_rss,
        "rss_curve_n": len(curve),
        "rss_curve": curve,
        "success": True,
        "correct": correct,
        "detail": detail,
        "customers": len(expected),
        "total_sum": sum(expected.values()),
        "skipped": len(report["skipped"]),
    }
    results.append(result)
    _save_results(results)
    _teardown_field_config()
    st = "PASS" if correct else "FAIL"
    print(f"  [长时间运行] {n_files} 文件 / {sum(rows_list)} 行，"
          f"耗时 {result['elapsed']:.1f}s，峰值 {_mb(sampler.peak)}，{st}（{detail}）")


# ------------------------------------------------------------------
#  主入口
# ------------------------------------------------------------------

GROUPS = {
    "count": test_count,
    "single": test_single,
    "mixed": test_mixed,
    "abnormal": test_abnormal,
    "mapping": test_mapping,
    "leak": test_leak,
    "longrun": test_longrun,
}


def _dedupe(results):
    """按 name 去重，同名保留最后一次（重跑覆盖旧结果）。"""
    out = {}
    for r in results:
        out[r.get("name")] = r
    return list(out.values())


def main(argv=None):
    argv = argv if argv is not None else sys.argv[1:]
    args = [a for a in argv if not a.startswith("--")]
    tests = ",".join(args) if args else "all"

    _ensure_dirs()

    if tests == "all":
        selected = list(GROUPS.keys())
    else:
        selected = [t for t in tests.split(",") if t in GROUPS]

    if not selected:
        print(f"未知测试组：{tests}。可选：{', '.join(GROUPS)} / all")
        return 2

    print("=" * 92)
    print("Excel智能汇总助手 v1.5.4 — 正式压力测试")
    print("=" * 92)
    print(f"数据目录：{DATA_ROOT}")
    print(f"字段配置：{' / '.join(FIELD_NAMES)}")
    print(f"测试组：{', '.join(selected)}")
    print()

    results = _load_results()

    for name in selected:
        GROUPS[name](results)
        print()

    results = _dedupe(results)
    report = build_report(results)
    with open(REPORT_TXT, "w", encoding="utf-8") as f:
        f.write(report + "\n")
    _save_results(results)

    print(report)
    print(f"\n报告已写入：{REPORT_TXT}")
    print(f"结果 JSON：{RESULT_JSON}")
    print(f"数据保留在：{DATA_ROOT}（未删除）")

    all_pass = all(r.get("correct", True) for r in results)
    return 0 if all_pass else 1


if __name__ == "__main__":
    sys.exit(main())
