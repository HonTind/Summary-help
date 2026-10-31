"""Synthetic monthly date-filter acceptance tests. TemporaryDirectory only; never user data."""
import datetime
import hashlib
import importlib
import json
from pathlib import Path
import sys
import tempfile
import threading
import time
import tkinter as tk
import unittest
from unittest.mock import patch
from decimal import Decimal
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import MAC_EPOCH
import xlwt

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
PACKAGE = "Excel自动汇总助手"
analysis = importlib.import_module(PACKAGE + ".analysis")
monthly = importlib.import_module(PACKAGE + ".monthly")
export = importlib.import_module(PACKAGE + ".workbench_export")
ui = importlib.import_module(PACKAGE + ".workbench_gui")
Rules = analysis.SheetRules

EPOCH_1900 = datetime.date(1899, 12, 30)
EPOCH_1904 = datetime.date(1904, 1, 1)
SHEET_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def serial_1900(year, month, day):
    """Excel 1900-base serial for a modern (>= 1900-03-01) date."""
    return (datetime.date(year, month, day) - EPOCH_1900).days


def serial_1904(year, month, day):
    return (datetime.date(year, month, day) - EPOCH_1904).days


class DateFilterCoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ehelper-dates-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.rules = Rules("客户", "金额", None, 1, None, "日期")

    def xlsx(self, name, rows, *, headers=("客户", "日期", "金额"), sheet="账单", epoch=None):
        path = self.root / name
        book = Workbook()
        if epoch is not None:
            book.epoch = epoch
        ws = book.active
        ws.title = sheet
        ws.append(headers)
        for row in rows:
            ws.append(row)
        book.save(path)
        book.close()
        return path

    def xlsx_serials(self, name, serials, *, headers=("客户", "日期", "金额"), epoch=None,
                     date_format=True):
        """Real workbook whose date column holds raw Excel numbers (optionally date-styled)."""
        path = self.root / name
        book = Workbook()
        if epoch is not None:
            book.epoch = epoch
        ws = book.active
        ws.title = "账单"
        ws.append(headers)
        for index, serial in enumerate(serials, 2):
            ws.append([f"K{index}", serial, "1"])
            if date_format:
                ws.cell(row=index, column=2).number_format = "yyyy-mm-dd"
        book.save(path)
        book.close()
        return path

    def xls(self, name, rows, *, headers=("客户", "日期", "金额"), sheet="账单", dates_1904=False):
        path = self.root / name
        book = xlwt.Workbook()
        if dates_1904:
            book.dates_1904 = True
        ws = book.add_sheet(sheet)
        for j, header in enumerate(headers):
            ws.write(0, j, header)
        date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD")
        for i, row in enumerate(rows, 1):
            for j, value in enumerate(row):
                if value is None:
                    continue
                if isinstance(value, datetime.datetime):
                    ws.write(i, j, value, date_style)
                else:
                    ws.write(i, j, value)
        book.save(str(path))
        return path

    def xls_serials(self, name, serials, *, dates_1904=False, date_format=True):
        path = self.root / name
        book = xlwt.Workbook()
        if dates_1904:
            book.dates_1904 = True
        ws = book.add_sheet("账单")
        for j, header in enumerate(("客户", "日期", "金额")):
            ws.write(0, j, header)
        style = xlwt.easyxf(num_format_str="YYYY-MM-DD") if date_format else None
        for index, serial in enumerate(serials, 1):
            ws.write(index, 0, f"K{index}")
            if style is None:
                ws.write(index, 1, float(serial))              # Genuinely unformatted number cell.
            else:
                ws.write(index, 1, float(serial), style)
            ws.write(index, 2, "1")
        book.save(str(path))
        return path

    def rewrite_xml(self, original, name, transform):
        """Copy an xlsx with sheet1.xml edited, e.g. to forge a cached formula value."""
        target = self.root / name
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(target, "w") as output:
            for entry in source.infolist():
                content = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    tree = ET.fromstring(content)
                    transform(tree)
                    content = ET.tostring(tree, encoding="utf-8")
                output.writestr(entry, content)
        return target

    def read(self, path, month, rules=None):
        return analysis.read_inputs([path], rules or self.rules, month=month)

    # --- accepted input kinds -------------------------------------------------
    def test_text_real_and_serial_dates_in_month(self):
        serial = serial_1900(2026, 10, 8)
        path = self.xlsx("kinds.xlsx", [
            ("A", datetime.datetime(2026, 10, 5, 9, 30), "1"),
            ("B", "2026/10/06", "2"),
            ("C", "20261007", "3"),
            ("D", serial, "4"),
            ("E", datetime.datetime(2026, 10, 31, 23, 59), "5"),
            ("F", "2026-11-01", "6"),
        ])
        before = hashlib.sha256(path.read_bytes()).digest()
        report = self.read(path, "2026-10")
        self.assertTrue(report.complete, report.issues)
        self.assertEqual([r.parts for r in report.records], [("A",), ("B",), ("C",), ("D",), ("E",)])
        self.assertEqual(report.excluded_rows, 1)
        self.assertEqual(report.excluded[0].normalized, "2026-11-01")
        self.assertEqual(report.scanned_rows, 6)
        self.assertEqual(report.invalid_rows, 0)
        self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), before)

    def test_three_formats_and_1900_1904_bases(self):
        for extension in ("xlsx", "xlsm", "xls"):
            with self.subTest(extension=extension):
                rows = [("A", "2026-10-05", "1"), ("B", "2026-09-30", "2"), ("C", "2026-10-06", "3")]
                path = (self.xls if extension == "xls" else self.xlsx)("book." + extension, rows)
                report = self.read(path, "2026-10")
                self.assertTrue(report.complete, report.issues)
                self.assertEqual([r.parts for r in report.records], [("A",), ("C",)])
                self.assertEqual(report.excluded_rows, 1)
        # 1904 base, numeric serials, both containers.
        value = serial_1904(2026, 10, 5)
        for name, path in (("b1904.xlsx", self.xlsx("b1904.xlsx", [("A", value, "1")], epoch=MAC_EPOCH)),
                           ("b1904.xls", self.xls("b1904.xls", [("A", value, "1")], dates_1904=True))):
            with self.subTest(name=name):
                report = self.read(path, "2026-10")
                self.assertTrue(report.complete, report.issues)
                self.assertEqual([r.parts for r in report.records], [("A",)])

    # --- range boundaries -----------------------------------------------------
    def test_month_boundaries_cross_year_and_leap_february(self):
        path = self.xlsx("edges.xlsx", [
            ("first", "2026-12-01", "1"),
            ("last", "2026-12-31", "2"),
            ("next", "2027-01-01", "3"),
            ("prev", "2026-11-30", "4"),
        ])
        report = self.read(path, "2026-12")
        self.assertEqual([r.parts for r in report.records], [("first",), ("last",)])
        self.assertEqual(report.excluded_rows, 2)
        leap = self.xlsx("leap.xlsx", [("f29", "2024-02-29", "1"), ("mar", "2024-03-01", "2")])
        report = self.read(leap, "2024-02")
        self.assertEqual([r.parts for r in report.records], [("f29",)])
        self.assertEqual(report.excluded_rows, 1)

    # --- anomalies vs legal out-of-month -------------------------------------
    def test_out_of_month_bad_amount_is_skipped_not_an_anomaly(self):
        path = self.xlsx("skip.xlsx", [
            ("A", "2026-10-05", "bad"),   # in month, bad amount -> anomaly
            ("B", "2026-09-30", "NaN"),   # out of month with bad amount -> legal skip
            ("C", "2026-10-06", "1"),     # in month, valid
        ])
        report = self.read(path, "2026-10")
        self.assertEqual([r.parts for r in report.records], [("C",)])
        self.assertEqual(report.invalid_rows, 1)
        self.assertEqual(report.excluded_rows, 1)
        self.assertEqual([issue.code for issue in report.issues], ["invalid_row"])
        self.assertFalse(report.complete)

    def test_all_out_of_month_is_not_complete_and_not_anomaly(self):
        path = self.xlsx("outside.xlsx", [("A", "2026-09-01", "1"), ("B", "2026-11-30", "2")])
        report = self.read(path, "2026-10")
        self.assertEqual(report.records, [])
        self.assertEqual(report.excluded_rows, 2)
        self.assertEqual(report.issues, [])   # legal skips, not anomalies
        self.assertFalse(report.complete)

    def test_missing_and_bad_dates_are_anomalies(self):
        path = self.xlsx("bad-dates.xlsx", [
            ("missing", None, "1"),
            ("ambiguous", "01/02/2026", "2"),
            ("month13", "2026-13-01", "3"),
            ("day30feb", "2026-02-30", "4"),
            ("boolean", True, "5"),
            ("infinity", float("inf"), "6"),
            ("ratio", float("nan"), "7"),
            ("junk", "not-a-date", "8"),
        ])
        report = self.read(path, "2026-10")
        self.assertEqual(report.records, [])
        self.assertEqual(report.invalid_rows, 8)
        self.assertEqual(sum(issue.code == "invalid_date" for issue in report.issues), 8)
        self.assertFalse(report.complete)

    def test_error_and_uncached_formula_date_are_anomalies(self):
        path = self.xlsx("formula-date.xlsx", [("A", "=SUM(1,2)", "1")])
        report = self.read(path, "2026-10")
        self.assertEqual(report.invalid_rows, 1)
        self.assertIn("公式没有缓存", report.issues[0].message)
        error_path = self.xlsx("error-date.xlsx", [("A", None, "1")])
        book = load_workbook(error_path)
        book.active["B2"] = "#N/A"
        book.active["B2"].data_type = "e"
        book.save(error_path)
        book.close()
        report = self.read(error_path, "2026-10")
        self.assertEqual(report.invalid_rows, 1)
        self.assertIn("Excel 错误", report.issues[0].message)

    # --- row accounting -------------------------------------------------------
    def test_mutually_exclusive_row_counts(self):
        path = self.xlsx("counts.xlsx", [
            ("A", "2026-10-05", "1"),   # valid in month
            (None, None, None),         # blank
            ("B", "2026-10-06", "bad"), # in month, invalid amount
            ("C", "2026-10-07", "2"),   # valid in month
            ("D", "2026-09-30", "3"),   # legal out of month
        ])
        report = self.read(path, "2026-10")
        self.assertEqual(report.scanned_rows,
                         len(report.records) + report.blank_rows + report.invalid_rows + report.excluded_rows)
        self.assertEqual((len(report.records), report.blank_rows, report.invalid_rows, report.excluded_rows),
                         (2, 1, 1, 1))

    # --- 1900 fictitious leap day / serial zero -------------------------------
    def test_serial_boundaries_1900_and_1904(self):
        # 1900 base: 1..59 are real dates before the fake 1900-02-29; 60 is the fake day.
        self.assertEqual(analysis.parse_date_value(1, "", datemode=0), datetime.date(1900, 1, 1))
        self.assertEqual(analysis.parse_date_value(59, "", datemode=0), datetime.date(1900, 2, 28))
        self.assertEqual(analysis.parse_date_value(59.5, "", datemode=0), datetime.date(1900, 2, 28))
        self.assertEqual(analysis.parse_date_value(61, "", datemode=0), datetime.date(1900, 3, 1))
        self.assertEqual(analysis.parse_date_value(serial_1900(2026, 10, 5), "", datemode=0),
                         datetime.date(2026, 10, 5))
        for bad in (0, -1, 60, 60.5):
            with self.assertRaises(ValueError):
                analysis.parse_date_value(bad, "", datemode=0)
        # 1904 base: serial 0 is the real 1904-01-01; a time fraction keeps that same day.
        self.assertEqual(analysis.parse_date_value(0, "", datemode=1), datetime.date(1904, 1, 1))
        self.assertEqual(analysis.parse_date_value(0.5, "", datemode=1), datetime.date(1904, 1, 1))
        self.assertEqual(analysis.parse_date_value(59, "", datemode=1), datetime.date(1904, 2, 29))
        self.assertEqual(analysis.parse_date_value(60, "", datemode=1), datetime.date(1904, 3, 1))
        self.assertEqual(analysis.parse_date_value(61, "", datemode=1), datetime.date(1904, 3, 2))
        self.assertEqual(analysis.parse_date_value(serial_1904(2026, 10, 5), "", datemode=1),
                         datetime.date(2026, 10, 5))
        with self.assertRaises(ValueError):
            analysis.parse_date_value(-1, "", datemode=1)
        for non_finite in (float("inf"), float("-inf"), float("nan")):
            with self.assertRaises(ValueError):
                analysis.parse_date_value(non_finite, "", datemode=0)

    def test_date_serials_workbook_both_bases_and_formats(self):
        serials = [0, 1, 59, 60, 60.5, 61]
        for label, build in (("xlsx", lambda: self.xlsx_serials("b1900s.xlsx", serials)),
                             ("xls", lambda: self.xls_serials("b1900s.xls", serials))):
            with self.subTest(base="1900", kind=label):
                report = self.read(build(), "1900-01")
                self.assertEqual([r.normalized_date for r in report.records], ["1900-01-01"])
                self.assertEqual(report.invalid_rows, 3)   # 0, 60 and 60.5
                self.assertEqual(report.excluded_rows, 2)  # serials 59 and 61 are legal but out of month
                self.assertTrue(all(issue.code == "invalid_date" for issue in report.issues))
        for label, build in (("xlsx", lambda: self.xlsx_serials("b1904s.xlsx", serials, epoch=MAC_EPOCH)),
                             ("xls", lambda: self.xls_serials("b1904s.xls", serials, dates_1904=True))):
            with self.subTest(base="1904", kind=label):
                report = self.read(build(), "1904-01")
                self.assertEqual({r.normalized_date for r in report.records},
                                 {"1904-01-01", "1904-01-02"})   # serial 0 and 1
                self.assertEqual(report.invalid_rows, 0)
                self.assertEqual(report.excluded_rows, 4)       # serials 59, 60, 60.5 and 61

    def test_unformatted_numeric_serials_follow_workbook_base(self):
        serials = [0, 60, 61]
        for label, build in (("xlsx", lambda: self.xlsx_serials("u1900.xlsx", serials, date_format=False)),
                             ("xls", lambda: self.xls_serials("u1900.xls", serials, date_format=False))):
            with self.subTest(kind=label):
                report = self.read(build(), "1900-03")
                self.assertEqual([r.normalized_date for r in report.records], ["1900-03-01"])
                self.assertEqual(report.invalid_rows, 2)   # 0 and the fictitious 60
                self.assertEqual(report.excluded_rows, 0)

    def test_fictitious_serial_60_workbook_rejected_all_formats(self):
        # openpyxl decodes serial 60 to 1900-02-28; the raw workbook number must win.
        for label, build in (("xlsx", lambda: self.xlsx_serials("s60.xlsx", [60])),
                             ("xlsm", lambda: self.xlsx_serials("s60.xlsm", [60])),
                             ("xls", lambda: self.xls_serials("s60.xls", [60]))):
            with self.subTest(kind=label):
                report = self.read(build(), "1900-02")
                self.assertEqual(report.records, [])
                self.assertEqual(report.invalid_rows, 1)
                self.assertEqual(report.issues[0].code, "invalid_date")

    def test_real_1900_02_28_serial_59_not_rejected(self):
        for label, build in (("xlsx", lambda: self.xlsx_serials("r59.xlsx", [59])),
                             ("xls", lambda: self.xls_serials("r59.xls", [59]))):
            with self.subTest(kind=label):
                report = self.read(build(), "1900-02")
                self.assertEqual([r.normalized_date for r in report.records], ["1900-02-28"])
                self.assertEqual(report.invalid_rows, 0)

    def test_fictitious_serial_60_in_cached_formula_rejected(self):
        base = self.xlsx_serials("formula60-base.xlsx", [59])   # B2 already carries a date style

        def transform(tree):
            cell = tree.find(".//" + SHEET_NS + "c[@r='B2']")
            for child in list(cell):
                if child.tag == SHEET_NS + "v":
                    child.text = "60"                            # fictitious cached value
            formula = ET.Element(SHEET_NS + "f")
            formula.text = "1+58"
            cell.insert(0, formula)

        path = self.rewrite_xml(base, "formula60.xlsx", transform)
        report = self.read(path, "1900-02")
        self.assertEqual(report.records, [])
        self.assertEqual(report.invalid_rows, 1)
        self.assertEqual(report.issues[0].code, "invalid_date")
        self.assertEqual(report.issues[0].raw_date, "=1+58")     # formula text kept intact

    # --- configuration errors and safety --------------------------------------
    def test_configuration_errors_are_explicit(self):
        with self.assertRaises(ValueError):
            Rules("客户", "金额", None, 1, None, 0)          # invalid date selection
        with self.assertRaises(ValueError):
            Rules("客户", "金额", None, 1, None, True)       # bool is not a selection
        path = self.xlsx("overlap.xlsx", [("A", "2026-10-05", "1")])
        for selection in ("客户", "金额"):
            report = self.read(path, "2026-10", Rules("客户", "金额", None, 1, None, selection))
            self.assertTrue(report.issues)
            self.assertEqual(report.issues[0].code, "file_error")
            self.assertIn("日期字段", report.issues[0].message)
        with self.assertRaises(ValueError):
            analysis.read_inputs([path], self.rules)  # month is required when filtering

    def test_cancel_before_read_with_filter(self):
        event = threading.Event()
        event.set()
        path = self.xlsx("cancel.xlsx", [("A", "2026-10-05", "1")])
        with self.assertRaises(analysis.CancelledError):
            analysis.read_inputs([path], self.rules, month="2026-10", cancel=event)

    def test_zero_negative_and_high_precision_amounts(self):
        path = self.xlsx("amounts.xlsx", [
            ("zero", "2026-10-01", 0),
            ("neg", "2026-10-02", "-1.001"),
            ("prec", "2026-10-03", "123456789012345678901234567890.001"),
            ("out", "2026-09-30", "999"),
        ])
        result = monthly.run_monthly([path], monthly.make_template("金额", self.rules), "2026-10")
        rows = {row.parts: row.amount for row in result.rows}
        self.assertEqual(rows[("zero",)], Decimal(0))
        self.assertEqual(rows[("neg",)], Decimal("-1.001"))
        self.assertEqual(rows[("prec",)], Decimal("123456789012345678901234567890.001"))
        self.assertNotIn(("out",), rows)
        self.assertEqual(result.report.excluded_rows, 1)

    # --- date traceability (round 2) ------------------------------------------
    def test_valid_records_and_export_keep_date_trace(self):
        serial = serial_1900(2026, 10, 5)
        path = self.xlsx("trace.xlsx", [
            ("A", datetime.datetime(2026, 10, 5, 9, 30), "1"),
            ("B", "2026/10/06", "2"),
            ("C", serial, "3"),
            ("D", "2026-09-30", "4"),
        ])
        result = monthly.run_monthly([path], monthly.make_template("溯源", self.rules), "2026-10")
        by_key = {record.key: record for record in result.report.records}
        self.assertEqual(by_key["A"].normalized_date, "2026-10-05")
        self.assertIn("2026-10-05", by_key["A"].raw_date)          # datetime source text kept
        self.assertEqual(by_key["B"].normalized_date, "2026-10-06")
        self.assertEqual(by_key["C"].normalized_date, "2026-10-05")
        output = export.export_result(result, self.root, prefix="溯源")
        book = load_workbook(output)
        self.addCleanup(book.close)
        details = book["来源明细"]
        traced = {row[5]: row[7] for row in details.values if row[1] == "月度汇总"}
        self.assertIn("2026-10-05", traced["A"])
        self.assertIn("2026/10/06", traced["B"])                   # original kept
        self.assertIn("2026-10-06", traced["B"])                   # normalised kept
        excluded = [row for row in details.values if row[1] == "日期排除"]
        self.assertEqual(len(excluded), 1)
        self.assertEqual(excluded[0][5], "D")                      # genuine key, not the date
        self.assertEqual(excluded[0][6], "未校验")                 # no fabricated zero
        self.assertEqual(excluded[0][7], "2026-09-30")

    def test_invalid_rows_keep_verified_date_and_real_key(self):
        path = self.xlsx("badamount.xlsx", [
            ("A", "2026-10-05", "bad"),   # valid date, bad amount
            ("B", "not-a-date", "1"),     # bad date
        ])
        report = self.read(path, "2026-10")
        by_code = {issue.code: issue for issue in report.issues}
        self.assertEqual(by_code["invalid_row"].normalized_date, "2026-10-05")
        self.assertEqual(by_code["invalid_row"].raw_date, "2026-10-05")
        self.assertEqual(by_code["invalid_row"].raw_key, "A")
        self.assertEqual(by_code["invalid_date"].raw_key, "B")     # real key, not the date text
        self.assertEqual(by_code["invalid_date"].raw_date, "not-a-date")
        self.assertEqual(by_code["invalid_date"].normalized_date, "")

    def test_export_issue_columns_keep_dates(self):
        path = self.xlsx("exp-issues.xlsx", [("A", "2026-10-05", "bad"), ("B", "not-a-date", "1")])
        result = monthly.run_monthly([path], monthly.make_template("异常", self.rules), "2026-10")
        output = export.export_result(result, self.root, prefix="异常")
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book["异常清单"]["A1"].value, "侧别")     # leading columns unchanged
        self.assertEqual(book["异常清单"]["G1"].value, "原始键")
        rows = {row[6]: row for row in book["异常清单"].values}
        self.assertEqual(rows["A"][8], "2026-10-05")
        self.assertEqual(rows["A"][9], "2026-10-05")
        self.assertEqual(rows["B"][6], "B")
        self.assertEqual(rows["B"][8], "not-a-date")
        self.assertIn(rows["B"][9], (None, ""), "空规范日期单元格读回应为空")

    # --- original-number traceability (round 3) --------------------------------
    def test_fictitious_serials_keep_original_number_in_issues(self):
        # Date-styled serials 60/60.5 are anomalies; the issue keeps the workbook number,
        # not openpyxl's silently-corrected 1900-02-28 display (xls keeps the raw number too).
        serials = [60, 60.5]
        for label, build in (("xlsx", lambda: self.xlsx_serials("s60trace.xlsx", serials)),
                             ("xlsm", lambda: self.xlsx_serials("s60trace.xlsm", serials)),
                             ("xls", lambda: self.xls_serials("s60trace.xls", serials))):
            with self.subTest(kind=label):
                report = self.read(build(), "1900-02")
                self.assertEqual(report.records, [])
                self.assertEqual(report.invalid_rows, 2)
                invalid = [issue for issue in report.issues if issue.code == "invalid_date"]
                self.assertEqual(len(invalid), 2)
                for issue, number in zip(invalid, ("60", "60.5")):
                    self.assertEqual(issue.normalized_date, "")
                    # The original number leads; the corrected display never does.
                    self.assertTrue(issue.raw_date.startswith(number),
                                    f"{label} {number}: {issue.raw_date!r}")

    def test_1904_serial_zero_record_and_export_keep_original_number(self):
        # 1904 serial 0 is the real 1904-01-01; openpyxl shows only 00:00:00, so the raw
        # number must stay traceable in the record and in the exported date column.
        for label, name in (("xlsx", "b1904zero.xlsx"), ("xlsm", "b1904zero.xlsm")):
            with self.subTest(kind=label):
                result = monthly.run_monthly([self.xlsx_serials(name, [0], epoch=MAC_EPOCH)],
                                             monthly.make_template("零", self.rules), "1904-01")
                self.assertEqual(len(result.report.records), 1)
                record = result.report.records[0]
                self.assertEqual(record.normalized_date, "1904-01-01")
                self.assertEqual(record.raw_date, "0")             # 原序列 0，不是 00:00:00
                self.assertEqual(result.report.invalid_rows, 0)
                output = export.export_result(result, self.root, prefix="零" + label)
                book = load_workbook(output)
                self.addCleanup(book.close)
                row = next(r for r in book["来源明细"].values if r[5] == record.key)
                self.assertEqual(row[7], "原 0｜规范 1904-01-01")

    # --- templates ------------------------------------------------------------
    def test_schema1_read_execute_default_off_then_new_format_roundtrip(self):
        path = self.xlsx("tpl.xlsx", [("A", "2026-10-05", "1"), ("B", "2026-09-30", "2")])
        raw = self.root / "legacy.json"
        raw.write_text(json.dumps({"schema_version": 1, "templates": [{
            "schema_version": 1, "name": "旧", "output_prefix": "月度汇总",
            "rules": {"key_column": "客户", "amount_column": "金额", "sheet": None, "header_row": 1},
        }]}, ensure_ascii=False), encoding="utf-8")
        before = raw.read_bytes()
        template = monthly.load_templates(raw)["旧"]
        # Default off: both months aggregate, results unchanged.
        result = monthly.run_monthly([path], template, "2026-10")
        self.assertEqual(result.total, Decimal("3"))
        self.assertEqual(result.report.excluded_rows, 0)
        self.assertEqual(raw.read_bytes(), before)  # Reading never rewrites the file.
        # New format carries the date rule and filters after reload.
        new = monthly.make_template("旧", self.rules)
        monthly.save_template(raw, new, allow_update=True)
        reloaded = monthly.load_templates(raw)["旧"]
        self.assertEqual(reloaded["rules"]["date_column"], "日期")
        filtered = monthly.run_monthly([path], reloaded, "2026-10")
        self.assertEqual(filtered.total, Decimal("1"))
        self.assertEqual(filtered.report.excluded_rows, 1)

    # --- export traceability --------------------------------------------------
    def test_export_records_date_rule_and_excluded_rows(self):
        path = self.xlsx("exp.xlsx", [("A", "2026-10-05", "1"), ("B", "2026-09-30", "2")])
        before = hashlib.sha256(path.read_bytes()).digest()
        result = monthly.run_monthly([path], monthly.make_template("导出", self.rules), "2026-10")
        output = export.export_result(result, self.root, prefix="日期导出")
        self.assertNotEqual(output, str(path))
        self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), before)
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["月度汇总", "来源明细", "异常清单", "运行配置"])
        config = dict(book["运行配置"].values)
        self.assertEqual(config["日期筛选"], "开启")
        self.assertEqual(config["日期字段"], "日期")
        self.assertEqual(config["外月跳过行数"], 1)
        self.assertIn("2026-10", config["日期范围"])
        details = book["来源明细"]
        excluded = [row for row in details.values if row[1] == "日期排除"]
        self.assertEqual(len(excluded), 1)
        self.assertEqual(excluded[0][5], "B")                      # genuine key
        self.assertEqual(excluded[0][6], "未校验")                 # no fabricated zero
        self.assertEqual(excluded[0][7], "2026-09-30")             # explicit date column


class DateFilterGuiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ehelper-dates-ui-")
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.root = tk.Tk()
        self.root.withdraw()
        self.tk_errors = []
        self.root.report_callback_exception = lambda *args: self.tk_errors.append(args)
        self.window = ui.WorkbenchWindow(self.root, self.folder / "config" / "monthly_templates.json")
        self.window.withdraw()
        self.root.update()
        self.addCleanup(self.cleanup_window)
        for name, value in (("showinfo", None), ("showerror", None), ("showwarning", None), ("askyesno", True)):
            mock = patch.object(ui.messagebox, name, return_value=value)
            mock.start()
            self.addCleanup(mock.stop)

    def cleanup_window(self):
        self.window.runner.cancel()
        if self.window.runner.thread:
            self.window.runner.thread.join(2)
        try:
            self.root.destroy()
        except tk.TclError:
            pass
        self.assertFalse(self.window.runner.busy, "测试线程未释放")
        self.assertFalse(self.tk_errors, self.tk_errors)

    def make_file(self, name, rows, headers=("客户", "日期", "金额")):
        path = self.folder / name
        book = Workbook()
        sheet = book.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        book.save(path)
        book.close()
        return path

    def drain(self):
        deadline = time.monotonic() + 5
        while self.window._busy and not self.window._closed and time.monotonic() < deadline:
            self.root.update()
            time.sleep(0.005)
        self.root.update()
        self.assertFalse(self.window._busy, "工作台未结束")
        self.assertFalse(self.tk_errors)

    def enable_filter(self, field="日期"):
        self.window.date_filter.set(True)
        self.window._date_toggle()
        self.window.date_field.set(field)

    def test_switch_requires_field_and_invalidates(self):
        self.assertEqual(str(self.window.date_entry.cget("state")), "disabled")
        self.enable_filter()
        self.assertEqual(str(self.window.date_entry.cget("state")), "normal")
        self.window.monthly_input.key.set("客户")
        self.window.monthly_input.amount.set("金额")
        rules = self.window.monthly_rules()
        self.assertEqual(rules.date_column, "日期")
        self.window.date_field.set("")
        with self.assertRaises(ValueError):
            self.window.monthly_rules()

    def test_preview_filters_reports_skip_and_invalidates(self):
        path = self.make_file("m.xlsx", [("A", "2026-10-05", "1"), ("B", "2026-09-30", "2"),
                                         ("C", "2026-10-06", "3")])
        self.window.monthly_input.set_files([path])
        self.enable_filter()
        self.window.month.set("2026-10")
        self.window.preview()
        self.drain()
        self.assertEqual({row.parts for row in self.window.result.rows}, {("A",), ("C",)})
        self.assertEqual(self.window.result.report.excluded_rows, 1)
        self.assertIn("外月跳过", self.window.status.get())
        self.assertEqual(len(self.window.sources.rows), 3)  # two valid rows plus the skip
        self.window.date_field.set("金额")
        self.assertIsNone(self.window.result)
        self.assertEqual(str(self.window.export_button.cget("state")), "disabled")

    def test_run_locks_controls_and_freezes_date_snapshot(self):
        path = self.make_file("snap.xlsx", [("A", "2026-10-05", "1"), ("B", "2026-09-30", "2")])
        self.window.monthly_input.set_files([path])
        self.enable_filter()
        self.window.month.set("2026-10")
        release = threading.Event()
        original = ui.run_monthly

        def delayed(paths, template, month, **kwargs):
            release.wait(2)
            return original(paths, template, month, **kwargs)

        with patch.object(ui, "run_monthly", side_effect=delayed):
            self.window.preview()
            self.assertEqual(str(self.window.date_check.cget("state")), "disabled")
            self.assertEqual(str(self.window.date_entry.cget("state")), "disabled")
            self.window.date_field.set("金额")  # Must not leak into the running snapshot.
            release.set()
            self.drain()
        self.assertEqual(self.window.result.template["rules"]["date_column"], "日期")
        self.assertEqual({row.parts for row in self.window.result.rows}, {("A",)})

    def test_template_roundtrip_and_export_with_date(self):
        path = self.make_file("tpl.xlsx", [("A", "2026-10-05", "1"), ("B", "2026-09-30", "2")])
        self.window.monthly_input.set_files([path])
        self.window.monthly_input.key.set("客户")
        self.window.monthly_input.amount.set("金额")
        self.enable_filter()
        self.window.template_name.set("日期模板")
        self.window.save_template()
        self.assertIn("date_column", monthly.load_templates(self.window.template_path)["日期模板"]["rules"])
        self.window.date_filter.set(False)
        self.window._date_toggle()
        self.window.date_field.set("")
        self.window.load_template()
        self.assertTrue(self.window.date_filter.get())
        self.assertEqual(self.window.date_field.get(), "日期")
        self.window.month.set("2026-10")
        self.window.preview()
        self.drain()
        with patch.object(ui.filedialog, "askdirectory", return_value=str(self.folder)):
            self.window.export()
            self.drain()
        output = next(self.folder.glob("月度汇总-2026-10*.xlsx"))
        book = load_workbook(output)
        self.addCleanup(book.close)
        config = dict(book["运行配置"].values)
        self.assertEqual(config["日期筛选"], "开启")
        self.assertEqual(config["外月跳过行数"], 1)

    def test_filter_invalidate_then_restore_full_month(self):
        path = self.make_file("mix.xlsx", [("A", "2026-10-05", "1"), ("B", "2026-09-30", "2"),
                                           ("C", "2026-10-06", "3")])
        self.window.monthly_input.set_files([path])
        self.enable_filter()
        self.window.month.set("2026-10")
        self.window.preview()
        self.drain()
        self.assertEqual({row.parts for row in self.window.result.rows}, {("A",), ("C",)})
        self.window.date_filter.set(False)
        self.window._date_toggle()
        self.assertIsNone(self.window.result)
        self.assertEqual(str(self.window.export_button.cget("state")), "disabled")
        self.window.preview()
        self.drain()
        self.assertEqual({row.parts for row in self.window.result.rows}, {("A",), ("B",), ("C",)})
        self.assertEqual(self.window.result.report.excluded_rows, 0)

    def assert_date_controls_reachable(self, geometry):
        """Real Tk layout evidence: the date controls and footer buttons stay on-screen."""
        self.root.deiconify()
        self.window.deiconify()
        self.window.geometry(geometry)
        self.window.update_idletasks()
        self.window.update()
        try:
            width, height = self.window.winfo_width(), self.window.winfo_height()
            widgets = {
                "按日期筛选": self.window.date_check,
                "日期字段": self.window.date_entry,
                "预览结果": self.window.preview_button,
                "导出完整 Excel": self.window.export_button,
            }
            for name, widget in widgets.items():
                x = widget.winfo_rootx() - self.window.winfo_rootx()
                y = widget.winfo_rooty() - self.window.winfo_rooty()
                w, h = widget.winfo_width(), widget.winfo_height()
                bounds = (x, y, w, h)
                self.assertTrue(widget.winfo_ismapped(), f"{geometry} {name} 未映射: {bounds}")
                self.assertGreaterEqual(x, 0, f"{geometry} {name} 越出左边界: {bounds}")
                self.assertGreaterEqual(y, 0, f"{geometry} {name} 越出上边界: {bounds}")
                self.assertLessEqual(x + w, width, f"{geometry} {name} 越出右边界: {bounds}")
                self.assertLessEqual(y + h, height, f"{geometry} {name} 越出下边界: {bounds}")
        finally:
            self.window.withdraw()
            self.root.withdraw()
            self.window.update_idletasks()

    def test_date_controls_reachable_at_min_and_default_size(self):
        self.window.tabs.select(0)
        self.root.update()
        for geometry in ("840x580", "1040x720"):
            with self.subTest(geometry=geometry):
                self.assert_date_controls_reachable(geometry)


if __name__ == "__main__":
    unittest.main()
