"""Synthetic composite-key acceptance tests. Real workbooks only; never user data."""
import hashlib
import importlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
import datetime
from decimal import Decimal
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
import xlwt

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
PACKAGE = "Excel自动汇总助手"
analysis = importlib.import_module(PACKAGE + ".analysis")
monthly = importlib.import_module(PACKAGE + ".monthly")
reconciliation = importlib.import_module(PACKAGE + ".reconciliation")
export = importlib.import_module(PACKAGE + ".workbench_export")
Rules = analysis.SheetRules


class CompositeWorkbenchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ehelper-composite-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def workbook(self, name, rows, headers=("客户", "单据号", "金额"), sheet="账单"):
        path = self.root / name
        if path.suffix == ".xls":
            book = xlwt.Workbook()
            ws = book.add_sheet(sheet)
            for i, row in enumerate([headers, *rows]):
                for j, value in enumerate(row):
                    if value is not None:
                        ws.write(i, j, value)
            book.save(str(path))
        else:
            book = Workbook()
            ws = book.active
            ws.title = sheet
            ws.append(headers)
            for row in rows:
                ws.append(row)
            book.save(path)
            book.close()
        return path

    def read(self, path, rules):
        return analysis.read_inputs([path], rules)

    def rewrite_xml(self, original, name, transform):
        """Return a copy of an xlsx with sheet1.xml edited, to forge cached values."""
        target = self.root / name
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(target, "w") as output:
            for entry in source.infolist():
                content = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    tree = ET.fromstring(content)
                    transform(tree)
                    content = ET.tostring(tree, encoding="utf-8")
                output.writestr(entry, content)
        return target

    # --- matching structure -------------------------------------------------
    def test_composite_customer_categories_and_column_order(self):
        left = self.workbook("L1.xlsx", [("甲", "D1", "1"), ("甲", "D2", "1.001"),
                                        ("乙", "D1", "2"), ("丙", "X", "3")])
        right = self.workbook("R1.xlsx", [("D1", "甲", "1"), ("D2", "甲", "1.002"),
                                         ("D1", "乙", "2"), ("Y", "丁", "4")],
                              headers=("流水号", "买方", "货款"))
        left_rules = Rules("客户", "金额", None, 1, ("客户", "单据号"))
        right_rules = Rules("买方", "货款", None, 1, ("买方", "流水号"))
        result = reconciliation.run_reconciliation([left], [right], left_rules, right_rules)
        rows = {row.parts: row for row in result.rows}
        self.assertEqual(rows[("甲", "D1")].category, "consistent")
        self.assertEqual(rows[("甲", "D2")].category, "different")
        self.assertEqual(rows[("甲", "D2")].difference, Decimal("-0.001"))
        self.assertEqual(rows[("乙", "D1")].category, "consistent")
        self.assertEqual(rows[("丙", "X")].category, "left_only")
        self.assertEqual(rows[("丁", "Y")].category, "right_only")
        self.assertFalse(result.all_consistent)
        # Display is decomposable JSON; matching used the structured tuple above.
        self.assertEqual(rows[("甲", "D1")].key, '["甲", "D1"]')

    def test_composite_three_fields_with_mixed_headers_and_indices(self):
        left = self.workbook("L3.xlsx", [("甲", "D1", "B1", "1"), ("甲", "D1", "B2", "2"),
                                        ("乙", "D1", "B1", "0")],
                             headers=("客户", "单据号", "批次", "金额"))
        right = self.workbook("R3.xlsx", [("B1", "D1", "甲", "1"), ("B2", "D1", "甲", "2"),
                                         ("B1", "D1", "乙", "0")],
                              headers=("批次", "流水号", "买方", "货款"))
        left_rules = Rules(1, 4, None, 1, (1, 2, 3))            # #1/#2/#3, amount #4
        right_rules = Rules("买方", "货款", None, 1, ("买方", "流水号", "批次"))
        result = reconciliation.run_reconciliation([left], [right], left_rules, right_rules)
        self.assertTrue(result.all_consistent, [(r.key, r.category) for r in result.rows])
        self.assertEqual({row.parts for row in result.rows},
                         {("甲", "D1", "B1"), ("甲", "D1", "B2"), ("乙", "D1", "B1")})

    def test_separator_order_and_text_exactness(self):
        report = self.read(
            self.workbook("sep.xlsx", [("A|B", "C", "1"), ("A", "B|C", "2"), ("甲", "001", "3"),
                                       ("甲", " 1 ", "4"), (" 甲 ", " D1 ", "5")]),
            Rules("客户", "金额", None, 1, ("客户", "单据号")))
        self.assertTrue(report.complete, report.issues)
        parts = [record.parts for record in report.records]
        self.assertEqual(parts, [("A|B", "C"), ("A", "B|C"), ("甲", "001"), ("甲", "1"), ("甲", "D1")])
        # Collision and order never merge into one group.
        self.assertNotEqual(report.records[0].key, report.records[1].key)
        a = self.read(self.workbook("oa.xlsx", [("A", "B", "1")]),
                      Rules("客户", "金额", None, 1, ("客户", "单据号")))
        b = self.read(self.workbook("ob.xlsx", [("A", "B", "1")]),
                      Rules("单据号", "金额", None, 1, ("单据号", "客户")))
        flipped = reconciliation.compare_reports(a, b)
        self.assertEqual({row.category for row in flipped.rows}, {"left_only", "right_only"})
        self.assertFalse(flipped.all_consistent)

    def test_invalid_components_keep_raw_parts(self):
        path = self.workbook("bad.xlsx", [
            (None, "D1", "1"), (True, "D1", "2"), (datetime.datetime(2026, 10, 8), "D1", "3"),
            ("甲", None, "4"), ("乙", "D2", "=SUM(1,2)"),
        ])
        report = self.read(path, Rules("客户", "金额", None, 1, ("客户", "单据号")))
        self.assertFalse(report.records)
        self.assertEqual(report.invalid_rows, 5)
        messages = [issue.message for issue in report.issues]
        self.assertIn("第1个键字段缺失或为布尔值", messages[0])
        self.assertIn("第1个键字段缺失或为布尔值", messages[1])
        self.assertIn("Excel 错误值或日期", messages[2])
        self.assertIn("第2个键字段缺失或为布尔值", messages[3])
        self.assertIn("公式没有缓存", messages[4])
        self.assertEqual(report.issues[0].raw_key, '["", "D1"]')
        self.assertEqual(report.issues[3].raw_key, '["甲", ""]')

    def test_duplicate_and_overlap_columns_rejected(self):
        path = self.workbook("dup.xlsx", [("甲", "D1", "1")])
        dup = self.read(path, Rules("客户", "金额", None, 1, ("客户", 1)))   # name and #1 -> same column
        self.assertTrue(dup.issues)
        self.assertIn("同一列", dup.issues[0].message)
        overlap = self.read(path, Rules("客户", "金额", None, 1, ("客户", 3)))  # key column == amount column
        self.assertTrue(overlap.issues)
        self.assertIn("金额字段", overlap.issues[0].message)
        named = self.workbook("named.xlsx", [("甲", "D1", "1")], headers=("1", "单据号", "金额"))
        clash = self.read(named, Rules("1", "金额", None, 1, ("1", 1)))    # header "1" and #1 -> same column
        self.assertTrue(clash.issues)
        self.assertIn("同一列", clash.issues[0].message)

    def test_field_count_rules_and_left_right_length(self):
        for bad in ((), ("客户", "客户"), ("客户", True), tuple(f"k{i}" for i in range(9))):
            with self.assertRaises(ValueError):
                Rules("客户", "金额", None, 1, bad)
        Rules("客户", "金额", None, 1, tuple(f"k{i}" for i in range(8)))  # 8 is allowed
        with self.assertRaises(ValueError):
            reconciliation.run_reconciliation([], [], Rules("客户", "金额"),
                                              Rules("买方", "货款", None, 1, ("买方", "流水号")))

    # --- document mode ------------------------------------------------------
    def test_document_duplicates_traceable_and_shared_document_numbers(self):
        rules = Rules("客户", "金额", None, 1, ("客户", "单据号"))
        left = self.read(self.workbook("dl.xlsx", [("甲", "D1", "1"), ("甲", "D1", "2"),
                                                   ("乙", "D1", "3")]), rules)
        right = self.read(self.workbook("dr.xlsx", [("甲", "D1", "5"), ("乙", "D1", "3")]), rules)
        result = reconciliation.compare_reports(left, right, mode="document")
        rows = {row.parts: row for row in result.rows}
        self.assertEqual(rows[("甲", "D1")].category, "ambiguous")
        self.assertEqual(len(rows[("甲", "D1")].left) + len(rows[("甲", "D1")].right), 3)
        self.assertIsNone(rows[("甲", "D1")].difference)
        self.assertEqual(rows[("乙", "D1")].category, "consistent")  # same document, other customer
        self.assertFalse(result.all_consistent)
        self.assertFalse(result.complete)
        classified = [(side, record.source) for row in result.rows
                      for side, records in (("L", row.left), ("R", row.right)) for record in records]
        self.assertEqual(len(classified), len(set(classified)))
        self.assertEqual(len(classified), len(left.records) + len(right.records))

    def test_single_and_composite_never_mismatch(self):
        single = self.read(self.workbook("s.xlsx", [("甲", "1")], headers=("客户", "金额")),
                           Rules("客户", "金额"))
        composite = self.read(self.workbook("c.xlsx", [("甲", "D1", "1")]),
                              Rules("客户", "金额", None, 1, ("客户", "单据号")))
        result = reconciliation.compare_reports(single, composite)
        self.assertEqual({row.category for row in result.rows}, {"left_only", "right_only"})
        self.assertFalse(result.all_consistent)

    # --- templates and export ----------------------------------------------
    def test_legacy_snapshot_and_composite_template_roundtrip(self):
        self.assertEqual(set(Rules().snapshot()), {"key_column", "amount_column", "sheet", "header_row"})
        composite = Rules("客户", "金额", None, 1, ("客户", "单据号"))
        self.assertEqual(composite.snapshot()["key_fields"], ("客户", "单据号"))
        path = self.root / "config" / "monthly_templates.json"
        monthly.save_template(path, monthly.make_template("组合模板", composite))
        loaded = monthly.load_templates(path)["组合模板"]
        self.assertEqual(loaded["rules"]["key_fields"], ["客户", "单据号"])
        self.assertEqual(analysis.SheetRules(**loaded["rules"]).field_selections(), ("客户", "单据号"))
        legacy = monthly.make_template("旧模板", Rules("客户", "金额"))
        self.assertEqual(set(legacy["rules"]), {"key_column", "amount_column", "sheet", "header_row"})
        raw = self.root / "legacy.json"
        raw.write_text(json.dumps({"schema_version": 1, "templates": [{
            "schema_version": 1, "name": "旧", "output_prefix": "月度汇总",
            "rules": {"key_column": "客户", "amount_column": "金额", "sheet": None, "header_row": 1},
        }]}, ensure_ascii=False), encoding="utf-8")
        self.assertEqual(monthly.load_templates(raw)["旧"]["rules"]["key_column"], "客户")

    def test_composite_export_traceability_and_config(self):
        rules = Rules("客户", "金额", None, 1, ("客户", "单据号"))
        template = monthly.make_template("组合月报", rules)
        source = self.workbook("src.xlsx", [("甲", "D1", "1.001"), ("甲", "D1", "0.002"),
                                            ("乙", "D2", "-1"), ("bad", "D3", "NaN")])
        before = hashlib.sha256(source.read_bytes()).digest()
        result = monthly.run_monthly([source], template, "2026-10")
        rows = {row.parts: row for row in result.rows}
        self.assertEqual(rows[("甲", "D1")].amount, Decimal("1.003"))
        self.assertEqual(rows[("乙", "D2")].amount, Decimal("-1"))
        output = export.export_result(result, self.root, prefix="组合")
        self.assertEqual(hashlib.sha256(source.read_bytes()).digest(), before)
        book = load_workbook(output, data_only=False)
        self.addCleanup(book.close)
        config = dict(book["运行配置"].values)
        self.assertEqual(config["key_fields[1]"], "客户")
        self.assertEqual(config["key_fields[2]"], "单据号")
        self.assertEqual(book["来源明细"]["A1"].value, "侧别")
        self.assertEqual(book["来源明细"].max_row, 4)  # header + 3 valid records
        self.assertEqual(book["来源明细"]["F2"].data_type, "s")
        for sheet in book:
            self.assertFalse(any(cell.data_type == "f" for row in sheet for cell in row))

    def test_xls_composite_precision_negative_and_tolerance(self):
        rules = Rules("客户", "金额", None, 1, ("客户", "单据号"))
        xls = self.workbook("c.xls", [("甲", "D1", 0.1), ("甲", "D1", 0.2)])
        report = self.read(xls, rules)
        self.assertTrue(report.complete, report.issues)
        self.assertEqual(report.records[0].parts, ("甲", "D1"))
        result = monthly.run_monthly([xls], monthly.make_template("X", rules), "2026-10")
        self.assertEqual(result.total, Decimal("0.3"))
        left = self.read(self.workbook("nl.xlsx", [("甲", "D1", "-1.001"), ("乙", "Z", "0")]), rules)
        right = self.read(self.workbook("nr.xlsx", [("甲", "D1", "-1.002"), ("乙", "Z", "0")]), rules)
        self.assertFalse(reconciliation.compare_reports(left, right, tolerance="0").all_consistent)
        self.assertTrue(reconciliation.compare_reports(left, right, tolerance="0.001").all_consistent)


    # --- formats, special characters and extra-component errors -------------
    def test_composite_reads_all_three_formats(self):
        for extension in ("xlsx", "xlsm", "xls"):
            with self.subTest(extension=extension):
                path = self.workbook("fmt." + extension, [("甲", "D1", "1"), ("乙", "D2", "0")],
                                     headers=("客户", "单据号", "金额"))
                report = self.read(path, Rules("客户", "金额", None, 1, ("客户", "单据号")))
                self.assertTrue(report.complete, report.issues)
                self.assertEqual([r.parts for r in report.records], [("甲", "D1"), ("乙", "D2")])

    def test_composite_special_characters_and_case_exactness(self):
        left = self.workbook(
            "spec-l.xlsx",
            [("甲,乙", 'D"1', "B\n1", "1"), ("AA", "d", "b", "2"), ("aa", "d", "b", "3")],
            headers=("客户,名", '单据"号', "批次\n号", "金额"))
        right = self.workbook(
            "spec-r.xlsx",
            [("B\n1", "甲,乙", 'D"1', "1"), ("b", "AA", "d", "2"), ("b", "aa", "d", "3")],
            headers=("批|次", "买方", "流水", "货款"))
        left_rules = Rules("客户,名", "金额", None, 1, ("客户,名", '单据"号', "批次\n号"))
        right_rules = Rules("买方", "货款", None, 1, ("买方", "流水", "批|次"))
        result = reconciliation.run_reconciliation([left], [right], left_rules, right_rules)
        self.assertTrue(result.all_consistent, [(r.key, r.category) for r in result.rows])
        self.assertEqual({row.parts for row in result.rows},
                         {("甲,乙", 'D"1', "B\n1"), ("AA", "d", "b"), ("aa", "d", "b")})
        # Case and embedded separators never collapse into one key.
        self.assertNotEqual(("AA", "d", "b"), ("aa", "d", "b"))
        self.assertNotEqual(("A|B", "C"), ("A", "B|C"))

    def test_composite_extra_component_error_and_formula(self):
        namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        rules = Rules("客户", "金额", None, 1, ("客户", "单据号"))
        # Second key component is an Excel error value.
        error_path = self.workbook("err.xlsx", [("甲", None, "1")])
        book = load_workbook(error_path)
        book.active["B2"] = "#N/A"
        book.active["B2"].data_type = "e"
        book.save(error_path)
        book.close()
        report = self.read(error_path, rules)
        self.assertEqual(report.invalid_rows, 1)
        self.assertIn("Excel 错误", report.issues[0].message)
        self.assertEqual(report.issues[0].raw_key, '["甲", "#N/A"]')
        # Second key component is a formula without a cached value.
        uncached = self.read(self.workbook("nc.xlsx", [("甲", "=SUM(1,2)", "1")]), rules)
        self.assertEqual(uncached.invalid_rows, 1)
        self.assertIn("公式没有缓存", uncached.issues[0].message)
        # Second key component is a formula whose cached value is an Excel error.
        original = self.workbook("ce.xlsx", [("甲", "=NA()", "1")])
        def transform(tree):
            cell = tree.find(".//" + namespace + "c[@r='B2']")
            cell.set("t", "e")
            cell.find(namespace + "v").text = "#N/A"
        cached = self.read(self.rewrite_xml(original, "cached-err.xlsx", transform), rules)
        self.assertEqual(cached.invalid_rows, 1)
        self.assertIn("Excel 错误", cached.issues[0].message)
        # raw_key keeps the original formula text; the error is taken from the cached value.
        self.assertEqual(cached.issues[0].raw_key, '["甲", "=NA()"]')

    def test_composite_customer_aggregation_and_document_ambiguity_both_sides(self):
        rules = Rules("客户", "金额", None, 1, ("客户", "单据号"))
        left = self.read(self.workbook("agg-l.xlsx", [("甲", "D1", "0.1"), ("甲", "D1", "0.2"),
                                                      ("乙", "D1", "1"), ("甲", "D2", "5")]), rules)
        right = self.read(self.workbook("agg-r.xlsx", [("甲", "D1", "0.3"), ("乙", "D1", "1"),
                                                       ("甲", "D2", "5")]), rules)
        customer = reconciliation.compare_reports(left, right, mode="customer")
        rows = {row.parts: row for row in customer.rows}
        self.assertEqual(rows[("甲", "D1")].category, "consistent")
        self.assertEqual(rows[("甲", "D1")].left_amount, Decimal("0.3"))  # whole combo aggregated
        self.assertEqual(len(rows[("甲", "D1")].left), 2)                 # every source row kept
        self.assertEqual(customer.counts["consistent"], 3)
        self.assertTrue(customer.all_consistent)

        # Document mode: the same complete combo repeats on BOTH sides.
        dla = self.read(self.workbook("d-l.xlsx", [("甲", "D1", "1"), ("甲", "D1", "2")]), rules)
        dra = self.read(self.workbook("d-r.xlsx", [("甲", "D1", "3"), ("甲", "D1", "4")]), rules)
        document = reconciliation.compare_reports(dla, dra, mode="document")
        duplicate = {row.parts: row for row in document.rows}[("甲", "D1")]
        self.assertEqual(duplicate.category, "ambiguous")
        self.assertEqual((len(duplicate.left), len(duplicate.right)), (2, 2))
        self.assertIsNone(duplicate.difference)
        self.assertFalse(document.complete)
        self.assertFalse(document.all_consistent)
        classified = [(side, record.source) for row in document.rows
                      for side, records in (("L", row.left), ("R", row.right)) for record in records]
        self.assertEqual(len(classified), len(set(classified)))
        self.assertEqual(len(classified), len(dla.records) + len(dra.records))


if __name__ == "__main__":
    unittest.main()
