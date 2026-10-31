from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

# ========== 应用名称与版本（唯一来源，全工程引用） ==========
APP_NAME = "E助手"
APP_VERSION = "v1.5.4"
APP_FULL_NAME = f"{APP_NAME} {APP_VERSION}"

# ========== 路径常量 ==========
FOLDER_NAME = "对账单"
OUTPUT_FILE = "汇总.xlsx"
OUTPUT_FOLDER_NAME = "Excel汇总结果"
ERROR_LOG = "error.log"

# ========== 表头常量 ==========
HEADERS = ["客户", "总金额"]

# ========== 表头样式：浅绿色背景 + 加粗 + 居中 + 边框 ==========
header_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
header_font = Font(bold=True, size=12)
header_alignment = Alignment(horizontal="center", vertical="center")

# ========== 通用边框 ==========
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

# ========== 合计行样式：浅黄色背景 + 加粗 ==========
sum_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
sum_font = Font(bold=True, size=11)

# ========== 字段识别关键词 ==========
# 客户字段（按优先级排列，常用词在前）
CUSTOMER_KEYWORDS = [
    "客户",
    "客户名称",
    "购货单位",
    "单位名称",
    "公司",
    "顾客",
    "收货人",
    "买方",
]

# 金额字段（按优先级排列，常用词在前）
MONEY_KEYWORDS = [
    "金额",
    "总金额",
    "合计金额",
    "价税合计",
    "货款",
    "费用",
    "小计",
    "应收金额",
]

# ========== 自定义表头映射 ==========
# 当自动关键词匹配无法识别你的 Excel 表头时，在这里添加映射即可。
# 格式：{"你的表头名": "customer" 或 "money"}
#
# 示例：如果你的 Excel 表头是 "姓名" 和 "价格"：
#   CUSTOM_HEADER_MAP = {
#       "姓名": "customer",
#       "价格": "money",
#   }
CUSTOM_HEADER_MAP = {
    # "姓名": "customer",
    # "价格": "money",
}
