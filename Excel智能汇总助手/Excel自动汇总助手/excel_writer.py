from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from .edition_policy import guarded_export
from .config import (
    HEADERS,
    header_fill,
    header_font,
    header_alignment,
    thin_border,
    sum_fill,
    sum_font,
)


@guarded_export
def generate_summary(sorted_data, total_dict, output_path, field_names=None):
    """生成格式化汇总表并保存（v1.5.0 动态字段版）。

    Args:
        sorted_data: [(分组键, {字段名: 值, ...}), ...] 已排序的汇总数据
        total_dict: {分组键: {字段名: 值, ...}} 原始汇总字典
        output_path: 输出文件完整路径
        field_names: 字段名列表（用于表头，默认使用 config.HEADERS）
    """
    if field_names is None:
        field_names = list(HEADERS)
    if len(field_names) < 2:
        raise ValueError("汇总至少需要分组字段和金额字段。")

    new_wb = Workbook()
    new_ws = new_wb.active

    # 写入表头
    def append_literal(values):
        if any(isinstance(value, str) and len(value) > 32767 for value in values):
            raise ValueError("单元格文字超出 Excel 长度限制，未导出。")
        new_ws.append(values)
        for cell in new_ws[new_ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    append_literal(field_names)

    # ---------- 表头美化 ----------
    for cell in new_ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # ---------- 写入数据 ----------
    # 求和字段索引（第 2 列，用于格式化和合计）
    sum_col_idx = 1  # 0-based，第 2 个字段

    for key, row_data in sorted_data:
        row_values = [row_data.get(fn, "") for fn in field_names]
        append_literal(row_values)

    # ---------- 合计行 ----------
    sum_total = 0
    for row_data in total_dict.values():
        if len(field_names) >= 2:
            v = row_data.get(field_names[1], 0)
        else:
            v = 0
        if isinstance(v, (int, float)):
            sum_total += v
        else:
            try:
                sum_total += float(v)
            except (ValueError, TypeError):
                pass

    sum_row = ["合计"] + [""] * (len(field_names) - 1)
    sum_row[sum_col_idx] = sum_total
    append_literal(sum_row)

    # ---------- 金额列格式化为千分位两位小数 ----------
    sum_col_letter = get_column_letter(sum_col_idx + 1)
    for cell in new_ws[sum_col_letter]:
        cell.number_format = "#,##0.00"

    # ---------- 合计行美化 ----------
    last_row = new_ws.max_row
    for col_idx in range(len(field_names)):
        col_letter = get_column_letter(col_idx + 1)
        cell = new_ws[f"{col_letter}{last_row}"]
        cell.font = sum_font
        cell.fill = sum_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # ---------- 数据区域边框 ----------
    for row in new_ws.iter_rows(min_row=2, max_row=last_row - 1,
                                 min_col=1, max_col=len(field_names)):
        for cell in row:
            cell.border = thin_border
            if cell.column == sum_col_idx + 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # ---------- 自动列宽 ----------
    for column in new_ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            if cell.value:
                value_length = len(str(cell.value))
                if column_letter == sum_col_letter:
                    value_length += 4
                max_length = max(max_length, value_length)
        new_ws.column_dimensions[column_letter].width = max_length + 4

    # ---------- 保存 ----------
    new_wb.save(output_path)
